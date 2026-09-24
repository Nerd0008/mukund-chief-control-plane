#!/usr/bin/env python3
"""Company Watch test suite.

Runs fully offline: every network path is exercised through
``ats_endpoints.FakeFetcher``. The only real workbooks touched are copies made by
the tests themselves; the canonical regional workbooks are opened read-only.

    python -m pytest company-watch/tests/test_company_watch.py -v
"""

from __future__ import annotations

import datetime as dt
import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
WATCH_DIR = HERE.parent
REPO_ROOT = WATCH_DIR.parent
sys.path.insert(0, str(WATCH_DIR))

import ats_endpoints  # noqa: E402
import company_registry  # noqa: E402
import company_watch as cw  # noqa: E402

CFG = cw.load_config()
TW, PROFILES, UK_CFG = cw.resolve_region_config(CFG, "uk")
CANONICAL_UK = Path(UK_CFG["tracker"])

FIXTURE_HISTORY = """# UK job-application company history — past 18 months

**Mailbox reviewed:** test@example.invalid
**Period:** 1 January 2026 to 1 February 2026
**Scope:** synthetic fixture used by the Company Watch test suite.

## Summary

- **Confirmed or strongly evidenced employers:** 2
- **Recruiters / intermediaries that received applications or CVs:** 1
- **Total job-search organisations with application/CV evidence:** 3
- **Other non-employment applications:** 1
- **Possible applications needing verification:** 1

## Confirmed or strongly evidenced employers (2)

- Testco
- Example Bank — rejected after interview

## Recruiters and intermediaries (1)

- Test Recruiters

## Other applications kept separate (1)

- Example Charity — volunteering application

## Possible applications needing verification (1)

- Example Motors — candidate-account creation only; no submission confirmation found
"""


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #

def gh_payload(name: str, jobs: list[dict]) -> str:
    return json.dumps({
        "jobs": [
            {
                "id": j["id"],
                "title": j["title"],
                "absolute_url": j.get("url", f"https://job-boards.greenhouse.io/testco/jobs/{j['id']}"),
                "company_name": j.get("company_name", name),
                "location": {"name": j.get("location", "London, United Kingdom")},
                "updated_at": j.get("updated_at", "2026-09-20T10:00:00-04:00"),
                "first_published": j.get("posted_at", "2026-09-20T10:00:00-04:00"),
            }
            for j in jobs
        ]
    })


def registry_company(name: str = "Testco", **kw) -> dict:
    return {
        "name": name,
        "evidence_class": kw.get("evidence_class", "employer"),
        "prior_application_evidence": True,
        "name_variants": company_registry.name_variants(name),
        "ats_slug_candidates": company_registry.slug_candidates(name),
    }


def copy_uk_tracker(tmp_path: Path) -> Path:
    dest = tmp_path / "uk-cyber-job-tracker.xlsx"
    shutil.copy2(CANONICAL_UK, dest)
    return dest


def write_manifest_cli(tmp_path: Path, records: list[dict], tracker: Path,
                       archive_dir: Path, apply: bool, region: str = "uk") -> dict:
    manifest = tmp_path / f"manifest-{region}.json"
    manifest.write_text(json.dumps({"records": records}), encoding="utf-8")
    cmd = [sys.executable, str(REPO_ROOT / "career-ops" / "career_ops_cli.py"),
           "write", "--region", region, "--manifest", str(manifest),
           "--tracker", str(tracker), "--backup-dir", str(tmp_path / "backups"),
           "--archive-dir", str(archive_dir)]
    if apply:
        cmd.append("--apply")
    proc = subprocess.run(cmd, cwd=str(REPO_ROOT), capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=300)
    assert proc.stdout.strip(), proc.stderr
    return json.loads(proc.stdout)


# --------------------------------------------------------------------------- #
# registry
# --------------------------------------------------------------------------- #

def test_registry_parses_fixture_and_flags_unverified(tmp_path):
    src = tmp_path / "history.md"
    src.write_text(FIXTURE_HISTORY, encoding="utf-8")
    reg = company_registry.build_registry(src)

    counts = reg["counts"]
    assert counts["parsed_employers"] == 2
    assert counts["parsed_recruiters"] == 1
    assert counts["parsed_other_applications"] == 1
    assert counts["parsed_unverified_candidates"] == 1
    assert counts["parsed_organizations_with_application_or_cv_evidence"] == 3
    assert counts["count_mismatch"] == {}

    by_name = {c["name"]: c for c in reg["companies"]}
    assert by_name["Testco"]["evidence_class"] == "employer"
    assert by_name["Testco"]["prior_application_evidence"] is True
    assert by_name["Example Bank"]["evidence_note"] == "rejected after interview"
    assert by_name["Test Recruiters"]["evidence_class"] == "recruiter"
    assert by_name["Example Motors"]["evidence_class"] == "unverified_candidate"
    assert by_name["Example Motors"]["prior_application_evidence"] is False


def test_registry_detects_declared_count_mismatch(tmp_path):
    src = tmp_path / "history.md"
    src.write_text(FIXTURE_HISTORY.replace(
        "**Confirmed or strongly evidenced employers:** 2",
        "**Confirmed or strongly evidenced employers:** 9"), encoding="utf-8")
    reg = company_registry.build_registry(src)
    assert reg["counts"]["count_mismatch"]["confirmed_employers"] == {"declared": 9, "parsed": 2}


def test_registry_never_invents_companies(tmp_path):
    src = tmp_path / "history.md"
    src.write_text(FIXTURE_HISTORY, encoding="utf-8")
    reg = company_registry.build_registry(src)
    names = {c["name"] for c in reg["companies"]}
    assert names == {"Testco", "Example Bank", "Test Recruiters", "Example Charity", "Example Motors"}
    assert all(c["source_section"] for c in reg["companies"])
    assert reg["source"]["sha256"] == company_registry.sha256_file(src)


@pytest.mark.skipif(not Path(CFG["registry_source"]).exists(),
                    reason="owner company-history source not present on this machine")
def test_real_source_reproduces_historical_counts():
    reg = company_registry.build_registry(Path(CFG["registry_source"]))
    counts = reg["counts"]
    assert counts["parsed_employers"] == 203
    assert counts["parsed_recruiters"] == 18
    assert counts["parsed_organizations_with_application_or_cv_evidence"] == 221
    assert counts["count_mismatch"] == {}
    assert counts["parsed_total_rows"] == 229


def test_name_variants_and_slug_candidates():
    assert company_registry.name_variants("AVEVA / ETAP / RIB")[0] == "avevaetaprib"
    assert "aveva" in company_registry.name_variants("AVEVA / ETAP / RIB")
    assert company_registry.slug_candidates("Hewlett Packard Enterprise Limited") == [
        "hewlettpackardenterpriselimited", "hewlettpackardenterprise"]
    assert company_registry.slug_candidates("8x8") == ["8x8"]


# --------------------------------------------------------------------------- #
# matching primitives
# --------------------------------------------------------------------------- #

def test_word_boundary_matching_avoids_substring_false_positives():
    assert cw.term_in_text("Intern", "Cybersecurity Intern")
    assert not cw.term_in_text("Intern", "Internal Audit Manager")
    assert cw.term_in_text("UK", "Remote, UK")
    assert not cw.term_in_text("UK", "Ukraine")
    assert not cw.term_in_text("Lead ", "Team Leader")


def test_url_classification():
    assert cw.classify_url("https://job-boards.greenhouse.io/x/jobs/1", "vendor") == "posting"
    assert cw.classify_url("https://api.smartrecruiters.com/v1/companies/x/postings/1",
                           "vendor") == "api_endpoint"
    assert cw.classify_url("https://jobs.smartrecruiters.com/x/1", "derived_pattern") == "derived_public"
    assert cw.classify_url(None, "vendor") == "missing"
    assert cw.classify_url("not-a-url", "vendor") == "missing"


def test_route_region_uses_owner_terms_and_refuses_unknowns():
    filters = cw.load_owner_filters(CFG, "uk")
    assert cw.route_region("London, United Kingdom", CFG, filters, "uk") == "uk"
    assert cw.route_region("Dubai", CFG) == "dubai"
    assert cw.route_region("Tokyo, Japan", CFG) == "japan"
    assert cw.route_region("Singapore", CFG) == "singapore"
    assert cw.route_region("Milton Keynes", CFG) is None
    assert cw.route_region("London, United States", CFG) is None


def test_owner_filters_come_from_the_career_ops_install():
    filters = cw.load_owner_filters(CFG, "uk")
    assert filters["available"] is True
    assert filters["path"].endswith("portals.yml")
    assert "Intern" in filters["title_positive"]
    assert filters["max_posting_age_days"] == 14
    # regions without a lane config must not inherit the UK filters
    for region in ("dubai", "japan", "singapore"):
        assert cw.load_owner_filters(CFG, region)["available"] is False


def test_minimal_portals_parser_matches_pyyaml_for_filters():
    text = (Path(CFG["career_ops_root"]) / "portals.yml").read_text(encoding="utf-8")
    minimal = cw._parse_portals_filters_minimal(text)
    try:
        import yaml
    except ImportError:  # pragma: no cover
        pytest.skip("PyYAML not installed")
    parsed, _ = cw._parse_portals_filters(text, Path("portals.yml"))
    assert minimal["title_positive"] == parsed["title_positive"]
    assert minimal["title_negative"] == parsed["title_negative"]
    assert minimal["location_allow"] == parsed["location_allow"]
    assert minimal["location_block"] == parsed["location_block"]
    assert minimal["max_posting_age_days"] == parsed["max_posting_age_days"]
    assert yaml.safe_load(text)["title_filter"]["positive"] == parsed["title_positive"]


# --------------------------------------------------------------------------- #
# ATS resolution / attribution
# --------------------------------------------------------------------------- #

def test_resolution_requires_name_confirmation():
    company = registry_company("Testco")
    jobs = [{"id": 1, "title": "Security Intern"}]
    fetch = ats_endpoints.FakeFetcher({
        "https://boards-api.greenhouse.io/v1/boards/testco/jobs": (200, gh_payload("Testco", jobs)),
    })
    rec = ats_endpoints.resolve_company(company, fetch)
    assert rec["resolved"] is True
    assert rec["attribution_confidence"] == "high"
    assert rec["attribution_basis"] == "vendor_payload_company_name:Testco"
    assert rec["jobs"][0]["title"] == "Security Intern"

    # Payload name belongs to a different organisation → never attributed.
    fetch2 = ats_endpoints.FakeFetcher({
        "https://boards-api.greenhouse.io/v1/boards/wise/jobs": (
            200, gh_payload("Wise Worksite Field Sales", jobs)),
        "https://boards.greenhouse.io/wise": (200, "<html>Wise Worksite Field Sales</html>"),
    })
    rec2 = ats_endpoints.resolve_company(registry_company("Wise"), fetch2)
    assert rec2["resolved"] is True
    assert rec2["attribution_confidence"] == "low"
    assert rec2["manual_attribution_required"] is True
    assert rec2["attribution_basis"] == "vendor_payload_company_name_mismatch:Wise Worksite Field Sales"
    assert rec2["jobs"] == []
    assert rec2["jobs_discarded_attribution_unconfirmed"] == 1


def test_board_page_confirmation_and_short_name_refusal():
    jobs = [{"id": 5, "title": "Junior Analyst"}]
    page = "<html><body><h1>Testco careers</h1></body></html>"
    fetch = ats_endpoints.FakeFetcher({
        "https://api.ashbyhq.com/posting-api/job-board/testco": (
            200, json.dumps({"jobs": [{"id": "5", "title": "Junior Analyst",
                                       "location": "London", "jobUrl": "https://jobs.ashbyhq.com/testco/5",
                                       "publishedAt": "2026-09-19T00:00:00Z"}]})),
        "https://jobs.ashbyhq.com/testco": (200, page),
    })
    rec = ats_endpoints.resolve_company(registry_company("Testco"), fetch)
    assert rec["attribution_confidence"] == "high"
    assert rec["attribution_basis"].startswith("board_page_name_match")
    assert len(rec["jobs"]) == 1

    # A short acronym cannot be confirmed from page text.
    fetch2 = ats_endpoints.FakeFetcher({
        "https://api.ashbyhq.com/posting-api/job-board/ey": (
            200, json.dumps({"jobs": [{"id": "1", "title": "Analyst", "location": "London",
                                       "jobUrl": "https://jobs.ashbyhq.com/ey/1"}]})),
        "https://jobs.ashbyhq.com/ey": (200, "<html>they are hiring analysts</html>"),
    })
    rec2 = ats_endpoints.resolve_company(registry_company("EY"), fetch2)
    assert rec2["attribution_confidence"] == "low"
    assert rec2["manual_attribution_required"] is True


def test_resolution_reports_unresolved_without_claiming_absence():
    fetch = ats_endpoints.FakeFetcher({})
    rec = ats_endpoints.resolve_company(registry_company("Nonexistent"), fetch)
    assert rec["resolved"] is False
    assert rec["unresolved_reason"] == "no_vendor_board_resolved_via_slug_guess"
    assert rec["jobs"] == []
    assert all(a["outcome"] == "no_board" for a in rec["attempts"])


def test_resolution_survives_malformed_payload():
    fetch = ats_endpoints.FakeFetcher({
        "https://boards-api.greenhouse.io/v1/boards/testco/jobs": (200, "not json"),
    })
    rec = ats_endpoints.resolve_company(registry_company("Testco"), fetch)
    assert rec["resolved"] is False


def test_budget_stops_the_sweep():
    fetch = ats_endpoints.FakeFetcher({})
    companies = [registry_company(f"Co{i}") for i in range(10)]
    ticks = iter([0, 0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 100, 101])
    doc = ats_endpoints.resolve_companies(companies, fetch, budget_s=5, clock=lambda: next(ticks))
    assert doc["budget_exhausted"] is True
    assert doc["companies_probed"] < 10


# --------------------------------------------------------------------------- #
# findings + shared dedupe
# --------------------------------------------------------------------------- #

def synthetic_resolution(jobs: list[dict], company: str = "Testco") -> dict:
    return {
        "resolved": [{
            "company": company,
            "evidence_class": "employer",
            "vendor": "greenhouse",
            "board_slug": company.casefold(),
            "board_url": f"greenhouse board '{company.casefold()}'",
            "jobs_listed": len(jobs),
            "attribution_confidence": "high",
            "attribution_basis": f"vendor_payload_company_name:{company}",
            "manual_attribution_required": False,
            "jobs": [{**j, "vendor": "greenhouse", "board_slug": company.casefold()} for j in jobs],
        }],
        "unresolved": [],
    }


FULL_FILTERS = {
    "available": True,
    "location_always_allow": ["United Kingdom", "UK"],
    "location_allow": ["United Kingdom", "England", "Scotland", "Wales", "Northern Ireland", "UK"],
    "location_block": ["United States", "USA", "Canada", "India", "Singapore", "Australia",
                       "Dubai", "United Arab Emirates", "UAE"],
    "title_positive": ["Intern", "Internship"],
    "title_negative": ["Senior", "Principal", "Lead ", "Manager", "Director", "Head of"],
    "max_posting_age_days": 14,
    "path": "portals.yml",
}
NOW = dt.datetime(2026, 9, 23, 12, 0, tzinfo=dt.timezone.utc)


def make_finding_doc(tmp_path, jobs, *, filters=FULL_FILTERS, tracker=None, company="Testco"):
    tracker = tracker or copy_uk_tracker(tmp_path)
    dedupe = cw.build_shared_dedupe(CFG, "uk", tracker_override=str(tracker))
    res = synthetic_resolution(jobs, company=company)
    return cw.build_findings(res, region="uk", cfg=CFG, dedupe=dedupe, tw=TW,
                             filters=filters, now=NOW), dedupe


def test_findings_apply_owner_filters_and_age(tmp_path):
    jobs = [
        {"job_id": "1", "title": "Cybersecurity Intern", "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/1",
         "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"},
        {"job_id": "2", "title": "Internal Audit Manager", "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/2",
         "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"},
        {"job_id": "3", "title": "Security Internship", "location": "New York, United States",
         "url": "https://job-boards.greenhouse.io/testco/jobs/3",
         "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"},
        {"job_id": "4", "title": "Security Internship", "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/4",
         "posted_at": "2026-01-01T00:00:00+00:00", "url_source": "vendor"},
    ]
    doc, _ = make_finding_doc(tmp_path, jobs)
    by_id = {f["vendor_job_id"]: f for f in doc["findings"]}
    assert by_id["1"]["owner_filter_eligible"] is True
    assert by_id["1"]["tracker_eligible"] is True
    assert by_id["2"]["title_rule_pass"] is False          # "Internal Audit" != "Intern"
    assert by_id["3"]["location_rule_pass"] is False       # blocked country
    assert by_id["3"]["region_route"] is None
    assert by_id["4"]["age_rule_pass"] is False            # older than the 14-day owner rule
    assert by_id["4"]["tracker_eligible"] is False
    assert doc["counts"]["owner_filter_eligible"] == 1
    assert doc["counts"]["tracker_eligible"] == 1


def test_findings_never_write_api_endpoints_to_a_tracker(tmp_path):
    jobs = [{"job_id": "9", "title": "Security Intern", "location": "London, United Kingdom",
             "url": "https://api.smartrecruiters.com/v1/companies/testco/postings/9",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc, _ = make_finding_doc(tmp_path, jobs)
    f = doc["findings"][0]
    assert f["url_quality"] == "api_endpoint"
    assert f["tracker_eligible"] is False


def test_findings_discard_unattributed_boards(tmp_path):
    tracker = copy_uk_tracker(tmp_path)
    dedupe = cw.build_shared_dedupe(CFG, "uk", tracker_override=str(tracker))
    res = synthetic_resolution([{"job_id": "1", "title": "Security Intern",
                                 "location": "London, United Kingdom",
                                 "url": "https://job-boards.greenhouse.io/testco/jobs/1"}])
    res["resolved"][0]["attribution_confidence"] = "low"
    doc = cw.build_findings(res, region="uk", cfg=CFG, dedupe=dedupe, tw=TW,
                            filters=FULL_FILTERS, now=NOW)
    assert doc["findings"] == []


def test_dedupe_is_shared_with_career_ops(tmp_path):
    tracker = copy_uk_tracker(tmp_path)
    wb = openpyxl.load_workbook(tracker)
    ws = wb["Jobs"]
    existing_url = ws[f"{UK_CFG['dedupe']['url_column']}10"].value
    existing_company = ws[f"{UK_CFG['dedupe']['company_column']}10"].value
    existing_title = ws[f"{UK_CFG['dedupe']['title_column']}10"].value
    wb.close()

    jobs = [
        {"job_id": "1", "title": "Whatever", "location": "London, United Kingdom",
         "url": existing_url},                                     # exact URL duplicate
        {"job_id": "2", "title": existing_title, "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/2"},  # company+title duplicate
        {"job_id": "3", "title": "SOC Intern", "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/brand-new-3"},
    ]
    doc, dedupe = make_finding_doc(tmp_path, jobs, tracker=tracker, company=existing_company)
    by_id = {f["vendor_job_id"]: f for f in doc["findings"]}
    assert by_id["1"]["decision"] == "duplicate"
    assert by_id["2"]["decision"] == "duplicate"
    assert by_id["2"]["decision_reason"] == "same company+title already present (posting URL may have moved)"
    assert by_id["3"]["decision"] == "new"
    assert dedupe["tracker_writer"] == "career-ops/tracker_writer.py"
    assert existing_company  # sanity: the fixture row really has a company


def test_cross_month_dedupe_uses_the_region_ledger(tmp_path):
    tracker = copy_uk_tracker(tmp_path)
    ledger_url = None
    ledger = Path(CFG["career_ops_root"]) / "data" / "scan-history.tsv"
    if ledger.exists():
        first = ledger.read_text(encoding="utf-8").splitlines()[1].split("\t")[0]
        ledger_url = first
    if not ledger_url:
        pytest.skip("career-ops scan-history ledger unavailable")
    jobs = [{"job_id": "1", "title": "SOC Intern", "location": "London, United Kingdom",
             "url": ledger_url}]
    doc, _ = make_finding_doc(tmp_path, jobs, tracker=tracker)
    assert doc["findings"][0]["decision"] == "duplicate-cross-month"


def test_manifest_cannot_express_application_state(tmp_path):
    jobs = [{"job_id": "1", "title": "SOC Intern", "location": "London, United Kingdom",
             "url": "https://job-boards.greenhouse.io/testco/jobs/1",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc, _ = make_finding_doc(tmp_path, jobs)
    manifest = cw.build_manifest(doc, region="uk", region_cfg=UK_CFG, cfg=CFG)
    assert manifest["ok"] is True
    assert manifest["counts"]["selected"] == 1
    record = manifest["records"][0]
    forbidden = {"application_status", "priority", "date_applied", "tailored_cv_path",
                 "personal_notes", "notes"}
    assert not (forbidden & set(record))
    assert record["discovery"].startswith("COMPANY WATCH")
    assert "https://job-boards.greenhouse.io/testco/jobs/1" == record["url"]
    assert "NOT human-verified" in record["discovery"]


def test_manifest_skips_other_region_and_ineligible_findings(tmp_path):
    jobs = [{"job_id": "1", "title": "SOC Intern", "location": "Tokyo, Japan",
             "url": "https://job-boards.greenhouse.io/testco/jobs/tokyo",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc, _ = make_finding_doc(tmp_path, jobs)
    manifest = cw.build_manifest(doc, region="uk", region_cfg=UK_CFG, cfg=CFG)
    assert manifest["counts"]["selected"] == 0
    assert cw.build_manifest(doc, region="uk", region_cfg=UK_CFG, cfg=CFG,
                             include_ineligible_as_test=True)["counts"]["selected"] == 0


# --------------------------------------------------------------------------- #
# handoff + tracker write (real writer, copy workbook)
# --------------------------------------------------------------------------- #

def test_handoff_appends_provenance_then_dedupes_on_repeat(tmp_path):
    tracker = copy_uk_tracker(tmp_path)
    archive_dir = tmp_path / "copies"
    archive_dir.mkdir()
    jobs = [
        {"job_id": "1", "title": "SOC Analyst Intern", "location": "London, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/cw-acceptance-1",
         "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"},
        {"job_id": "2", "title": "IT Support Intern", "location": "Manchester, United Kingdom",
         "url": "https://job-boards.greenhouse.io/testco/jobs/cw-acceptance-2",
         "posted_at": "2026-09-22T00:00:00+00:00", "url_source": "vendor"},
    ]
    doc, _ = make_finding_doc(tmp_path, jobs, tracker=tracker)
    manifest = cw.build_manifest(doc, region="uk", region_cfg=UK_CFG, cfg=CFG)
    assert manifest["counts"]["selected"] == 2

    before = TW.sha256_file(tracker)
    dry = write_manifest_cli(tmp_path, manifest["records"], tracker, archive_dir, apply=False)
    assert dry["applied"] is False
    assert dry["counts"]["appended"] == 2
    assert TW.sha256_file(tracker) == before

    applied = write_manifest_cli(tmp_path, manifest["records"], tracker, archive_dir, apply=True)
    assert applied["applied"] is True
    assert applied["counts"]["appended"] == 2
    assert TW.sha256_file(tracker) != before
    assert applied["verification"]["ok"] is True

    # provenance + prior-company signal landed, application state came from defaults
    tr = TW.Tracker(tracker, UK_CFG)
    last = tr.last_data_row()
    assert tr.ws[f"Q{last}"].value.startswith("COMPANY WATCH")
    assert "greenhouse" in tr.ws[f"Q{last}"].value
    assert tr.ws[f"R{last}"].value.startswith("Prior application/CV evidence")
    assert tr.ws[f"J{last}"].value == "To Review"
    assert tr.ws[f"N{last}"].value in (None, "")   # unknown gap is never invented
    tr.wb.close()

    # the same handoff again must append nothing (shared dedupe)
    repeat = write_manifest_cli(tmp_path, manifest["records"], tracker, archive_dir, apply=True)
    assert repeat["counts"]["appended"] == 0
    assert repeat["counts"]["duplicates"] == 2

    # canonical workbook untouched
    assert TW.sha256_file(CANONICAL_UK) != TW.sha256_file(tracker)
    assert TW.sha256_file(CANONICAL_UK) == cw.build_shared_dedupe(
        CFG, "uk", tracker_override=str(CANONICAL_UK))["tracker_sha256"]


def test_handoff_refuses_when_no_provenance_column_exists(tmp_path):
    region_cfg = json.loads(json.dumps(UK_CFG))
    region_cfg["field_map"] = {k: v for k, v in region_cfg["field_map"].items()
                               if k not in ("discovery", "source")}
    doc = {"findings": [{"decision": "new", "attribution_confidence": "high",
                         "region_route": "uk", "owner_filter_eligible": True,
                         "company": "Testco", "title": "T", "url": "https://x.invalid/1",
                         "location": "London", "vendor": "greenhouse", "vendor_job_id": "1",
                         "posted_at": None, "source": "s", "source_timestamp": "t",
                         "attribution_basis": "b", "evidence_class": "employer"}]}
    manifest = cw.build_manifest(doc, region="uk", region_cfg=region_cfg, cfg=CFG)
    assert manifest["ok"] is False
    assert "provenance" in manifest["reason"]


def test_region_trackers_receive_provenance_without_touching_owner_columns():
    for region in ("dubai", "japan", "singapore"):
        _tw, _profiles, cfg = cw.resolve_region_config(CFG, region)
        targets = cw.provenance_targets(cfg)
        assert targets["provenance"] == "source"
        owner_keys = set(targets["owner_column_keys_excluded"])
        assert "notes" in owner_keys
        assert targets["prior_signal"] is None


# Regional end-to-end handoff: a region-routed finding must actually land in that
# region's workbook (provenance in `Source`, owner columns untouched) and must
# dedupe on repeat. Only workbook COPIES are written; the canonical regional
# trackers are opened read-only and their hash is re-checked afterwards.
REGIONAL_CASES = {
    "dubai": {"location": "Dubai, United Arab Emirates",
              "allow": ["Dubai", "United Arab Emirates", "Abu Dhabi"]},
    "japan": {"location": "Tokyo, Japan", "allow": ["Japan", "Tokyo"]},
    "singapore": {"location": "Singapore, Singapore", "allow": ["Singapore"]},
}


def regional_filters(allow: list[str]) -> dict:
    return {
        "available": True,
        "location_always_allow": allow,
        "location_allow": allow,
        "location_block": [],
        "title_positive": ["Intern", "Internship"],
        "title_negative": ["Senior", "Principal", "Lead ", "Manager", "Director", "Head of"],
        "max_posting_age_days": 14,
        "path": "portals.yml",
    }


@pytest.mark.parametrize("region", sorted(REGIONAL_CASES))
def test_regional_handoff_writes_provenance_to_a_tracker_copy(region, tmp_path):
    case = REGIONAL_CASES[region]
    tw, _profiles, region_cfg = cw.resolve_region_config(CFG, region)
    canonical = Path(region_cfg["tracker"])
    if not canonical.exists():
        pytest.skip(f"canonical {region} tracker unavailable")

    copy_path = tmp_path / canonical.name
    shutil.copy2(canonical, copy_path)
    archive_dir = tmp_path / "copies"
    archive_dir.mkdir()

    dedupe = cw.build_shared_dedupe(CFG, region, tracker_override=str(copy_path))
    jobs = [{"job_id": f"cw-{region}-1", "title": "Security Operations Intern",
             "location": case["location"],
             "url": f"https://job-boards.greenhouse.io/testco/jobs/cw-{region}-1",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc = cw.build_findings(synthetic_resolution(jobs), region=region, cfg=CFG, dedupe=dedupe,
                            tw=tw, filters=regional_filters(case["allow"]), now=NOW)
    assert doc["counts"]["findings"] == 1
    assert doc["counts"]["routed_in_scope"] == 1
    assert doc["findings"][0]["tracker_eligible"] is True

    manifest = cw.build_manifest(doc, region=region, region_cfg=region_cfg, cfg=CFG)
    assert manifest["provenance_column"] == "source"
    assert manifest["counts"]["selected"] == 1
    record = manifest["records"][0]
    assert record["source"].startswith(f"{cw.PROVENANCE_PREFIX} — ")
    assert "notes" not in record and "discovery" not in record
    assert record["last_checked"] == NOW.date().isoformat()

    tr = tw.Tracker(copy_path, region_cfg)
    before_owned = tr.owned_snapshot()
    last_before = tr.last_data_row()
    tr.wb.close()
    canonical_before = tw.sha256_file(canonical)

    dry = write_manifest_cli(tmp_path, manifest["records"], copy_path, archive_dir,
                             apply=False, region=region)
    assert dry["applied"] is False
    assert dry["counts"]["appended"] == 1
    assert tw.sha256_file(copy_path) == dedupe["tracker_sha256"]

    applied = write_manifest_cli(tmp_path, manifest["records"], copy_path, archive_dir,
                                 apply=True, region=region)
    assert applied["applied"] is True
    assert applied["counts"]["appended"] == 1
    assert applied["verification"]["ok"] is True

    tr = tw.Tracker(copy_path, region_cfg)
    last = tr.last_data_row()
    assert last == last_before + 1
    assert tr.ws[f"X{last}"].value.startswith(cw.PROVENANCE_PREFIX)
    assert tr.ws[f"Z{last}"].value in (None, "")          # owner notes never automated
    after_owned = tr.owned_snapshot()
    for row, cols in before_owned.items():
        for col, value in cols.items():
            assert after_owned[row][col] == value, f"{region} owner column {col}{row} changed"
    # the appended row carries only the writer's own documented defaults
    assert after_owned[last]["R"] == region_cfg["defaults"]["application_status"]
    tr.wb.close()

    # shared dedupe: the same region-routed finding must not append twice
    repeat = write_manifest_cli(tmp_path, manifest["records"], copy_path, archive_dir,
                                apply=True, region=region)
    assert repeat["counts"]["appended"] == 0
    assert repeat["counts"]["duplicates"] == 1

    # the canonical regional workbook was never written
    assert tw.sha256_file(canonical) == canonical_before
    assert tw.sha256_file(canonical) != tw.sha256_file(copy_path)


def test_regional_handoff_requires_an_explicit_region(tmp_path):
    """A finding routed to another region is never handed off implicitly."""
    tw, _profiles, uk_cfg = cw.resolve_region_config(CFG, "uk")
    tracker = copy_uk_tracker(tmp_path)
    dedupe = cw.build_shared_dedupe(CFG, "uk", tracker_override=str(tracker))
    jobs = [{"job_id": "1", "title": "SOC Analyst Intern", "location": "Tokyo, Japan",
             "url": "https://job-boards.greenhouse.io/testco/jobs/tokyo-1",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc = cw.build_findings(synthetic_resolution(jobs), region="uk", cfg=CFG, dedupe=dedupe,
                            tw=tw, filters=FULL_FILTERS, now=NOW)
    assert doc["findings"][0]["region_route"] == "japan"
    assert doc["counts"]["routed_other_region"] == 1
    manifest = cw.build_manifest(doc, region="uk", region_cfg=uk_cfg, cfg=CFG)
    assert manifest["counts"]["selected"] == 0
    assert manifest["records"] == []


# --------------------------------------------------------------------------- #
# monthly operational workbook
# --------------------------------------------------------------------------- #

def test_monthly_workbook_creation_and_conflict_safety(tmp_path):
    jobs = [{"job_id": "1", "title": "SOC Intern", "location": "London, United Kingdom",
             "url": "https://job-boards.greenhouse.io/testco/jobs/1",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc, _ = make_finding_doc(tmp_path, jobs)
    res = synthetic_resolution(jobs)
    result = cw.write_workbook(
        CFG, month="2026-09", directory=tmp_path / "wb", registry={"source": {"path": "x"},
                                                                   "_prior_evidence": {"Testco": True}},
        resolutions=res, findings_doc=doc,
        handoff_result={"region": "uk", "mode": "dry-run", "tracker": "canonical",
                        "manifest": "m", "exit_code": 0,
                        "result": {"counts": {"appended": 1}}},
        run_meta={"run_id": "r1", "started_at": "t", "region": "uk", "companies_probed": 1,
                  "boards_resolved": 1, "findings": 1, "duplicates": 0, "eligible": 0,
                  "http_requests": 1, "budget_exhausted": False, "duration_s": 1.0, "mode": "live"},
    )
    assert result["ok"] is True
    assert set(["Watch List", "Findings", "Tracker Handoff", "Run Log", "README"]) <= set(result["sheets"])
    assert result["forbidden_application_state_columns"] == []
    assert result["conflict_check"]["safe"] is True

    wb = openpyxl.load_workbook(result["path"])
    assert wb["Findings"].max_row == 2
    assert wb["Watch List"].max_row == 2
    assert wb["Run Log"].max_row == 2
    wb.close()


def test_monthly_workbook_is_invisible_to_career_ops_dedupe(tmp_path):
    path = cw.workbook_path(CFG, "2026-09", tmp_path)
    dedupe = cw.build_shared_dedupe(CFG, "uk", tracker_override=str(CANONICAL_UK),
                                    extra_archive_dirs=[str(tmp_path)])
    assert path.name == "Company_Watch_2026-09.xlsx"
    assert dedupe["cross_month_keys"] == cw.build_shared_dedupe(
        CFG, "uk", tracker_override=str(CANONICAL_UK))["cross_month_keys"]
    for region in ("japan", "singapore", "dubai"):
        glob = (f"{region.capitalize()}_Cybersecurity_Job_Tracker*.xlsx")
        assert not Path(path.name).match(glob)


def call_workbook(tmp_path, doc, res, run_id, handoff=None):
    return cw.write_workbook(
        CFG, month="2026-09", directory=tmp_path / "wb",
        registry={"source": {"path": "x"}, "_prior_evidence": {"Testco": True}},
        resolutions=res, findings_doc=doc, handoff_result=handoff,
        run_meta={"run_id": run_id, "started_at": "t", "region": "uk", "companies_probed": 1,
                  "boards_resolved": 1, "findings": len(doc.get("findings", [])),
                  "duplicates": 0, "eligible": 0, "http_requests": 1,
                  "budget_exhausted": False, "duration_s": 1.0, "mode": "live"})


def test_monthly_workbook_is_a_ledger_not_a_run_log(tmp_path):
    jobs = [{"job_id": "1", "title": "SOC Intern", "location": "London, United Kingdom",
             "url": "https://job-boards.greenhouse.io/testco/jobs/1",
             "posted_at": "2026-09-21T00:00:00+00:00", "url_source": "vendor"}]
    doc, _ = make_finding_doc(tmp_path, jobs)
    res = synthetic_resolution(jobs)
    first = call_workbook(tmp_path, doc, res, "run-1")
    second = call_workbook(tmp_path, doc, res, "run-2")
    assert first["watch_list_rows"] == 1
    assert second["watch_list_rows"] == 1          # refreshed, not duplicated
    assert second["watch_list_refreshed"] == 1
    assert first["findings_rows"] == 1
    assert second["findings_rows"] == 1            # same posting not re-recorded
    assert second["findings_added"] == 0
    wb = openpyxl.load_workbook(second["path"])
    assert wb["Run Log"].max_row == 3              # one row per distinct run id
    wb.close()


def test_workbook_refuses_to_overwrite_a_regional_tracker(tmp_path):
    cfg = json.loads(json.dumps(CFG))
    cfg["workbook"]["filename_template"] = "uk-cyber-job-tracker.xlsx"
    dest = cw.workbook_path(cfg, "2026-09", tmp_path)
    shutil.copy2(CANONICAL_UK, dest)
    before = TW.sha256_file(dest)
    result = cw.write_workbook(cfg, month="2026-09", directory=tmp_path)
    assert result["ok"] is False
    assert result["conflict"]["archive_glob_conflicts"]
    assert TW.sha256_file(dest) == before              # the copy is untouched
    assert not dest.with_suffix(".write-tmp.xlsx").exists()


# --------------------------------------------------------------------------- #
# registry-driven sweep ordering
# --------------------------------------------------------------------------- #

def test_ordered_companies_is_deterministic_and_sliceable():
    reg = {"companies": [
        {"name": "Zeta", "evidence_class": "recruiter"},
        {"name": "Alpha", "evidence_class": "employer"},
        {"name": "Beta", "evidence_class": "employer"},
        {"name": "Gamma", "evidence_class": "unverified_candidate"},
    ]}
    order = [c["name"] for c in cw.ordered_companies(reg)]
    assert order == ["Alpha", "Beta", "Zeta", "Gamma"]
    assert [c["name"] for c in cw.ordered_companies(reg, 1, 2)] == ["Beta", "Zeta"]
    assert [c["name"] for c in cw.ordered_companies(reg, 3)] == ["Gamma"]
    assert [c["name"] for c in cw.ordered_companies(reg, 99)] == []


def test_age_parsing_is_conservative():
    now = dt.datetime(2026, 9, 23, tzinfo=dt.timezone.utc)
    assert cw.parse_age_days("2026-09-22T00:00:00Z", now) == 1.0
    assert cw.parse_age_days("Posted 30+ days ago", now) is None
    assert cw.parse_age_days(None, now) is None
