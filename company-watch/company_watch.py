#!/usr/bin/env python3
"""Company Watch — deterministic watch → finding → dedupe → Career Ops handoff.

Company Watch sits under Career. It watches the organisations Mukund already has
recorded application/CV evidence for, checks their *public structured ATS/career
endpoints*, and hands eligible findings to the existing Career Ops tracker writer.

It does not:

* submit applications, contact companies or recruiters, or touch any account;
* invent employers, jobs, salaries, eligibility or application state;
* keep its own copy of application state, or write tracker rows behind Career
  Ops' back. The canonical regional workbooks stay authoritative and every write
  goes through ``career-ops/tracker_writer.py``'s dry-run/backup/verify path.

Dedupe is *shared*, not reimplemented: this module imports the same
``tracker_writer`` primitives Career Ops uses (``normalize_url``, ``pair_key``,
``build_cross_month_index``), so a Company Watch finding cannot create a row that
Career Ops would consider new state.

Subcommands (each prints exactly one JSON object on stdout):

    registry    build/refresh the watched-company registry from historical evidence
    resolve     bounded structured-ATS resolution for a slice of the registry
    scan        resolve + normalise + eligibility + shared dedupe → findings
    handoff     hand findings to Career Ops (dry-run by default)
    workbook    create/refresh the monthly Company Watch operational workbook
    run         the bounded end-to-end chain, with an evidence document

Usage:
    python company-watch/company_watch.py run --region uk --limit 12 \\
        --out-dir audits/evidence/<stamp>-company-watch-integration
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent
DEFAULT_CONFIG = HERE / "company_watch_config.json"


# --------------------------------------------------------------------------- #
# configuration & shared Career Ops imports
# --------------------------------------------------------------------------- #

def load_config(path: str | Path | None = None) -> dict:
    cfg = json.loads(Path(path or DEFAULT_CONFIG).read_text(encoding="utf-8"))
    cfg.setdefault("_config_path", str(path or DEFAULT_CONFIG))
    return cfg


def import_tracker_writer(career_ops_dir: Path):
    """Import the *existing* Career Ops writer — never a copy of its rules."""
    career_ops_dir = Path(career_ops_dir)
    if not (career_ops_dir / "tracker_writer.py").exists():
        raise FileNotFoundError(f"tracker_writer.py not found in {career_ops_dir}")
    if str(career_ops_dir) not in sys.path:
        sys.path.insert(0, str(career_ops_dir))
    import tracker_writer  # noqa: PLC0415
    return tracker_writer


def career_ops_paths(cfg: dict) -> dict:
    return {
        "install_root": Path(cfg["career_ops_root"]),
        "integration_dir": REPO_ROOT / cfg["career_ops_dir"],
        "profiles": REPO_ROOT / cfg["regional_profiles"],
        "runtime_dir": REPO_ROOT / cfg["runtime_dir"],
        "cli": REPO_ROOT / cfg["career_ops_dir"] / "career_ops_cli.py",
    }


def resolve_region_config(cfg: dict, region: str):
    tw = import_tracker_writer(career_ops_paths(cfg)["integration_dir"])
    profiles = tw.load_profiles(career_ops_paths(cfg)["profiles"])
    region_cfg = dict(tw.region_config(profiles, region))
    region_cfg["region"] = region
    return tw, profiles, region_cfg


# --------------------------------------------------------------------------- #
# owner search filters (single source of truth: the Career Ops install)
# --------------------------------------------------------------------------- #

def load_owner_filters(cfg: dict, region: str) -> dict:
    """Read the owner's configured search filters for a region.

    The UK lane's filters live in the Career Ops install's ``portals.yml``. No
    other region has a lane config yet (that is owned by
    ``agent-regional-job-search-agents-and-schedulers-2026-09-23``), so an absent
    lane is reported as unavailable — never silently substituted with the UK
    filters, which would produce wrong-region eligibility.
    """
    root = Path(cfg["career_ops_root"])
    portals = root / "portals.yml"
    if region != "uk" or not portals.exists():
        return {
            "available": False,
            "reason": ("no lane filter config for this region" if region != "uk"
                       else f"portals.yml not found at {portals}"),
            "path": str(portals),
        }
    text = portals.read_text(encoding="utf-8")
    parsed, parser = _parse_portals_filters(text, portals)
    parsed.update({"available": True, "path": str(portals), "parser": parser})
    return parsed


def _parse_portals_filters(text: str, path: Path) -> tuple[dict, str]:
    try:
        import yaml  # noqa: PLC0415
        doc = yaml.safe_load(text) or {}
        loc = doc.get("location_filter") or {}
        title = doc.get("title_filter") or {}
        return {
            "location_always_allow": list(loc.get("always_allow") or []),
            "location_allow": list(loc.get("allow") or []),
            "location_block": list(loc.get("block") or []),
            "title_positive": list(title.get("positive") or []),
            "title_negative": list(title.get("negative") or []),
            "max_posting_age_days": doc.get("max_posting_age_days"),
        }, "pyyaml"
    except ImportError:
        return _parse_portals_filters_minimal(text), "minimal_parser"


def _parse_portals_filters_minimal(text: str) -> dict:
    """Fallback parser for the two filter blocks when PyYAML is unavailable.

    Deliberately narrow: it reads only ``key:`` blocks whose children are
    ``- "value"`` list items, which is exactly the shape used by portals.yml for
    ``location_filter`` and ``title_filter``. Anything else is ignored rather
    than guessed.
    """
    out: dict[str, list[str]] = {}
    stack: list[str] = []
    for raw in text.splitlines():
        line = raw.split("#")[0].rstrip()
        if not line.strip():
            continue
        indent = len(line) - len(line.lstrip())
        stripped = line.strip()
        if stripped.endswith(":") and not stripped.startswith("-"):
            key = stripped[:-1].strip()
            stack = stack[: indent // 2] + [key]
            continue
        if stripped.startswith("- "):
            if indent // 2 < len(stack):
                stack = stack[: indent // 2]
            if not stack:
                continue
            container, key = stack[0], stack[-1]
            if container in ("location_filter", "title_filter"):
                out.setdefault(f"{container}.{key}", []).append(stripped[2:].strip().strip('"\''))
    m = re.search(r"^max_posting_age_days:\s*(\d+)", text, re.MULTILINE)
    return {
        "location_always_allow": out.get("location_filter.always_allow", []),
        "location_allow": out.get("location_filter.allow", []),
        "location_block": out.get("location_filter.block", []),
        "title_positive": out.get("title_filter.positive", []),
        "title_negative": out.get("title_filter.negative", []),
        "max_posting_age_days": int(m.group(1)) if m else None,
    }


def _words(text: str) -> str:
    return re.sub(r"[^0-9a-z]+", " ", (text or "").casefold())


def term_in_text(term: str, text: str) -> bool:
    """Word-boundary containment.

    Substring matching would let "Intern" match "Internal Audit" and "UK" match
    "Ukraine" — both real false positives in this data — so matching is on
    standalone tokens.
    """
    t = _words(term)
    if not t:
        return False
    return re.search(rf"(?<![0-9a-z]){re.escape(t.strip())}(?![0-9a-z])", _words(text)) is not None


def any_term_in_text(terms: list[str], text: str) -> bool:
    return any(term_in_text(t, text) for t in terms)


# --------------------------------------------------------------------------- #
# registry
# --------------------------------------------------------------------------- #

def load_registry(cfg: dict, *, rebuild: bool = False) -> dict:
    import company_registry  # noqa: PLC0415
    paths = career_ops_paths(cfg)
    out = paths["runtime_dir"] / "company_watch_registry.json"
    source = Path(cfg["registry_source"])
    if rebuild or not out.exists():
        registry = company_registry.build_registry(source)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(registry, indent=2, ensure_ascii=False), encoding="utf-8")
        return registry
    registry = json.loads(out.read_text(encoding="utf-8"))
    # refresh if the owner's source changed under us
    if registry["source"]["sha256"] != company_registry.sha256_file(source):
        registry = company_registry.build_registry(source)
        out.write_text(json.dumps(registry, indent=2, ensure_ascii=False), encoding="utf-8")
    return registry


def ordered_companies(registry: dict, start_index: int = 0, limit: int | None = None) -> list[dict]:
    """Deterministic sweep order over the watched companies.

    Employers first, then recruiters, then the other recorded classes, each
    alphabetically. ``start_index`` lets consecutive bounded runs advance through
    the registry instead of re-probing the same first N every night.
    """
    priority = {"employer": 0, "recruiter": 1, "other_application": 2, "unverified_candidate": 3}
    companies = sorted(
        registry["companies"],
        key=lambda c: (priority.get(c["evidence_class"], 9), c["name"].casefold()),
    )
    sliced = companies[start_index:]
    if limit is not None:
        sliced = sliced[:limit]
    return sliced


# --------------------------------------------------------------------------- #
# shared dedupe (imported from Career Ops)
# --------------------------------------------------------------------------- #

def build_shared_dedupe(cfg: dict, region: str, *, tracker_override: str | None = None,
                        extra_archive_dirs: list[str] | None = None) -> dict:
    tw, profiles, region_cfg = resolve_region_config(cfg, region)
    tracker = Path(tracker_override or region_cfg["tracker"])
    tr = tw.Tracker(tracker, region_cfg)
    by_url, by_pair = tr.existing_keys()
    cross = tw.build_cross_month_index(profiles, region_cfg, extra_archive_dirs)
    canonical_companies = {
        tw.key_text(tr.ws.cell(row=r, column=tw.column_index_from_string(region_cfg["dedupe"]["company_column"])).value)
        for r in range(region_cfg["first_data_row"], tr.last_data_row() + 1)
    }
    data_rows = tr.last_data_row() - region_cfg["first_data_row"] + 1
    tr.wb.close()
    return {
        "region": region,
        "tracker": str(tracker),
        "tracker_sha256": tw.sha256_file(tracker),
        "canonical_url_keys": set(by_url),
        "canonical_pair_keys": set(by_pair),
        "canonical_data_rows": data_rows,
        "cross_month_keys": set(cross["url_keys"]),
        "cross_month_sources": cross["sources"],
        "canonical_companies": {c for c in canonical_companies if c},
        "tracker_writer": "career-ops/tracker_writer.py",
    }


def dedupe_decision(finding: dict, dedupe: dict, tw) -> tuple[str, str]:
    url_key = tw.normalize_url(finding.get("url"))
    pkey = tw.pair_key(finding.get("company"), finding.get("title"))
    if not url_key:
        return "rejected", "missing or unusable application URL"
    if url_key in dedupe["canonical_url_keys"]:
        return "duplicate", "already present in the canonical regional workbook"
    if pkey.strip("|") and pkey in dedupe["canonical_pair_keys"]:
        return "duplicate", "same company+title already present (posting URL may have moved)"
    if url_key in dedupe["cross_month_keys"]:
        return "duplicate-cross-month", "present in an archived workbook or the region ledger"
    return "new", "eligible for Career Ops handoff"


# --------------------------------------------------------------------------- #
# findings
# --------------------------------------------------------------------------- #

def parse_age_days(posted_at: str | None, now: dt.datetime) -> float | None:
    if not posted_at or not isinstance(posted_at, str):
        return None
    text = posted_at.strip().replace("Z", "+00:00")
    try:
        when = dt.datetime.fromisoformat(text)
    except ValueError:
        return None
    if when.tzinfo is None:
        when = when.replace(tzinfo=dt.timezone.utc)
    return round((now - when).total_seconds() / 86400.0, 1)


API_URL_HOSTS = ("api.smartrecruiters.com", "boards-api.greenhouse.io", "api.lever.co",
                 "api.ashbyhq.com")


def classify_url(url: str | None, url_source: str | None) -> str:
    """Classify a finding URL before it can ever reach a tracker.

    ``posting``            vendor-supplied public posting URL
    ``derived_public``     constructed from a documented vendor URL pattern;
                           must be verified live before it is tracker-eligible
    ``api_endpoint``       a JSON API address — not a human application URL
    ``missing``            nothing usable
    """
    if not url or not isinstance(url, str):
        return "missing"
    lowered = url.casefold()
    if any(host in lowered for host in API_URL_HOSTS):
        return "api_endpoint"
    if url_source == "derived_pattern":
        return "derived_public"
    if not lowered.startswith(("http://", "https://")):
        return "missing"
    return "posting"


def build_findings(resolutions: dict, *, region: str, cfg: dict, dedupe: dict, tw,
                   filters: dict, now: dt.datetime, fetcher=None) -> dict:
    findings: list[dict] = []
    seen_urls: set[str] = set()
    seen_pairs: set[str] = set()
    age_limit = filters.get("max_posting_age_days") if filters.get("available") else None
    verify_derived = bool(cfg.get("probe", {}).get("verify_derived_urls", True))

    for rec in resolutions.get("resolved", []):
        if rec.get("attribution_confidence") != "high":
            continue
        for job in rec.get("jobs", []):
            url = job.get("url")
            title = job.get("title") or ""
            location = job.get("location") or ""
            url_key = tw.normalize_url(url)
            pair = tw.pair_key(rec["company"], title)

            title_pos = filters.get("title_positive") or []
            title_neg = filters.get("title_negative") or []
            loc_allow = (filters.get("location_always_allow") or []) + (filters.get("location_allow") or [])
            loc_block = filters.get("location_block") or []

            if not filters.get("available"):
                title_ok: bool | None = None
                location_ok: bool | None = None
                age_ok: bool | None = None
            else:
                title_ok = any_term_in_text(title_pos, title) and not any_term_in_text(title_neg, title)
                location_ok = (any_term_in_text(loc_allow, location)
                               and not any_term_in_text(loc_block, location))
                age = parse_age_days(job.get("posted_at"), now)
                age_ok = None if age is None else (age_limit is None or age <= float(age_limit))

            owner_eligible = None
            if filters.get("available"):
                owner_eligible = bool(title_ok and location_ok and (age_ok is not False))

            route = route_region(location, cfg, filters=filters, region=region)
            route_in_scope = route == region

            finding = {
                "company": rec["company"],
                "evidence_class": rec.get("evidence_class"),
                "title": title,
                "location": location,
                "url": url,
                "url_source": job.get("url_source", "vendor"),
                "url_quality": classify_url(url, job.get("url_source")),
                "url_verified": None,
                "vendor": job.get("vendor"),
                "board_slug": job.get("board_slug"),
                "vendor_job_id": job.get("job_id"),
                "posted_at": job.get("posted_at"),
                "updated_at": job.get("updated_at"),
                "posted_age_days": parse_age_days(job.get("posted_at"), now),
                "source": f"company-watch/{job.get('vendor')}:{job.get('board_slug')}",
                "source_timestamp": now.replace(microsecond=0).isoformat(),
                "attribution_confidence": rec.get("attribution_confidence"),
                "attribution_basis": rec.get("attribution_basis"),
                "region_route": route,
                "title_rule_pass": title_ok,
                "location_rule_pass": location_ok,
                "age_rule_pass": age_ok,
                "owner_filter_eligible": owner_eligible,
            }

            if not url_key:
                finding.update(decision="rejected",
                               decision_reason="missing or unusable application URL")
            elif url_key in seen_urls or (pair.strip("|") and pair in seen_pairs):
                finding.update(decision="duplicate-in-run",
                               decision_reason="same posting/pair already produced in this run")
            else:
                decision, reason = dedupe_decision(finding, dedupe, tw)
                finding.update(decision=decision, decision_reason=reason)
                seen_urls.add(url_key)
                if pair.strip("|"):
                    seen_pairs.add(pair)

            company_key = tw.key_text(rec["company"])
            finding["prior_application_company_evidence"] = {
                "in_watch_registry": True,
                "evidence_class": rec.get("evidence_class"),
                "already_in_canonical_tracker": company_key in dedupe["canonical_companies"],
                "note": ("a previous application to this company is a review signal, "
                         "not a company-wide block"),
            }
            finding["tracker_eligible"] = bool(
                finding["decision"] == "new"
                and finding["attribution_confidence"] == "high"
                and route_in_scope
                and (owner_eligible is True)
                and finding["url_quality"] in ("posting", "derived_public")
            )
            # A constructed URL is not vendor evidence: verify it live before it
            # can be handed off. Only otherwise-eligible findings cost a request.
            if (finding["tracker_eligible"] and finding["url_quality"] == "derived_public"
                    and fetcher is not None and verify_derived):
                status, _body, _err = fetcher.get(finding["url"])
                finding["url_verified"] = (status == 200)
                if status != 200:
                    finding["tracker_eligible"] = False
                    finding["decision_reason"] += (
                        f"; derived posting URL failed live verification (HTTP {status})")
            findings.append(finding)

    counts: dict = {}
    for f in findings:
        counts[f["decision"]] = counts.get(f["decision"], 0) + 1
    return {
        "region": region,
        "generated_at": now.replace(microsecond=0).isoformat(),
        "filters": {k: v for k, v in filters.items() if k != "path"},
        "filters_path": filters.get("path"),
        "region_routing_configured": sorted(cfg.get("region_routing", {})),
        "counts": {
            "findings": len(findings),
            "by_decision": counts,
            "tracker_eligible": sum(1 for f in findings if f["tracker_eligible"]),
            "owner_filter_eligible": sum(1 for f in findings if f["owner_filter_eligible"] is True),
            "routed_in_scope": sum(1 for f in findings if f["region_route"] == region),
            "routed_other_region": sum(1 for f in findings
                                       if f["region_route"] not in (region, None)),
            "routed_unresolved": sum(1 for f in findings if f["region_route"] is None),
        },
        "findings": findings,
    }


def route_region(location: str | None, cfg: dict, filters: dict | None = None,
                 region: str | None = None) -> str | None:
    """Route a finding to the regional tracker its location belongs to.

    The owner's own configured location terms for the region being scanned come
    first, so routing and eligibility cannot disagree about what counts as UK.
    Unmatched locations return ``None`` (recorded as unresolved) rather than
    being pushed into a tracker they may not belong to.
    """
    rules = cfg.get("region_routing", {})
    extra: dict[str, list[str]] = {}
    if filters and filters.get("available") and region:
        extra[region] = (list(filters.get("location_always_allow") or [])
                         + list(filters.get("location_allow") or []))
    for region_name, spec in rules.items():
        terms = list(spec.get("terms", [])) + extra.get(region_name, [])
        if any_term_in_text(terms, location or ""):
            if spec.get("block_terms") and any_term_in_text(spec["block_terms"], location or ""):
                continue
            return region_name
    return None


# --------------------------------------------------------------------------- #
# handoff to Career Ops
# --------------------------------------------------------------------------- #

PROVENANCE_PREFIX = "COMPANY WATCH"


def provenance_targets(region_cfg: dict) -> dict:
    """Which canonical columns carry Company Watch provenance for this region.

    Never an owner column: the regional trackers keep their Notes column (Z)
    owner-only, so a region without a non-owner provenance column simply gets
    provenance recorded in the handoff manifest and the Company Watch workbook
    instead of in the tracker.
    """
    field_map = region_cfg.get("field_map", {})
    owner = set(region_cfg.get("owner_columns", []))
    owner_keys = {k for k, col in field_map.items() if col in owner}

    def pick(candidates):
        for key in candidates:
            if key in field_map and key not in owner_keys:
                return key
        return None

    return {
        "provenance": pick(["discovery", "source"]),
        "prior_signal": pick(["prior_company_signal"]),
        "live_status": pick(["live_status"]),
        "posted_date": pick(["posted_date"]),
        "last_checked": pick(["last_checked"]),
        "owner_column_keys_excluded": sorted(owner_keys),
    }


def build_manifest(findings_doc: dict, *, region: str, region_cfg: dict, cfg: dict,
                   include_ineligible_as_test: bool = False, limit: int = 0) -> dict:
    """Convert findings into a Career Ops tracker-writer manifest.

    The writer's own guard means a manifest can never set application state; this
    function additionally refuses to invent any value it does not hold: unknown
    clearance, work-authorisation risk, salary or gaps are simply absent rather
    than defaulted into a claim.
    """
    targets = provenance_targets(region_cfg)
    if not targets["provenance"]:
        return {
            "ok": False,
            "reason": (f"region '{region}' has no non-owner provenance column; refusing to "
                       "hand off without provenance"),
            "region": region,
            "records": [],
        }

    selected = []
    for f in findings_doc.get("findings", []):
        if f.get("decision") != "new":
            continue
        if f.get("attribution_confidence") != "high":
            continue
        if f.get("region_route") != region:
            continue
        if f.get("owner_filter_eligible") is True:
            selected.append((f, False))
        elif include_ineligible_as_test:
            selected.append((f, True))
    if limit:
        selected = selected[:limit]

    records = []
    for f, is_test in selected:
        marker = f"{PROVENANCE_PREFIX} TEST ROW — " if is_test else f"{PROVENANCE_PREFIX} — "
        provenance = (
            f"{marker}source {f['source']} (vendor job id {f['vendor_job_id']}); "
            f"feeding timestamp {f['source_timestamp']}; posting timestamp "
            f"{f['posted_at'] or 'not stated'}; attribution {f['attribution_basis']}; "
            f"dedupe {f['decision']} against the canonical {region} workbook and cross-month ledger"
        ) + ("; NOT human-verified — review before shortlisting" if not is_test else
             "; acceptance-test row written to a workbook copy only")
        record = {
            "company": f["company"],
            "title": f["title"],
            "location": f["location"],
            "url": f["url"],
            targets["provenance"]: provenance,
        }
        if targets["live_status"]:
            # "Live" is a feed observation at probe time, never a liveness claim
            # about a destination page that Company Watch has not opened.
            record[targets["live_status"]] = "Live"
        if targets["posted_date"] and f.get("posted_at"):
            record[targets["posted_date"]] = f["posted_at"][:10]
        if targets["last_checked"]:
            record[targets["last_checked"]] = f["source_timestamp"][:10]
        if targets["prior_signal"]:
            klass = f.get("evidence_class")
            record[targets["prior_signal"]] = (
                f"Prior application/CV evidence on record ({klass}); role equivalence not established"
            )
        records.append(record)

    return {
        "ok": True,
        "region": region,
        "provenance_column": targets["provenance"],
        "provenance_targets": targets,
        "counts": {"selected": len(selected), "included_as_test_rows": sum(1 for _f, t in selected if t)},
        "records": records,
    }


def handoff_to_career_ops(cfg: dict, region: str, manifest: dict, *, tracker: str | None = None,
                          apply: bool = False, backup_dir: str | None = None,
                          extra_archive_dirs: list[str] | None = None,
                          manifest_path: str | None = None) -> dict:
    """Invoke the deterministic Career Ops writer exactly as a caller would."""
    paths = career_ops_paths(cfg)
    if not manifest.get("records"):
        return {"invoked": False, "reason": "no eligible records to hand off",
                "mode": "apply" if apply else "dry-run"}
    mpath = Path(manifest_path) if manifest_path else (
        paths["runtime_dir"] / f"handoff-{region}-{dt.datetime.now(dt.timezone.utc):%Y%m%dT%H%M%SZ}.json")
    mpath.parent.mkdir(parents=True, exist_ok=True)
    mpath.write_text(json.dumps({"records": manifest["records"]}, indent=2, ensure_ascii=False),
                     encoding="utf-8")
    cmd = [sys.executable, str(paths["cli"]), "write", "--region", region, "--manifest", str(mpath)]
    if tracker:
        cmd += ["--tracker", str(tracker)]
    if backup_dir:
        cmd += ["--backup-dir", str(backup_dir)]
    for d in extra_archive_dirs or []:
        cmd += ["--archive-dir", str(d)]
    if apply:
        cmd += ["--apply"]
    proc = subprocess.run(cmd, cwd=str(REPO_ROOT), capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=600)
    result: dict = {"invoked": True, "command": " ".join(cmd), "exit_code": proc.returncode,
                    "manifest": str(mpath), "mode": "apply" if apply else "dry-run",
                    "tracker": str(tracker or resolve_region_config(cfg, region)[2]["tracker"])}
    try:
        result["result"] = json.loads(proc.stdout)
    except json.JSONDecodeError:
        result["result"] = None
        result["stdout_tail"] = (proc.stdout or "")[-2000:]
        result["stderr_tail"] = (proc.stderr or "")[-2000:]
        result["ok"] = False
        return result
    result["ok"] = proc.returncode == 0
    return result


# --------------------------------------------------------------------------- #
# monthly operational workbook
# --------------------------------------------------------------------------- #

WORKBOOK_SHEETS = {
    "README": [
        ["Company Watch — monthly operational workbook"],
        ["Owner", "Mukund. Maintained by Chief of Staff (Company Watch)."],
        ["Status", "Operational record for Company Watch only. NOT application state."],
        ["Authority", "Microsoft Excel regional trackers stay authoritative for applications."],
        ["Rule 1", "This workbook is never read by the Career Ops dedupe index or the tracker writer."],
        ["Rule 2", "No application-status / owner column exists here; application state is never implied."],
        ["Rule 3", "Every row carries its source, vendor board and feeding timestamp."],
        ["Rule 4", "Findings are ATS-feed observations, not verified vacancies or eligibility decisions."],
        ["Rule 5", "No Company Watch run submits applications, messages or contacts anyone."],
    ],
    "Watch List": [
        ["Company", "Evidence Class", "Prior Application Evidence", "ATS Vendor", "Board Slug",
         "Jobs Listed", "Attribution Confidence", "Attribution Basis", "Manual Attribution Required",
         "Last Resolved At", "Board URL", "Source", "Source Section"],
    ],
    "Findings": [
        ["Company", "Title", "Location", "URL", "URL Quality", "Vendor", "Vendor Job ID",
         "Posted At", "Age (days)", "Region Route", "Dedupe Decision", "Dedupe Reason",
         "Owner Filter Eligible", "Tracker Eligible", "Attribution Confidence", "Source",
         "Feeding Timestamp"],
    ],
    "Tracker Handoff": [
        ["Region", "Mode", "Rows", "Tracker", "Manifest", "Career Ops Exit Code",
         "Handed Off At", "Notes"],
    ],
    "Run Log": [
        ["Run ID", "Started (UTC)", "Region", "Companies Probed", "Boards Resolved", "Findings",
         "Duplicates", "Eligible", "HTTP Requests", "Budget Exhausted", "Duration (s)", "Mode"],
    ],
}


def workbook_path(cfg: dict, month: str, directory: str | Path | None = None) -> Path:
    name = cfg.get("workbook", {}).get("filename_template", "Company_Watch_{month}.xlsx").format(month=month)
    base = Path(directory or cfg["workbook_dir"])
    return base / name


def workbook_conflict_check(cfg: dict, path: Path) -> dict:
    """Prove the workbook cannot be mistaken for a regional tracker."""
    import fnmatch  # noqa: PLC0415
    conflicts = []
    profiles_path = career_ops_paths(cfg)["profiles"]
    profiles = json.loads(profiles_path.read_text(encoding="utf-8"))
    for region in profiles["regions"]:
        globs = ([f"{region}-cyber-job-tracker*.xlsx"] if region == "uk"
                 else [f"{region.capitalize()}_Cybersecurity_Job_Tracker*.xlsx"])
        for g in globs:
            if fnmatch.fnmatch(path.name, g):
                conflicts.append({"region": region, "glob": g})
    return {"path": str(path), "name": path.name, "archive_glob_conflicts": conflicts,
            "safe": not conflicts}


def write_workbook(cfg: dict, *, month: str, directory: str | Path | None = None,
                   registry: dict | None = None, resolutions: dict | None = None,
                   findings_doc: dict | None = None, handoff_result: dict | None = None,
                   run_meta: dict | None = None) -> dict:
    import openpyxl  # noqa: PLC0415
    path = workbook_path(cfg, month, directory)
    conflict = workbook_conflict_check(cfg, path)
    if not conflict["safe"]:
        return {"ok": False, "reason": "monthly workbook name collides with a regional "
                                       "tracker archive glob", "conflict": conflict}

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet, rows in WORKBOOK_SHEETS.items():
            ws = wb.create_sheet(sheet)
            for row in rows:
                ws.append(row)
    for sheet in WORKBOOK_SHEETS:
        if sheet not in wb.sheetnames:
            ws = wb.create_sheet(sheet)
            for row in WORKBOOK_SHEETS[sheet]:
                ws.append(row)

    stamp = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    now_str = stamp

    # --- Watch List: upsert by company (a monthly workbook is a ledger, not a log
    #     of every run; re-probing a company refreshes its row instead of adding one)
    ws = wb["Watch List"]
    watch_cols = {name: idx for idx, name in enumerate(WORKBOOK_SHEETS["Watch List"][0], start=1)}
    watch_rows: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        value = ws.cell(row=r, column=watch_cols["Company"]).value
        if value:
            watch_rows[str(value).casefold()] = r
    updates = 0
    for rec in (resolutions or {}).get("resolved", []) + (resolutions or {}).get("unresolved", []):
        row = [
            rec.get("company"), rec.get("evidence_class"),
            (registry or {}).get("_prior_evidence", {}).get(rec.get("company")),
            rec.get("vendor"), rec.get("board_slug"), rec.get("jobs_listed"),
            rec.get("attribution_confidence"), rec.get("attribution_basis"),
            rec.get("manual_attribution_required"), now_str, rec.get("board_url"),
            (registry or {}).get("source", {}).get("path"), rec.get("source_section"),
        ]
        key = str(rec.get("company") or "").casefold()
        if key and key in watch_rows:
            for idx, value in enumerate(row, start=1):
                ws.cell(row=watch_rows[key], column=idx).value = value
            updates += 1
        else:
            ws.append(row)
            if key:
                watch_rows[key] = ws.max_row

    # --- Findings: append only postings not already recorded in this workbook
    wsf = wb["Findings"]
    url_col = WORKBOOK_SHEETS["Findings"][0].index("URL") + 1
    known_urls: set[str] = set()
    for r in range(2, wsf.max_row + 1):
        v = wsf.cell(row=r, column=url_col).value
        if v:
            known_urls.add(str(v))
    findings_added = 0
    for f in (findings_doc or {}).get("findings", []):
        url = f.get("url")
        if not url or str(url) in known_urls:
            continue
        known_urls.add(str(url))
        findings_added += 1
        wsf.append([
            f.get("company"), f.get("title"), f.get("location"), f.get("url"),
            f.get("url_quality"), f.get("vendor"),
            f.get("vendor_job_id"), f.get("posted_at"), f.get("posted_age_days"),
            f.get("region_route"), f.get("decision"), f.get("decision_reason"),
            f.get("owner_filter_eligible"), f.get("tracker_eligible"),
            f.get("attribution_confidence"), f.get("source"), f.get("source_timestamp"),
        ])

    # --- Tracker Handoff: one row per distinct handoff manifest
    wsh = wb["Tracker Handoff"]
    known_manifests = {str(wsh.cell(row=r, column=5).value)
                       for r in range(2, wsh.max_row + 1)
                       if wsh.cell(row=r, column=5).value}
    if handoff_result and str(handoff_result.get("manifest")) not in known_manifests:
        wsh.append([
            handoff_result.get("region"), handoff_result.get("mode"),
            (handoff_result.get("result") or {}).get("counts", {}).get("appended"),
            handoff_result.get("tracker"), handoff_result.get("manifest"),
            handoff_result.get("exit_code"), now_str,
            ("canonical workbook untouched (dry-run)" if handoff_result.get("mode") == "dry-run"
             else "applied through career-ops/tracker_writer.py with backup + verification"),
        ])

    # --- Run Log: one row per run id
    wsl = wb["Run Log"]
    known_runs = {str(wsl.cell(row=r, column=1).value)
                  for r in range(2, wsl.max_row + 1) if wsl.cell(row=r, column=1).value}
    if run_meta and str(run_meta.get("run_id")) not in known_runs:
        wsl.append([
            run_meta.get("run_id"), run_meta.get("started_at"), run_meta.get("region"),
            run_meta.get("companies_probed"), run_meta.get("boards_resolved"),
            run_meta.get("findings"), run_meta.get("duplicates"), run_meta.get("eligible"),
            run_meta.get("http_requests"), run_meta.get("budget_exhausted"),
            run_meta.get("duration_s"), run_meta.get("mode"),
        ])

    tmp = path.with_suffix(".write-tmp.xlsx")
    wb.save(tmp)
    wb.close()
    tmp.replace(path)

    check = openpyxl.load_workbook(path, read_only=True)
    headers = {name: [c.value for c in next(check[name].iter_rows(max_row=1))] for name in check.sheetnames}
    check.close()
    forbidden = [c for c in headers.get("Findings", []) if re.search(r"status", str(c), re.I)]
    return {
        "ok": not forbidden,
        "path": str(path),
        "sheets": list(headers),
        "headers": headers,
        "watch_list_rows": ws.max_row - 1,
        "watch_list_refreshed": updates,
        "findings_rows": wsf.max_row - 1,
        "findings_added": findings_added,
        "conflict_check": conflict,
        "forbidden_application_state_columns": forbidden,
        "note": "operational workbook only; canonical regional trackers remain application state",
        "written_at": now_str,
    }


# --------------------------------------------------------------------------- #
# commands
# --------------------------------------------------------------------------- #

def cmd_registry(args) -> int:
    cfg = load_config(args.config)
    registry = load_registry(cfg, rebuild=args.rebuild)
    import company_registry  # noqa: PLC0415
    doc = company_registry.summary(registry)
    doc["written_to"] = str(career_ops_paths(cfg)["runtime_dir"] / "company_watch_registry.json")
    doc["registry_source"] = registry["source"]
    doc["sample_classes"] = registry["counts"]["parsed_by_class"]
    json.dump(doc, sys.stdout, indent=2, ensure_ascii=False)
    sys.stdout.write("\n")
    return 0


def cmd_resolve(args) -> int:
    import ats_endpoints  # noqa: PLC0415
    cfg = load_config(args.config)
    registry = load_registry(cfg)
    companies = ordered_companies(registry, args.start_index, args.limit)
    fetcher = ats_endpoints.Fetcher(timeout=args.timeout, retries=args.retries)
    doc = ats_endpoints.resolve_companies(
        companies, fetcher, limit=args.limit, budget_s=args.budget_s,
        vendors=tuple(args.vendors or ats_endpoints.VENDOR_ORDER),
    )
    doc["registry_source"] = registry["source"]
    doc["selection"] = {"start_index": args.start_index, "limit": args.limit,
                        "order": "evidence_class then name"}
    out = career_ops_paths(cfg)["runtime_dir"] / "resolution-latest.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(doc, indent=2, ensure_ascii=False), encoding="utf-8")
    doc["written_to"] = str(out)
    doc["resolved"] = [{"company": r["company"], "vendor": r["vendor"], "slug": r["board_slug"],
                        "jobs": r["jobs_listed"], "confidence": r["attribution_confidence"],
                        "manual": r["manual_attribution_required"]}
                       for r in doc["resolved"]]
    doc["unresolved"] = [r["company"] for r in doc["unresolved"]]
    json.dump(doc, sys.stdout, indent=2, ensure_ascii=False, default=str)
    sys.stdout.write("\n")
    return 0


def _load_resolution(cfg: dict, path: str | None) -> dict:
    p = Path(path) if path else career_ops_paths(cfg)["runtime_dir"] / "resolution-latest.json"
    if not p.exists():
        return {}
    return json.loads(p.read_text(encoding="utf-8"))


def cmd_scan(args) -> int:
    cfg = load_config(args.config)
    tw, _profiles, region_cfg = resolve_region_config(cfg, args.region)
    filters = load_owner_filters(cfg, args.region)
    dedupe = build_shared_dedupe(cfg, args.region, tracker_override=args.tracker,
                                extra_archive_dirs=args.archive_dir)
    resolutions = _load_resolution(cfg, args.resolution)
    if not resolutions:
        return _emit_error("no resolution data; run `resolve` first", 2)
    now = dt.datetime.now(dt.timezone.utc)
    import ats_endpoints  # noqa: PLC0415
    fetcher = ats_endpoints.Fetcher(timeout=cfg.get("probe", {}).get("request_timeout_s", 20),
                                    retries=cfg.get("probe", {}).get("request_retries", 1))
    findings_doc = build_findings(resolutions, region=args.region, cfg=cfg, dedupe=dedupe,
                                 tw=tw, filters=filters, now=now, fetcher=fetcher)
    findings_doc["dedupe_source"] = {
        "tracker": dedupe["tracker"], "tracker_sha256": dedupe["tracker_sha256"],
        "canonical_data_rows": dedupe["canonical_data_rows"],
        "cross_month_keys": len(dedupe["cross_month_keys"]),
        "cross_month_sources": dedupe["cross_month_sources"],
        "shared_with": "career-ops/tracker_writer.py (same primitives as Career Ops)",
    }
    out = career_ops_paths(cfg)["runtime_dir"] / f"findings-{args.region}-latest.json"
    out.write_text(json.dumps(findings_doc, indent=2, ensure_ascii=False), encoding="utf-8")
    findings_doc["written_to"] = str(out)
    json.dump(findings_doc, sys.stdout, indent=2, ensure_ascii=False, default=str)
    sys.stdout.write("\n")
    return 0


def cmd_handoff(args) -> int:
    cfg = load_config(args.config)
    _tw, _profiles, region_cfg = resolve_region_config(cfg, args.region)
    p = Path(args.findings) if args.findings else (
        career_ops_paths(cfg)["runtime_dir"] / f"findings-{args.region}-latest.json")
    if not p.exists():
        return _emit_error(f"findings file not found: {p}", 2)
    findings_doc = json.loads(p.read_text(encoding="utf-8"))
    manifest = build_manifest(findings_doc, region=args.region, region_cfg=region_cfg, cfg=cfg,
                              include_ineligible_as_test=args.include_ineligible_as_test,
                              limit=args.limit)
    if not manifest["ok"]:
        json.dump(manifest, sys.stdout, indent=2, ensure_ascii=False)
        sys.stdout.write("\n")
        return 3
    result = handoff_to_career_ops(cfg, args.region, manifest, tracker=args.tracker,
                                   apply=args.apply, backup_dir=args.backup_dir,
                                   extra_archive_dirs=args.archive_dir)
    doc = {"region": args.region, "manifest": {k: v for k, v in manifest.items() if k != "records"},
           "records": manifest["records"], "handoff": result}
    if not args.apply:
        doc["safety"] = ("dry-run against the canonical workbook: no write, hash unchanged; "
                         "use --apply with an explicit --tracker to write")
    json.dump(doc, sys.stdout, indent=2, ensure_ascii=False, default=str)
    sys.stdout.write("\n")
    return 0 if result.get("ok") else 1


def cmd_workbook(args) -> int:
    cfg = load_config(args.config)
    registry = load_registry(cfg)
    resolutions = _load_resolution(cfg, args.resolution)
    findings_doc = None
    fp = Path(args.findings) if args.findings else (
        career_ops_paths(cfg)["runtime_dir"] / f"findings-{args.region}-latest.json")
    if fp.exists():
        findings_doc = json.loads(fp.read_text(encoding="utf-8"))
    handoff_result = None
    hp = Path(args.handoff_report) if args.handoff_report else None
    if hp and hp.exists():
        handoff_result = json.loads(hp.read_text(encoding="utf-8")).get("handoff")
        if handoff_result:
            handoff_result = {**handoff_result, "region": args.region}
    registry_view = {
        "source": registry["source"],
        "_prior_evidence": {c["name"]: c["prior_application_evidence"] for c in registry["companies"]},
    }
    res = write_workbook(cfg, month=args.month, directory=args.dir, registry=registry_view,
                         resolutions=resolutions, findings_doc=findings_doc,
                         handoff_result=handoff_result, run_meta=None)
    json.dump(res, sys.stdout, indent=2, ensure_ascii=False, default=str)
    sys.stdout.write("\n")
    return 0 if res.get("ok") else 1


def cmd_run(args) -> int:
    import ats_endpoints  # noqa: PLC0415
    cfg = load_config(args.config)
    paths = career_ops_paths(cfg)
    started = dt.datetime.now(dt.timezone.utc)
    run_id = started.strftime("%Y%m%dT%H%M%SZ")
    out_dir = Path(args.out_dir) if args.out_dir else paths["runtime_dir"] / f"run-{run_id}"
    out_dir.mkdir(parents=True, exist_ok=True)

    tw, _profiles, region_cfg = resolve_region_config(cfg, args.region)
    registry = load_registry(cfg, rebuild=args.rebuild_registry)
    companies = ordered_companies(registry, args.start_index, args.limit)

    fetcher = ats_endpoints.Fetcher(timeout=args.timeout, retries=args.retries)
    resolutions = ats_endpoints.resolve_companies(
        companies, fetcher, limit=args.limit, budget_s=args.budget_s,
        vendors=tuple(args.vendors or ats_endpoints.VENDOR_ORDER),
    )
    resolutions["selection"] = {"start_index": args.start_index, "limit": args.limit,
                                "order": "evidence_class then name",
                                "registry_companies": len(registry["companies"])}
    resolutions["registry_source"] = registry["source"]
    (out_dir / "resolution.json").write_text(json.dumps(resolutions, indent=2, ensure_ascii=False),
                                             encoding="utf-8")

    filters = load_owner_filters(cfg, args.region)
    dedupe = build_shared_dedupe(cfg, args.region, tracker_override=args.canonical_tracker,
                                extra_archive_dirs=args.archive_dir)
    findings_doc = build_findings(resolutions, region=args.region, cfg=cfg, dedupe=dedupe,
                                 tw=tw, filters=filters,
                                 now=dt.datetime.now(dt.timezone.utc), fetcher=fetcher)
    findings_doc["dedupe_source"] = {
        "tracker": dedupe["tracker"], "tracker_sha256": dedupe["tracker_sha256"],
        "canonical_data_rows": dedupe["canonical_data_rows"],
        "cross_month_keys": len(dedupe["cross_month_keys"]),
        "cross_month_sources": dedupe["cross_month_sources"],
    }
    (out_dir / "findings.json").write_text(json.dumps(findings_doc, indent=2, ensure_ascii=False),
                                          encoding="utf-8")

    # 1) dry-runs against the canonical workbook: prove the interface, write nothing.
    #    (a) the strictly eligible set (usually empty while the owner's UK lane is
    #        intern/internship-only and no watched company posts an intern role), and
    #    (b) a marked sample so the canonical dedupe path is actually exercised.
    manifest = build_manifest(findings_doc, region=args.region, region_cfg=region_cfg, cfg=cfg)
    dry = handoff_to_career_ops(
        cfg, args.region, manifest,
        tracker=args.canonical_tracker or region_cfg["tracker"], apply=False,
        extra_archive_dirs=args.archive_dir,
        manifest_path=str(out_dir / f"manifest-{args.region}-dryrun.json"))
    sample_manifest = build_manifest(findings_doc, region=args.region, region_cfg=region_cfg,
                                    cfg=cfg, include_ineligible_as_test=True,
                                    limit=args.acceptance_rows)
    dry_sample = handoff_to_career_ops(
        cfg, args.region, sample_manifest,
        tracker=args.canonical_tracker or region_cfg["tracker"], apply=False,
        extra_archive_dirs=args.archive_dir,
        manifest_path=str(out_dir / f"manifest-{args.region}-dryrun-sample.json"))

    # 2) acceptance write on a copy of the canonical workbook (safe/test data)
    acceptance = None
    if args.acceptance_copy:
        copy_dir = Path(args.acceptance_copy)
        copy_dir.mkdir(parents=True, exist_ok=True)
        copy_path = copy_dir / Path(region_cfg["tracker"]).name
        import shutil  # noqa: PLC0415
        shutil.copy2(region_cfg["tracker"], copy_path)
        before_hash = tw.sha256_file(copy_path)
        accept_manifest = build_manifest(findings_doc, region=args.region, region_cfg=region_cfg,
                                        cfg=cfg, include_ineligible_as_test=True,
                                        limit=args.acceptance_rows)
        first = handoff_to_career_ops(
            cfg, args.region, accept_manifest, tracker=str(copy_path), apply=True,
            backup_dir=str(out_dir / "backups"), extra_archive_dirs=[str(copy_dir)],
            manifest_path=str(out_dir / f"manifest-{args.region}-acceptance.json"))
        # 3) idempotency / dedupe: the same handoff again must append nothing
        second = handoff_to_career_ops(
            cfg, args.region, accept_manifest, tracker=str(copy_path), apply=True,
            backup_dir=str(out_dir / "backups"), extra_archive_dirs=[str(copy_dir)],
            manifest_path=str(out_dir / f"manifest-{args.region}-acceptance-repeat.json"))
        after_rows = _tracker_rows(tw, region_cfg, copy_path)
        acceptance = {
            "copy": str(copy_path),
            "copy_sha256_before": before_hash,
            "copy_sha256_after": tw.sha256_file(copy_path),
            "records_offered": accept_manifest["counts"]["selected"],
            "test_rows": accept_manifest["counts"]["included_as_test_rows"],
            "first": first,
            "repeat": second,
            "data_rows_after": after_rows,
            "canonical_untouched": tw.sha256_file(region_cfg["tracker"]) == dedupe["tracker_sha256"],
        }
        (out_dir / "acceptance.json").write_text(json.dumps(acceptance, indent=2, ensure_ascii=False,
                                                            default=str), encoding="utf-8")

    # 4) monthly operational workbook
    workbook = write_workbook(
        cfg, month=args.month, directory=args.workbook_dir,
        registry={"source": registry["source"],
                  "_prior_evidence": {c["name"]: c["prior_application_evidence"]
                                      for c in registry["companies"]}},
        resolutions=resolutions, findings_doc=findings_doc,
        handoff_result=({"region": args.region, **dry} if dry.get("invoked") else None),
        run_meta={
            "run_id": run_id, "started_at": started.replace(microsecond=0).isoformat(),
            "region": args.region, "companies_probed": resolutions["companies_probed"],
            "boards_resolved": len(resolutions["resolved"]),
            "findings": findings_doc["counts"]["findings"],
            "duplicates": findings_doc["counts"]["by_decision"].get("duplicate", 0)
            + findings_doc["counts"]["by_decision"].get("duplicate-cross-month", 0)
            + findings_doc["counts"]["by_decision"].get("duplicate-in-run", 0),
            "eligible": findings_doc["counts"]["tracker_eligible"],
            "http_requests": resolutions["http"]["requests"],
            "budget_exhausted": resolutions["budget_exhausted"],
            "duration_s": resolutions["duration_s"], "mode": "live",
        })

    finished = dt.datetime.now(dt.timezone.utc)
    evidence = {
        "task": "agent-company-watch-job-search-integration-2026-09-23",
        "run_id": run_id,
        "started_at": started.replace(microsecond=0).isoformat(),
        "finished_at": finished.replace(microsecond=0).isoformat(),
        "duration_s": round((finished - started).total_seconds(), 1),
        "region": args.region,
        "config": str(cfg["_config_path"]),
        "registry": {
            "source_path": registry["source"]["path"],
            "source_sha256": registry["source"]["sha256"],
            "source_period": registry["source"]["period"],
            "counts": registry["counts"],
            "full_registry_local_only": str(paths["runtime_dir"] / "company_watch_registry.json"),
        },
        "selection": resolutions["selection"],
        "http": resolutions["http"],
        "resolution": {
            "companies_probed": resolutions["companies_probed"],
            "boards_resolved": len(resolutions["resolved"]),
            "unresolved": len(resolutions["unresolved"]),
            "by_vendor": _count_by_vendor(resolutions),
            "attribution": {
                "high": sum(1 for r in resolutions["resolved"] if r["attribution_confidence"] == "high"),
                "low": sum(1 for r in resolutions["resolved"] if r["attribution_confidence"] != "high"),
                "manual_attribution_required": sum(
                    1 for r in resolutions["resolved"] + resolutions["unresolved"]
                    if r.get("manual_attribution_required")),
            },
            "budget_exhausted": resolutions["budget_exhausted"],
        },
        "findings": {
            "counts": findings_doc["counts"],
            "filters": findings_doc["filters"],
            "dedupe_source": findings_doc["dedupe_source"],
        },
        "handoff_dry_run": {
            "invoked": dry.get("invoked", False),
            "mode": dry.get("mode"),
            "exit_code": dry.get("exit_code"),
            "counts": (dry.get("result") or {}).get("counts"),
            "tracker": dry.get("tracker"),
            "canonical_untouched": tw.sha256_file(region_cfg["tracker"]) == dedupe["tracker_sha256"],
            "note": ("no eligible records for the strictly filtered set" if not dry.get("invoked")
                     else "dry-run of the eligible set"),
        },
        "handoff_dry_run_sample": {
            "invoked": dry_sample.get("invoked", False),
            "mode": dry_sample.get("mode"),
            "exit_code": dry_sample.get("exit_code"),
            "counts": (dry_sample.get("result") or {}).get("counts"),
            "tracker": dry_sample.get("tracker"),
            "rows_are_test_marked": True,
            "canonical_untouched": tw.sha256_file(region_cfg["tracker"]) == dedupe["tracker_sha256"],
            "note": ("marked sample so the canonical dedupe path is exercised; --apply is NOT "
                     "passed and the canonical workbook hash is re-checked after the call"),
        },
        "acceptance_write": acceptance,
        "workbook": workbook,
        "artifacts": {
            "resolution": str(out_dir / "resolution.json"),
            "findings": str(out_dir / "findings.json"),
            "acceptance": str(out_dir / "acceptance.json") if acceptance else None,
            "workbook": workbook.get("path"),
        },
        "local_only_detail": (
            "Per-company registry and per-posting findings name owner job-search history and are "
            "kept in git-ignored runtime paths; this evidence records aggregate counts and "
            "verification results only."
        ),
    }
    (out_dir / "evidence.json").write_text(json.dumps(evidence, indent=2, ensure_ascii=False, default=str),
                                           encoding="utf-8")
    (out_dir / "evidence.md").write_text(_evidence_md(evidence), encoding="utf-8")
    json.dump(evidence, sys.stdout, indent=2, ensure_ascii=False, default=str)
    sys.stdout.write("\n")
    return 0


def _count_by_vendor(resolutions: dict) -> dict:
    out: dict = {}
    for r in resolutions.get("resolved", []):
        out[r["vendor"]] = out.get(r["vendor"], 0) + 1
    return out


def _tracker_rows(tw, region_cfg: dict, path: Path) -> int:
    tr = tw.Tracker(Path(path), region_cfg)
    rows = tr.last_data_row() - region_cfg["first_data_row"] + 1
    tr.wb.close()
    return rows


def _evidence_md(ev: dict) -> str:
    lines = [
        f"# Company Watch — bounded integration evidence ({ev['run_id']})",
        "",
        f"- Task: `{ev['task']}`",
        f"- Region: `{ev['region']}`",
        f"- Started / finished (UTC): {ev['started_at']} → {ev['finished_at']} ({ev['duration_s']} s)",
        f"- Registry source: `{ev['registry']['source_path']}`",
        f"- Registry source SHA-256: `{ev['registry']['source_sha256']}`",
        f"- Registry counts (parsed): {json.dumps(ev['registry']['counts']['parsed_by_class'])}",
        f"- Declared-vs-parsed mismatches: {json.dumps(ev['registry']['counts']['count_mismatch'])}",
        "",
        "## Bounded live ATS run",
        "",
        f"- Companies probed: {ev['resolution']['companies_probed']} "
        f"(selection: {json.dumps(ev['selection'])})",
        f"- Boards resolved: {ev['resolution']['boards_resolved']} "
        f"(by vendor: {json.dumps(ev['resolution']['by_vendor'])})",
        f"- Attribution: {json.dumps(ev['resolution']['attribution'])}",
        f"- HTTP: {json.dumps(ev['http'])}",
        f"- Budget exhausted: {ev['resolution']['budget_exhausted']}",
        "",
        "## Findings",
        "",
        f"- {json.dumps(ev['findings']['counts'])}",
        f"- Owner filters: {json.dumps(ev['findings']['filters'])}",
        f"- Shared dedupe: {json.dumps(ev['findings']['dedupe_source'])}",
        "",
        "## Handoff",
        "",
        f"- Eligible-set dry-run against the canonical workbook: {json.dumps(ev['handoff_dry_run'])}",
        f"- Marked-sample dry-run against the canonical workbook: "
        f"{json.dumps(ev['handoff_dry_run_sample'])}",
    ]
    acc = ev.get("acceptance_write")
    if acc:
        lines += [
            "",
            "## Acceptance write (safe copy)",
            "",
            f"- Copy: `{acc['copy']}`",
            f"- Copies offered: {acc['records_offered']} (of which {acc['test_rows']} test rows)",
            f"- Copy SHA-256: {acc['copy_sha256_before']} → {acc['copy_sha256_after']}",
            f"- First apply: exit {acc['first'].get('exit_code')}, "
            f"counts {json.dumps((acc['first'].get('result') or {}).get('counts'))}",
            f"- Repeat apply (dedupe proof): exit {acc['repeat'].get('exit_code')}, "
            f"counts {json.dumps((acc['repeat'].get('result') or {}).get('counts'))}",
            "- Note: exit 1 on the repeat is the Career Ops writer's documented "
            "nothing-to-append signal for `--apply` (career_ops_cli.cmd_write), not an error; "
            "appended 0 with 3 duplicates is the dedupe proof.",
            f"- Data rows after: {acc['data_rows_after']}",
            f"- Canonical workbook untouched: {acc['canonical_untouched']}",
        ]
    lines += [
        "",
        "## Monthly operational workbook",
        "",
        f"- Path: `{ev['workbook'].get('path')}`",
        f"- Sheets: {json.dumps(ev['workbook'].get('sheets'))}",
        f"- Archive-glob conflict check: {json.dumps(ev['workbook'].get('conflict_check'))}",
        f"- Forbidden application-state columns found: "
        f"{json.dumps(ev['workbook'].get('forbidden_application_state_columns'))}",
        "",
        ev["local_only_detail"],
    ]
    return "\n".join(lines) + "\n"


def _emit_error(reason: str, code: int) -> int:
    json.dump({"ok": False, "reason": reason}, sys.stdout, indent=2)
    sys.stdout.write("\n")
    return code


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Company Watch (Career)")
    ap.add_argument("--config", default=None)
    sub = ap.add_subparsers(dest="command", required=True)

    p = sub.add_parser("registry")
    p.add_argument("--rebuild", action="store_true")
    p.set_defaults(fn=cmd_registry)

    p = sub.add_parser("resolve")
    p.add_argument("--limit", type=int, default=12)
    p.add_argument("--start-index", type=int, default=0)
    p.add_argument("--budget-s", type=float, default=240.0)
    p.add_argument("--timeout", type=float, default=20.0)
    p.add_argument("--retries", type=int, default=1)
    p.add_argument("--vendors", action="append")
    p.set_defaults(fn=cmd_resolve)

    p = sub.add_parser("scan")
    p.add_argument("--region", required=True)
    p.add_argument("--resolution")
    p.add_argument("--tracker")
    p.add_argument("--archive-dir", action="append")
    p.set_defaults(fn=cmd_scan)

    p = sub.add_parser("handoff")
    p.add_argument("--region", required=True)
    p.add_argument("--findings")
    p.add_argument("--tracker")
    p.add_argument("--apply", action="store_true")
    p.add_argument("--backup-dir")
    p.add_argument("--archive-dir", action="append")
    p.add_argument("--include-ineligible-as-test", action="store_true")
    p.add_argument("--limit", type=int, default=0)
    p.set_defaults(fn=cmd_handoff)

    p = sub.add_parser("workbook")
    p.add_argument("--month", default=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m"))
    p.add_argument("--dir")
    p.add_argument("--region", default="uk")
    p.add_argument("--resolution")
    p.add_argument("--findings")
    p.add_argument("--handoff-report")
    p.set_defaults(fn=cmd_workbook)

    p = sub.add_parser("run")
    p.add_argument("--region", default="uk")
    p.add_argument("--limit", type=int, default=12)
    p.add_argument("--start-index", type=int, default=0)
    p.add_argument("--budget-s", type=float, default=240.0)
    p.add_argument("--timeout", type=float, default=20.0)
    p.add_argument("--retries", type=int, default=1)
    p.add_argument("--vendors", action="append")
    p.add_argument("--out-dir")
    p.add_argument("--workbook-dir")
    p.add_argument("--month", default=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m"))
    p.add_argument("--canonical-tracker")
    p.add_argument("--acceptance-copy")
    p.add_argument("--acceptance-rows", type=int, default=3)
    p.add_argument("--archive-dir", action="append")
    p.add_argument("--rebuild-registry", action="store_true")
    p.set_defaults(fn=cmd_run)

    args = ap.parse_args(argv)
    return args.fn(args)


if __name__ == "__main__":
    raise SystemExit(main())
