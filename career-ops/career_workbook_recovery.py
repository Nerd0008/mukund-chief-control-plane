"""Conservative, explicit owner-approved workbook recovery; no mailbox calls.

Patch only designated XML cells/table extents, preserving all other ZIP members.
Private audit contains per-cell values; never commit it to source control.
"""
import copy, hashlib, json, os, shutil, tempfile, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',NS)
def tag(n): return '{'+NS+'}'+n
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def inline(address,value):
 c=ET.Element(tag('c'),{'r':address,'t':'inlineStr'});t=ET.SubElement(ET.SubElement(c,tag('is')),tag('t'));t.text=str(value);return c

def patch_package(path, changes, audit_dir, expected_hash):
 """Verified backup + compare-and-swap; changes map ZIP member to bytes."""
 path=Path(path);audit_dir=Path(audit_dir);audit_dir.mkdir(parents=True,exist_ok=True)
 lock=path.with_suffix(path.suffix+'.career-mail.lock')
 fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY);tmp=None
 try:
  if sha(path)!=expected_hash: raise ValueError('concurrent workbook drift before backup')
  backup=audit_dir/(path.stem+'-'+expected_hash+'.xlsx');shutil.copy2(path,backup)
  if sha(backup)!=expected_hash: raise ValueError('backup verification failed')
  handle,name=tempfile.mkstemp(suffix='.xlsx',dir=path.parent);os.close(handle);tmp=Path(name)
  with zipfile.ZipFile(backup) as source,zipfile.ZipFile(tmp,'w') as dest:
   for info in source.infolist(): dest.writestr(info,changes.get(info.filename,source.read(info.filename)))
  with zipfile.ZipFile(backup) as source,zipfile.ZipFile(tmp) as dest:
   for name in source.namelist():
    if name not in changes and source.read(name)!=dest.read(name): raise ValueError('unrelated workbook part changed')
  if sha(path)!=expected_hash: raise ValueError('concurrent workbook drift before commit')
  os.replace(tmp,path);tmp=None
  return {'before_hash':expected_hash,'after_hash':sha(path),'backup':str(backup),'changed_parts':list(changes)}
 finally:
  if tmp is not None: tmp.unlink(missing_ok=True)
  os.close(fd);lock.unlink()

def restore_uk(profiles,audit_dir):
 from career_mail_tracker import WorkbookTracker,HEADERS
 clean=copy.deepcopy(profiles);clean.pop('accepted_mail_state',None)
 current,fingerprints=WorkbookTracker(clean,audit_dir).read()
 accepted=WorkbookTracker(profiles,audit_dir);recovered,_=accepted.read()
 cfg=profiles['regions']['uk'];path=Path(cfg['tracker']);root=accepted.accepted_state
 receipt=json.loads((root/'result.json').read_text());snap=root/'staged/uk.xlsx'
 if sha(snap)!=receipt['after_hashes']['uk']: raise ValueError('snapshot source hash invalid')
 source_records,_=WorkbookTracker({'regions':{'uk':{**cfg,'tracker':str(snap)}}},audit_dir).read()
 source_records=[r for r in source_records if r.get('seen_message_ids')]
 if len(source_records)!=17: raise ValueError('accepted metadata record count changed')
 from openpyxl import load_workbook
 src=load_workbook(snap,read_only=True);sw=src[cfg['sheet']]
 headers={c.value:c.column for c in sw[cfg['header_row']] if c.value}
 owned=[cfg['mail_identity']['header']]+list(HEADERS.values())
 if any(h not in headers for h in owned): raise ValueError('snapshot metadata headers missing')
 with zipfile.ZipFile(path) as z: tree=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
 data=tree.find(tag('sheetData'));rows={int(r.attrib['r']):r for r in data};audit=[]
 def put(row,col,value):
  from openpyxl.utils import get_column_letter
  address=get_column_letter(col)+str(row);r=rows.get(row)
  if r is None: r=ET.SubElement(data,tag('row'),{'r':str(row)});rows[row]=r
  old=next((c for c in r if c.attrib.get('r')==address),None)
  if old is not None:
   # Never overwrite current metadata: restoration is for missing cells only.
   if old.find(tag('v')) is not None or old.find(tag('is')) is not None: raise ValueError('current metadata cell occupied: '+address)
   r.remove(old)
  if value is not None: r.append(inline(address,value))
  audit.append({'cell':address,'before':None,'after':value})
 for h in owned: put(cfg['header_row'],headers[h],h)
 targets={r['application_id']:r for r in recovered if r['region']=='uk' and r.get('seen_message_ids')}
 for old in source_records:
  target=targets.get(old['application_id'])
  if target is None: raise ValueError('accepted identity missing')
  for h in owned: put(target['row'],headers[h],sw.cell(old['row'],headers[h]).value)
 src.close()
 # SpreadsheetML cells must retain column order.
 from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
 for r in data:
  cells=list(r);cells.sort(key=lambda c:column_index_from_string(coordinate_from_string(c.attrib['r'])[0]));r[:]=cells
 dimension=tree.find(tag('dimension'));dimension.set('ref','A1:AA'+str(max(rows)))
 if sha(snap)!=receipt['after_hashes']['uk']: raise ValueError('snapshot changed before commit')
 result=patch_package(path,{'xl/worksheets/sheet1.xml':ET.tostring(tree,encoding='utf-8',xml_declaration=True)},audit_dir,fingerprints['uk'])
 records,_=WorkbookTracker(clean,audit_dir).read()
 persisted=[r for r in records if r['region']=='uk' and r.get('seen_message_ids')]
 if len(persisted)!=17 or len(set(m for r in persisted for m in r['seen_message_ids']))!=20: raise ValueError('persisted acceptance dedupe verification failed')
 result.update(restored=17,skipped=0,authorised_signals=20,per_cell_audit=audit)
 Path(audit_dir,'uk-metadata-audit.json').write_text(json.dumps(result,indent=2,default=str))
 return {k:v for k,v in result.items() if k!='per_cell_audit'}
