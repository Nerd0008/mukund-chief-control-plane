"""Bounded evidence-only Gmail status reconciliation for the UK tracker."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import re

import openpyxl
from career_mail_parser import classify_recruitment, mail_fields, identity_key
from career_mail_matching import has_full_role, company_named
from tracker_set_status import apply_changes
import tracker_writer as tw

RECEIPT = {'application_confirmation', 'application_received'}
REJECTION = re.compile(r'\b(?:we (?:have )?decided not to (?:move forward|progress|proceed)|'
                       r'(?:we (?:have )?)?decided to progress with other candidates|'
                       r'not (?:be |been |to )?(?:progressing|proceeding|successful|selected|move forward|moving forward)|(?:your application|you) (?:has |have )?'
                       r'(?:been )?(?:unsuccessful|not (?:been )?selected)|unfortunately.{0,100}'
                       r'(?:not (?:move forward|progress|proceed)|unsuccessful))\b', re.I)


def evidence_status(subject, body):
    kind = classify_recruitment(subject, body)
    if kind == 'rejection' and REJECTION.search(subject + '\n' + body):
        return 'Rejected'
    if kind in RECEIPT:
        return 'Applied'
    return None


def plan(messages, ws, *, first=3, processed=(), limit=200):
    if len(messages) > limit:
        raise ValueError('message limit exceeded')
    seen = set(processed)
    receipts = []
    for raw in messages:
        mail = mail_fields(raw)
        mid = raw.get('id') or raw.get('message_id') or mail.get('_gmail_id')
        if not mid or mid in seen:
            continue
        status = evidence_status(mail.get('subject') or '', mail.get('body') or '')
        if not status:
            continue
        matches = []
        text = (mail.get('subject') or '') + '\n' + (mail.get('body') or '')
        for row in range(first, ws.max_row + 1):
            company, role = ws.cell(row, 2).value, ws.cell(row, 3).value
            if company and role and company_named({'company': company, 'canonical_job_host': ''}, mail) and has_full_role(role, text):
                matches.append(row)
        if len(matches) != 1:
            receipts.append({'message_id': mid, 'decision': 'review', 'reason': f'{len(matches)} matching rows'})
            continue
        row = matches[0]
        current = ws.cell(row, 6).value
        # A receipt cannot reverse a later rejection. Gmail ordering is not a state oracle.
        if current == 'Rejected' and status == 'Applied':
            decision = 'unchanged'
        elif current == status:
            decision = 'unchanged'
        else:
            decision = 'update'
            ws.cell(row, 6).value = status
        receipts.append({'message_id': mid, 'row': row, 'company': ws.cell(row, 2).value,
                         'role': ws.cell(row, 3).value, 'from': current, 'to': status,
                         'decision': decision, 'subject': mail.get('subject') or '',
                         'evidence': (REJECTION.search(text).group(0) if status == 'Rejected'
                                      else classify_recruitment(mail.get('subject') or '', mail.get('body') or ''))})
        seen.add(mid)
    return receipts


def reconcile_file(path, messages, *, audit_path, apply=False, limit=200):
    path, audit_path = Path(path), Path(audit_path)
    audit = json.loads(audit_path.read_text(encoding='utf-8')) if audit_path.exists() else {'processed_ids': []}
    wb = openpyxl.load_workbook(path)
    try:
        receipts = plan(messages, wb['Jobs'], processed=audit['processed_ids'], limit=limit)
    finally:
        wb.close()
    result = {'applied': False, 'receipts': receipts}
    changes = [{'row': r['row'], 'status': r['to'], 'expect_status': r['from']}
               for r in receipts if r['decision'] == 'update']
    if not apply:
        return result
    if changes:
        write = apply_changes(path, changes, profiles=tw.load_profiles(),
                              backup_dir=audit_path.parent / 'backups')
        if not write.get('written'):
            raise ValueError('Status write failed: ' + repr(write))
        result['write'] = write
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit['processed_ids'] = sorted(set(audit['processed_ids']) | {r['message_id'] for r in receipts if r['decision'] != 'review'})
    audit.setdefault('receipts', []).extend(receipts)
    audit_path.write_text(json.dumps(audit, indent=2, default=str), encoding='utf-8')
    result['applied'] = True
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--messages', required=True, help='JSON export containing at most 200 full Gmail messages')
    parser.add_argument('--tracker', default=tw.region_config(tw.load_profiles(), 'uk')['tracker'])
    parser.add_argument('--audit', required=True)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    messages = json.loads(Path(args.messages).read_text(encoding='utf-8'))
    print(json.dumps(reconcile_file(args.tracker, messages, audit_path=args.audit, apply=args.apply), indent=2, default=str))


if __name__ == '__main__':
    main()
