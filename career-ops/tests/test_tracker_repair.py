import pathlib,sys
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
import openpyxl,pytest
import tracker_writer as tw
import tracker_repair as repair

def workbook(tmp_path,status='Not Applied'):
    p=tmp_path/'uk.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['UK jobs']);ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status'])
    ws.append(['2026-10-17','Galaxy','Cyber Internship','https://job-boards.greenhouse.io/galaxy/jobs/123?gh_src=one',None,status])
    ws.append(['2026-10-17','Galaxy','Cyber Internship','https://job-boards.greenhouse.io/galaxy/jobs/123',None,status])
    ws.append(['2026-10-16','Keep Owner','Cyber Intern','https://employer.com/job/456','2026-11-01','Applied'])
    from openpyxl.worksheet.datavalidation import DataValidation
    spec=tw.region_config(tw.load_profiles(),'uk')['validations']['F']
    validation=DataValidation(type='list',formula1=spec['formula1']);ws.add_data_validation(validation);validation.add('F3:F500')
    wb.save(p);wb.close();return p

def test_dry_run_preserves_source(tmp_path):
    p=workbook(tmp_path);before=tw.sha256_file(p)
    result=repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4])
    assert not result['applied'] and result['verification']['ok']
    assert tw.sha256_file(p)==before and not list(tmp_path.glob('*.repair-*.xlsx'))

def test_apply_archives_exact_duplicate_and_preserves_owner_cells(tmp_path):
    p=workbook(tmp_path);before=tw.sha256_file(p)
    result=repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    assert result['applied'] and result['verification']['ok']
    assert tw.sha256_file(result['backup'])==before
    wb=openpyxl.load_workbook(p);ws=wb['Jobs']
    assert ws.max_row==4 and ws.max_column==6
    assert [ws.cell(4,c).value for c in (2,3,5,6)]==['Keep Owner','Cyber Intern','2026-11-01','Applied']
    assert all(ws.cell(r,4).hyperlink.target==ws.cell(r,4).value for r in (3,4))
    wb.close()

def test_never_remove_an_applied_record(tmp_path):
    p=workbook(tmp_path,status='Applied');before=tw.sha256_file(p)
    with pytest.raises(ValueError,match='owner-status'):repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    assert tw.sha256_file(p)==before

def test_concurrent_owner_edit_aborts_commit(tmp_path,monkeypatch):
    p=workbook(tmp_path);original=repair.shutil.copy2
    def copy_then_edit(src,dest):
        value=original(src,dest)
        wb=openpyxl.load_workbook(p);wb['Jobs']['F5']='Owner edit';wb.save(p);wb.close()
        return value
    monkeypatch.setattr(repair.shutil,'copy2',copy_then_edit)
    with pytest.raises(ValueError,match='Concurrent workbook drift'):
        repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    wb=openpyxl.load_workbook(p);assert wb['Jobs']['F5'].value=='Owner edit';wb.close()

def test_existing_lock_prevents_repair(tmp_path):
    p=workbook(tmp_path);lock=p.with_suffix('.xlsx.career-mail.lock');lock.touch()
    with pytest.raises(FileExistsError):repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    assert lock.exists()

def test_guarded_append_can_preserve_existing_listing_and_passed_rows(tmp_path):
    p=workbook(tmp_path)
    repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    wb=openpyxl.load_workbook(p);ws=wb['Jobs'];ws['D4']='https://uk.indeed.com/jobs?q=cyber';ws['D4'].hyperlink=ws['D4'].value;ws['F4']='passed';wb.save(p);wb.close()
    rec={'company':'New Employer','title':'Cyber Intern','url':'https://employer.com/jobs/789','date_found':'2026-10-17'}
    result=tw.write_records(tw.load_profiles(),'uk',[rec],tracker_override=str(p),backup_dir=str(tmp_path/'backups'),apply=True,cleanup_existing=False)
    assert result['applied'] and result['counts']['passed_deleted']==result['counts']['listing_deleted']==0
    wb=openpyxl.load_workbook(p);assert wb['Jobs']['F4'].value=='passed' and wb['Jobs'].max_row==5;wb.close()

def requisition_workbook(tmp_path,other_employer='American Express'):
    p=tmp_path/'uk-req.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['UK jobs']);ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status'])
    ws.append(['2026-10-03','American Express','Information Security Engineer - 2027 (UK - Burgess Hill)','https://careers.americanexpress.com/en/sites/CX_1/job/26013761/','2026-10-23','Applied'])
    ws.append(['2026-10-17',other_employer,'Information Security Engineer - 2027','https://egug.fa.us2.oraclecloud.com/hcmUI/CandidateExperience/en/sites/CX_1/job/26013761',None,'Not Applied'])
    from openpyxl.worksheet.datavalidation import DataValidation
    spec=tw.region_config(tw.load_profiles(),'uk')['validations']['F']
    validation=DataValidation(type='list',formula1=spec['formula1']);ws.add_data_validation(validation);validation.add('F3:F500')
    wb.save(p);wb.close();return p

def test_requisition_match_removes_a_relisted_applied_role(tmp_path):
    """The reported regression: one Amex requisition tracked through the employer careers
    URL (Applied) and the Oracle URL (Not Applied) must collapse to a single row."""
    p=requisition_workbook(tmp_path);before=tw.sha256_file(p)
    result=repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    assert result['applied'] and result['verification']['ok']
    archived=result['duplicate_rows_archived'][0]
    assert archived['match_basis']=='same ATS requisition' and archived['retained_original_rows']==[3]
    assert tw.sha256_file(result['backup'])==before
    wb=openpyxl.load_workbook(p);ws=wb['Jobs']
    assert ws.max_row==3 and [ws.cell(3,c).value for c in (2,6)]==['American Express','Applied']
    wb.close()

def test_same_number_at_another_employer_is_not_a_duplicate(tmp_path):
    """A bare numeric requisition ID is not unique across employers; removing on it alone
    would delete a real vacancy."""
    p=requisition_workbook(tmp_path,other_employer='Some Other Bank');before=tw.sha256_file(p)
    with pytest.raises(ValueError,match='No exact compatible retained duplicate'):
        repair.repair(p,tw.load_profiles(),tmp_path/'backups',[4],apply=True)
    assert tw.sha256_file(p)==before
