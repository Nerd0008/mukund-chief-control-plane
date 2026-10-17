import datetime as dt
import pathlib,sys
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
import pytest
from application_deadlines import extract_application_deadline as extract,deadline_priority,deadline_label

@pytest.mark.parametrize('text,expected',[
    ('Applications close on 15 November 2026.','2026-11-15'),
    ('Closing date: 31 Dec 2026','2026-12-31'),
    ('Apply by 1st January 2027','2027-01-01'),
    ('Deadline: 30/06/2026','2026-06-30'),
    ('Application deadline: 2027-01-15','2027-01-15'),
    ('Apply no later than October 30, 2026.','2026-10-30'),
    ('Applications must be received by 30 October 2026','2026-10-30'),
    ('Closing date: 23:59 BST on 30 October 2026','2026-10-30'),
])
def test_exact_dates_are_from_application_wording(text,expected):
    info=extract(text,'https://employer.example/job/123')
    assert info['type']=='EXACT' and info['date']==expected
    assert info['evidence'] in text and info['source']=='JD'
    assert info['source_url']=='https://employer.example/job/123'

@pytest.mark.parametrize('text',[
    'Applications are reviewed on a rolling basis.',
    'We review applications as they are received.',
    'Applications are considered as received.',
    'Rolling applications are welcome.',
    'The vacancy remains open until filled.',
    'Applications may close early.',
    'We reserve the right to close this vacancy earlier.',
])
def test_explicit_rolling(text):
    info=extract(text)
    assert info['type']=='ROLLING' and info['date'] is None and info['rolling']
    assert deadline_label(info)=='ROLLING'

@pytest.mark.parametrize('text',[
    '', 'Start date: 15 November 2026.', 'We are rolling out a new security product.',
    'Applications are not considered on a rolling basis.',
    'Applications close 15 November.', 'Deadline: 03/04/2027',
    'Deadline: 31 February 2027.', 'Interview deadline: 30 November 2026.',
    'Assessment deadline: 30 November 2026.',
    'Closing date: 15 November 2026. Application deadline: 20 November 2026.',
])
def test_unknown_is_not_fabricated_or_treated_as_rolling(text):
    info=extract(text)
    assert info['type']=='UNKNOWN' and info['date'] is None
    assert deadline_label(info)=='UNKNOWN'

def test_exact_backstop_preserved_when_rolling():
    info=extract('Applications reviewed on a rolling basis. Closing date: 30 November 2026.')
    assert info['type']=='EXACT' and info['date']=='2026-11-30' and info['rolling']

def test_priority_urgent_then_rolling_then_future_then_unknown_then_expired():
    today=dt.date(2026,10,20)
    records=[{'name':'unknown','deadline_info':extract('')},
             {'name':'later','deadline':'2026-11-30'},
             {'name':'rolling','deadline_info':extract('Applications reviewed on a rolling basis.')},
             {'name':'urgent','deadline':'2026-10-22'},
             {'name':'expired','deadline':'2026-10-19'}]
    assert [r['name'] for r in sorted(records,key=lambda r:deadline_priority(r,today))]==['urgent','rolling','later','unknown','expired']

def test_failed_fetch_has_unknown_deadline(monkeypatch):
    import jd_fetcher
    monkeypatch.setattr(jd_fetcher.urllib.request,'urlopen',lambda *a,**k:(_ for _ in ()).throw(OSError('offline')))
    result=jd_fetcher.fetch_jd('https://employer.example/job')
    assert result['deadline_info']['type']=='UNKNOWN' and result['deadline'] is None

def test_actual_fetch_extracts_full_jd_not_result_title(monkeypatch):
    import jd_fetcher
    class Response:
        status=200
        class Headers:
            def get_content_charset(self):return 'utf-8'
        headers=Headers()
        def read(self,*a):return b'<main><h1>Cyber Intern</h1><p>Applications are reviewed on a rolling basis.</p></main>'
        def __enter__(self):return self
        def __exit__(self,*a):pass
    monkeypatch.setattr(jd_fetcher.urllib.request,'urlopen',lambda *a,**k:Response())
    result=jd_fetcher.fetch_jd('https://employer.example/job')
    assert result['deadline_info']['type']=='ROLLING'

def test_digest_reads_deadline_column_and_prioritises(tmp_path,monkeypatch):
    import openpyxl,daily_jobs_links as digest
    path=tmp_path/'jobs.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status'])
    for company,date in [('Unknown',None),('Later','2099-12-31'),('Rolling','ROLLING')]:
        ws.append([dt.date.today(),company,'Cyber Intern','https://employer.com/'+company,date,'Not Applied'])
    wb.save(path);wb.close()
    monkeypatch.setitem(digest.TRACKERS,'uk',path)
    monkeypatch.setitem(digest.TRACKER_SCHEMA,'uk',{'header_row':1,'sheet':'Jobs','id':None,'date':'Date Found','company':'Company','title':'Role Title','url':'Apply Link','status':'Status','deadline':'Application Deadline'})
    rows,error=digest.tracker_rows('uk')
    assert not error and [r['company'] for r in rows]==['Rolling','Later','Unknown']

def test_ambiguous_dates_do_not_get_rolling_priority():
    info=extract('Applications reviewed on a rolling basis. Closing date: 15 November 2026. Deadline: 20 November 2026.')
    assert info['type']=='UNKNOWN'
    assert deadline_priority({'deadline_info':info})[0]==3

def test_direct_search_persists_and_orders_jd_deadlines_in_temporary_workbook(tmp_path,monkeypatch,capsys):
    import json,openpyxl,codex_entry_level_search as search
    path=tmp_path/'tracker.xlsx';wb=openpyxl.Workbook();ws=wb.active
    ws.append(['UK jobs']);ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status'])
    wb.save(path);wb.close()
    monkeypatch.setitem(search.TRACKERS,'uk',path)
    monkeypatch.setattr(search.osp,'load_profile',lambda:{})
    monkeypatch.setattr(search,'_profile_terms',lambda:(search.ENTRY_MARKERS,(),None))
    rows=[{'company':name,'job_title':'Cyber Intern','url':'https://employer.com/'+name} for name in ('Unknown','Rolling','Urgent')]
    monkeypatch.setattr(search,'run_codex',lambda *a,**k:json.dumps(rows))
    tomorrow=(dt.date.today()+dt.timedelta(days=1)).isoformat()
    def fetch(url):
        text=('Applications reviewed on a rolling basis.' if url.endswith('Rolling') else 'Application deadline: '+tomorrow if url.endswith('Urgent') else 'No closing date provided.')
        info=extract(text,url)
        return {'text':text,'deadline_info':info,'deadline':info['date']}
    monkeypatch.setattr(search,'fetch_jd',fetch)
    monkeypatch.setattr(sys,'argv',['search','--batches','1'])
    assert search.main()==0
    wb=openpyxl.load_workbook(path);ws=wb.active
    assert [ws.cell(r,2).value for r in (3,4,5)]==['Urgent','Rolling','Unknown']
    assert ws.cell(3,5).value.date()==dt.date.fromisoformat(tomorrow)
    assert ws.cell(4,5).value=='ROLLING' and ws.cell(5,5).value=='UNKNOWN'
    assert ws.max_column==6
    assert all(ws.cell(r,4).hyperlink.target==ws.cell(r,4).value for r in (3,4,5))
    wb.close()

def test_regional_writer_record_preserves_jd_deadline_metadata():
    import regional_job_search as regional
    info=extract('Applications are reviewed on a rolling basis.','https://employer.com/job')
    candidate={'company':'Example','title':'Cyber Intern','url':'https://employer.com/job','location':'London, UK','deadline_info':info}
    record=regional.build_record('uk',candidate,regional.load_policy(),'fixture',{})
    assert record['deadline']=='ROLLING' and record['deadline_info']==info
