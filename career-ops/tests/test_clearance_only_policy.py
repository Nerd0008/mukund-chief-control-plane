"""Current owner rule, including both historical early-filter failure paths."""
import json
import sys
import datetime as dt
import pytest
import owner_rejection_policy as owner
import regional_job_search as regional
import pipeline

SCOPE = {'always_allow':['UK'], 'allow':['United Kingdom'], 'block':['India']}


@pytest.mark.parametrize('text', [
    'You must hold security clearance.',
    'This role requires SC clearance and five years UK residency.',
    'Additional requirements: Eligibility and willingness to obtain UK Government Security Clearance.',
    'What we are looking for: The ability to obtain HMG SC clearance.',
    'Applicants must be eligible to obtain SC clearance.',
    'Security clearance at SC level required.',
    'SC clearance is essential.',
    'This role is subject to national security vetting.',
    'You will need to obtain security clearance.',
    'No existing security clearance required but you must obtain SC clearance.',
])
def test_explicit_mandatory_clearance(text):
    assert owner.required_clearance_evidence(text)


@pytest.mark.parametrize('text', [
    'British citizenship required. Must have indefinite leave to remain.',
    'BPSS screening required.',
    'Security clearance is desirable.',
    'No security clearance is required.',
    'You do not need to obtain security clearance.',
    'Security clearance is not required.',
    'DV Trading offers an internship. SC-200 certification preferred.',
    'We help clients manage security clearance processes.',
    'SC clearance preferred, you must have three years of experience.',
    'Similar jobs: Applicants must hold security clearance.',
])
def test_mentions_and_other_barriers_do_not_reject(text):
    assert not owner.required_clearance_evidence(text)


@pytest.mark.parametrize('worker', ['regional','unified'])
@pytest.mark.parametrize('concern', [
    'Must have five years of professional experience.',
    'Must be a current student and return to university after placement.',
    'Must be a British citizen with indefinite right to work.',
    'No longer accepting applications. This vacancy is closed.',
    '',
    'Security clearance is desirable.',
])
def test_non_clearance_concerns_retained_by_both_gates(worker,concern):
    record={'company':'Acme','title':'Senior IT Support Engineer','location':'India',
            'url':'https://employer.example/jobs/123','jd_text':concern,
            'jd_screening_required':True,'vacancy_status':{'state':'CLOSED'}}
    call = pipeline.deterministic_gates if worker=='unified' else regional.evaluate_record
    args=('uk',record,regional.load_policy())
    result=call(*args,SCOPE) if worker=='unified' else call(*args,{},SCOPE)
    assert result['decision']=='accepted'
    assert result['reasons']==[]
    assert result['review_flags']
    assert result['eligibility_verified'] is False


def test_shared_gate_only_rejects_explicit_clearance():
    record={'company':'Acme','title':'Cyber Intern','location':'London, UK',
            'url':'https://employer.example/jobs/123','jd_text':'You must hold security clearance.'}
    result=pipeline.deterministic_gates('uk',record,regional.load_policy(),SCOPE)
    assert result['decision']=='rejected'
    assert all('Explicit required security clearance:' in reason for reason in result['reasons'])
    assert result['clearance_evidence'][0]['context']==record['jd_text'].rstrip('.')


def test_direct_search_retains_title_experience_closed_and_missing_country(tmp_path,monkeypatch,capsys):
    import openpyxl
    import codex_entry_level_search as search
    import job_scan_delta as report
    path=tmp_path/'tracker.xlsx'
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['UK jobs']);ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status'])
    wb.save(path);wb.close()
    monkeypatch.setitem(search.TRACKERS,'uk',path)
    monkeypatch.setattr(search.osp,'load_profile',lambda:{})
    monkeypatch.setattr(search,'_profile_terms',lambda:(search.ENTRY_MARKERS,(),None))
    rows=[{'company':name,'job_title':'Security Engineer (Early Career)','location':'London',
           'url':'https://employer.example/jobs/'+name} for name in ('Experience','Student','Closed','Clearance')]
    texts={'Experience':'Five years of experience required.',
           'Student':'You must be currently enrolled in university and return to study.',
           'Closed':'This vacancy is closed.',
           'Clearance':'You must obtain security clearance.'}
    monkeypatch.setattr(search,'run_codex',lambda *a,**k:json.dumps(rows))
    def fetch(url):
        name=url.rsplit('/',1)[-1]
        return {'text':texts[name],'is_closed':name=='Closed',
                'vacancy_status':{'state':'CLOSED' if name=='Closed' else 'OPEN','evidence':'offline fixture'}}
    monkeypatch.setattr(search,'fetch_jd',fetch)
    monkeypatch.setattr(sys,'argv',['search','--batches','1','--batch-offset','0'])
    assert search.main()==0
    output=capsys.readouterr().out
    reconciliation=report.parse_tagged(output,'SEARCH_RECONCILIATION')
    assert reconciliation['counts']['rejected']==1
    assert reconciliation['counts']['new_candidates']==3
    roles=report.current_run_roles(output)
    assert {row['company'] for row in roles}=={'Experience','Student','Closed'}
    assert all(row['review_flags'] for row in roles)
    assert 'Availability: CLOSED' in report.build_digest(roles,dt.datetime.now())
    wb=openpyxl.load_workbook(path)
    assert wb.active.max_row==5 and wb.active.max_column==6
    assert all('Review flags:' in wb.active.cell(row,4).comment.text for row in (3,4,5))
    wb.close()


def test_writer_integrity_holds_are_review_only_and_duplicates_stay_separate():
    import tracker_writer as writer
    result={'counts':{'appended':0,'duplicates':1,'rejected':1},'outcomes':[
        {'company':'Unknown','title':'Cyber Intern','url':'https://example.com/jobs','decision':'rejected','reason':'employer unknown'},
        {'company':'Acme','title':'Cyber Intern','url':'https://example.com/job/1','decision':'duplicate'}]}
    converted=writer._owner_writer_result(result,'uk',[])
    assert converted['counts']['rejected']==0
    assert converted['counts']['review_only']==1
    assert converted['outcomes'][1]['decision']=='duplicate'
    assert len(converted['review_records'])==1


def test_non_uk_policy_inherited():
    assert not owner.clearance_only('dubai')


def test_unified_keeps_prefilter_and_semantic_failures_for_clearance_screening(monkeypatch):
    import tracker_writer as writer
    records=[{'company':'Acme','title':'Senior Security Engineer','location':'London, UK',
              'url':'https://employer.example/job/123','source':'explicit_records'}]
    fetched=[]
    monkeypatch.setattr(pipeline,'fetch_jd',lambda url,**k: fetched.append(url) or {
        'text':'Five years of commercial experience required.',
        'vacancy_status':{'state':'OPEN','evidence':'offline fixture'}})
    monkeypatch.setattr(writer,'write_records',lambda *a,**k:{'counts':{'appended':1},'outcomes':[
        dict(records[0],decision='appended')],'applied':False})
    doc=pipeline.run_funnel(records,region='uk',mode='high_recall',semantic='deterministic',
        deepseek_model='unused',batch_size=5,codex_budget=0,codex_enabled=False,timeout=10,
        run_id='clearance-only-offline',collection=[{'source':'explicit_records','candidates':records,'coverage':{}}])
    assert fetched==[records[0]['url']]
    assert doc['funnel']['counts']['after_hard_negative_prefilter']==1
    assert doc['funnel']['counts']['deterministic_eligibility_pass']==1
    assert not doc['funnel']['rejections_by_reason']
    # Inspect the gate entries without relying on report labels.
    gated=doc['eligibility']['decisions']
    assert gated[0]['review_flags']
    assert gated[0]['decision']=='accepted'
