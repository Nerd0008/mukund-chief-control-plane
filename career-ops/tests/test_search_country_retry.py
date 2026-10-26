import json,sys
from pathlib import Path
import openpyxl,pytest
import codex_entry_level_search as search

@pytest.mark.parametrize('country_field,country_value',[('country','United Kingdom'),('country_code','GB')])
@pytest.mark.parametrize('evidence_source',['fresh','retained'])
@pytest.mark.parametrize('partial',[False,True])
def test_fresh_country_evidence_survives_pending_retry(tmp_path,monkeypatch,partial,country_field,country_value,evidence_source):
    path=tmp_path/'tracker.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['UK jobs']);ws.append(['Date Found','Company','Role Title','Apply Link','Application Deadline','Status']);wb.save(path);wb.close()
    monkeypatch.setitem(search.TRACKERS,'uk',path)
    monkeypatch.setattr(search.osp,'load_profile',lambda:{})
    monkeypatch.setattr(search,'_profile_terms',lambda:(search.ENTRY_MARKERS,(),None))
    state=tmp_path/'state';state.mkdir();monkeypatch.setattr(search,'SEARCH_STATE_DIR',state)
    url='https://employer.com/jobs/12345678'
    row={'company':'Example','job_title':'Cyber Security Intern','location':'Luton, Bedfordshire',country_field:country_value,'url':url}
    pending = dict(row, country=None, country_code=None) if evidence_source=='fresh' else dict(row)
    (state/'uk-blocked-candidates.json').write_text(json.dumps({'records':[pending]}))
    if evidence_source=='retained':row=dict(row,country=None,country_code=None)
    monkeypatch.setattr(search,'run_codex',lambda *a,**k: search.SearchResponse(json.dumps([row]),{'coverage_complete':not partial}))
    monkeypatch.setattr(search,'fetch_jd',lambda url:{'text':'Cyber security internship. No prior experience required.','vacancy_status':{'state':'UNKNOWN','evidence':'fixture automated access unavailable'}})
    import posting_resolution
    class Resolver:
        def __init__(self,*a,**k):pass
        def __call__(self,*a,**k):return None
    monkeypatch.setattr(posting_resolution,'PostingResolver',Resolver)
    monkeypatch.setattr(sys,'argv',['search','--batches','1','--batch-offset','0'])
    result=search.main()
    if partial:
        assert result!=0
        retained=json.loads((state/'uk-blocked-candidates.json').read_text())['records']
        assert retained[0][country_field]==country_value
    else:
        assert result==0
        wb=openpyxl.load_workbook(path);assert wb.active.cell(3,2).value=='Example';wb.close()


@pytest.mark.parametrize('difference',['company','title','url','country'])
def test_retry_country_never_crosses_identity_or_overrides_fresh_country(difference):
    old={'company':'Example','title':'Cyber Intern','url':'https://employer.com/jobs/123','country':'United Kingdom'}
    fresh=dict(old);fresh.pop('country')
    if difference=='country':fresh['country']='United States'
    elif difference=='title':fresh['title']='Different Internship'
    else:fresh[difference]='Different'
    result=search.retry_country_records([old],[fresh])[0]
    assert result.get('country')==('United States' if difference=='country' else None)
