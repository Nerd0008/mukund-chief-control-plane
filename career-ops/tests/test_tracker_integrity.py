import json,pathlib,sys,copy
import openpyxl,pytest
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
from tracker_integrity import assert_valid,bind_links
from career_mail_tracker import WorkbookTracker
from hermes_tracker_guard import pre_tool

@pytest.fixture
def workbook(tmp_path):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    headers=['Date Found','Company','Role Title','Apply Link','Application Deadline','Status']
    for c,h in enumerate(headers,1):ws.cell(2,c,h)
    for r in range(3,7):
        ws.cell(r,2,'Company'+str(r));ws.cell(r,3,'Role');ws.cell(r,4,'https://example.com/jobs/'+str(r));ws.cell(r,6,'Not Applied')
    cfg={'tracker':str(tmp_path/'tracker.xlsx'),'sheet':'Jobs','header_row':2,'first_data_row':3,'column_ceiling':6,
         'hyperlinks_required':True,'required_headers':{chr(65+i):h for i,h in enumerate(headers)},'dedupe':{'url_column':'D','company_column':'B','title_column':'C'},
         'field_map':{'company':'B','title':'C','url':'D'},'status_columns':{'application_status':'F'},'owner_columns':['E','F'],
         'mail_metadata_sidecar':'private-state.json'}
    bind_links(ws,cfg);wb.save(cfg['tracker']);return wb,ws,cfg

def test_delete_moves_link_refs_and_is_rebound_before_serialization(workbook,tmp_path):
    wb,ws,cfg=workbook;ws.delete_rows(3)
    with pytest.raises(ValueError,match='hyperlink'):assert_valid(ws,cfg)
    bind_links(ws,cfg);out=tmp_path/'fixed.xlsx';wb.save(out);loaded=openpyxl.load_workbook(out)
    assert_valid(loaded['Jobs'],cfg)
    assert loaded['Jobs']['D3'].hyperlink.target=='https://example.com/jobs/4'

def test_shifted_target_and_extra_column_fail_closed(workbook):
    wb,ws,cfg=workbook;ws['D3'].hyperlink='https://example.com/jobs/4'
    with pytest.raises(ValueError,match='hyperlink'):assert_valid(ws,cfg)
    bind_links(ws,cfg);ws['G2']='Career Ops Email Application ID'
    with pytest.raises(ValueError,match='column ceiling'):assert_valid(ws,cfg)

def test_sort_requires_rebinding_own_urls(workbook,tmp_path):
    wb,ws,cfg=workbook;ws['D3'].value,ws['D4'].value=ws['D4'].value,ws['D3'].value
    with pytest.raises(ValueError):assert_valid(ws,cfg)
    bind_links(ws,cfg);wb.save(tmp_path/'sort.xlsx');assert_valid(openpyxl.load_workbook(tmp_path/'sort.xlsx')['Jobs'],cfg)

def test_sidecar_preserves_workbook_and_deleted_rows_not_recreated(workbook,tmp_path,monkeypatch):
    wb,ws,cfg=workbook;monkeypatch.setenv('LOCALAPPDATA',str(tmp_path))
    adapter=WorkbookTracker({'regions':{'uk':cfg}},tmp_path/'backups');before=pathlib.Path(cfg['tracker']).read_bytes()
    rows,_=adapter.read();record=rows[0];record['seen_message_ids']=['m1'];adapter.upsert(record)
    assert pathlib.Path(cfg['tracker']).read_bytes()==before
    assert adapter.read()[0][0]['seen_message_ids']==['m1']
    ws.delete_rows(3);bind_links(ws,cfg);wb.save(cfg['tracker'])
    assert all(r['company']!=record['company'] for r in adapter.read()[0])
    with pytest.raises(ValueError,match='no row restored'):adapter.upsert(record)

def test_ad_hoc_mutations_blocked_but_unrelated_reads_work():
    assert pre_tool('execute_code',{'code':"load_workbook('uk-cyber-job-tracker.xlsx').save('uk-cyber-job-tracker.xlsx')"})['action']=='block'
    assert pre_tool('terminal',{'command':'python career_ops_cli.py write --region uk --apply'}) is None
    assert pre_tool('execute_code',{'code':"load_workbook('uk-cyber-job-tracker.xlsx').sheetnames"}) is None
    assert pre_tool('terminal',{'command':'python unrelated_task.py'}) is None


def test_real_writer_append_and_passed_deletion_have_correct_links(workbook,tmp_path):
    import tracker_writer as writer
    wb,ws,cfg=workbook;ws['F3']='passed';wb.save(cfg['tracker']);wb.close()
    profile={'regions':{'uk':cfg},'canonical_dir':str(tmp_path),'backup_dir':str(tmp_path/'backups'),'career_ops_root':str(tmp_path)}
    record={'company':'New Employer','title':'Cyber Security Intern','url':'https://example.com/jobs/new','location':'London'}
    result=writer.write_records(profile,'uk',[record],apply=True)
    assert result['applied'],result
    loaded=openpyxl.load_workbook(cfg['tracker']);assert_valid(loaded['Jobs'],cfg)
    assert loaded['Jobs']['D3'].value=='https://example.com/jobs/4'
    assert loaded['Jobs']['D6'].hyperlink.target=='https://example.com/jobs/new'
    assert result['backup']

def test_writer_lock_excludes_concurrent_mailbox_write(workbook,tmp_path):
    import tracker_writer as writer
    wb,ws,cfg=workbook;lock=pathlib.Path(cfg['tracker']+'.career-mail.lock');lock.write_text('busy')
    with pytest.raises(FileExistsError):writer.write_records({'regions':{'uk':cfg}},'uk',[],apply=True)

def test_fixed_schema_and_mail_identity_configuration_conflict_rejected(workbook):
    from career_tracker_layout import resolve_profiles
    wb,ws,cfg=workbook;cfg['mail_identity']={'column':'G'}
    with pytest.raises(ValueError,match='cannot reserve'):resolve_profiles({'regions':{'uk':cfg}})


def test_post_commit_failure_rolls_back_only_our_own_write(workbook,tmp_path,monkeypatch):
    import tracker_writer as writer
    wb,ws,cfg=workbook;wb.close();before=pathlib.Path(cfg['tracker']).read_bytes()
    profile={'regions':{'uk':cfg},'canonical_dir':str(tmp_path),'backup_dir':str(tmp_path/'backups'),'career_ops_root':str(tmp_path)}
    actual=writer.verify_workbook;calls=[]
    def verify(*args,**kwargs):
        result=actual(*args,**kwargs);calls.append(args[0])
        if len(calls)==2:result['ok']=False
        return result
    monkeypatch.setattr(writer,'verify_workbook',verify)
    result=writer.write_records(profile,'uk',[{'company':'New','title':'Cyber Security Intern','url':'https://example.com/new'}],apply=True)
    assert not result['applied'] and result['rolled_back']
    assert pathlib.Path(cfg['tracker']).read_bytes()==before

def test_concurrent_owner_edit_never_overwritten(workbook,tmp_path,monkeypatch):
    import tracker_writer as writer
    wb,ws,cfg=workbook;wb.close()
    profile={'regions':{'uk':cfg},'canonical_dir':str(tmp_path),'backup_dir':str(tmp_path/'backups'),'career_ops_root':str(tmp_path)}
    actual=writer.verify_workbook
    def verify(*args,**kwargs):
        result=actual(*args,**kwargs)
        changed=openpyxl.load_workbook(cfg['tracker']);changed['Jobs']['F3']='Applied';changed.save(cfg['tracker']);changed.close()
        return result
    monkeypatch.setattr(writer,'verify_workbook',verify)
    result=writer.write_records(profile,'uk',[{'company':'New','title':'Cyber Security Intern','url':'https://example.com/new'}],apply=True)
    assert not result['applied'] and 'changed during write' in result['error']
    assert openpyxl.load_workbook(cfg['tracker'])['Jobs']['F3'].value=='Applied'


def test_rollover_uses_shared_writer_lock(workbook):
    import tracker_rollover as rollover
    wb,ws,cfg=workbook;pathlib.Path(cfg['tracker']+'.career-mail.lock').write_text('busy')
    with pytest.raises(FileExistsError):rollover.rollover({'regions':{'uk':cfg}},'uk','2026-10',apply=True)


def test_hyperlink_created_phantom_row_cannot_be_committed(workbook):
    wb,ws,cfg=workbook;ws['D7'].hyperlink='https://example.com/orphan'
    with pytest.raises(ValueError,match='no application identity'):assert_valid(ws,cfg)
