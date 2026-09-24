#!/usr/bin/env python3
"""Deterministic tests for the Chief <-> Career Ops tracker writer.

Every test that writes operates on a *copy* of a canonical workbook; the
canonical files under C:\\Users\\mukun\\Downloads\\codex are only ever read here.
Run:  python -m pytest career-ops/tests/test_tracker_writer.py -v
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

CAREER_OPS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(CAREER_OPS))

import tracker_writer as tw  # noqa: E402

PROFILES = tw.load_profiles()
REGIONS = list(PROFILES["regions"])


def h(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def workspace(tmp_path_factory):
    d = tmp_path_factory.mktemp("career-ops-tests")
    (d / "backups").mkdir()
    return d


def copy_tracker(region: str, dest_dir: Path) -> Path:
    src = Path(PROFILES["regions"][region]["tracker"])
    dest = dest_dir / f"{region}-{src.name}"
    shutil.copy2(src, dest)
    return dest


def make_record(**kw) -> dict:
    base = {
        "date_found": "2026-09-23",
        "company": "Example Analytics Ltd",
        "title": "Cyber Security Intern",
        "location": "London, United Kingdom",
        "url": "https://job-boards.greenhouse.io/exampleanalytics/jobs/999000111",
        "salary": "£28,000",
        "posted_date": "2026-09-20",
        "fit_score": 4.4,
        "live_status": "Live",
        "clearance_check": "Pass — no clearance stated",
        "work_authorisation_risk": "Low; no restriction stated",
        "key_gap": "Internship contract date not yet confirmed",
        "recommendation": "Retained for owner review; official page verified 2026-09-23.",
        "discovery": "BOUNDED EVIDENCE RUN 2026-09-23",
        "source": "Bounded evidence run 2026-09-23",
        "notes": "Evidence-run record; no application created.",
    }
    base.update(kw)
    return base


# --------------------------------------------------------------------------- #
# structural / read-only tests
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("region", REGIONS)
def test_canonical_workbook_verifies(region):
    """Every canonical workbook passes structural + integrity verification."""
    cfg = dict(tw.region_config(PROFILES, region))
    cfg["region"] = region
    res = tw.verify_workbook(Path(cfg["tracker"]), cfg)
    assert res["ok"], res["problems"]


@pytest.mark.parametrize("region", REGIONS)
def test_profile_structure_matches_workbook(region):
    """Profile column letters/table/header facts still match the real workbook."""
    cfg = dict(tw.region_config(PROFILES, region))
    cfg["region"] = region
    tr = tw.Tracker(Path(cfg["tracker"]), cfg)
    try:
        assert tr.ws.title == cfg["sheet"]
        assert cfg["table"] in tr.ws.tables
        # header row must contain the dedupe url column header
        hdr = tr.headers()
        url_col = tw.column_index_from_string(cfg["dedupe"]["url_column"])
        assert hdr[url_col - 1] is not None
        # owner columns must be inside the used range
        for c in cfg["owner_columns"]:
            assert tw.column_index_from_string(c) <= tr.ws.max_column
        # ids are readable and non-empty
        assert tr.ids(), "no ids found"
    finally:
        tr.wb.close()


def test_no_artifact_tool_dependency_remains():
    """The replaced writer must not import the unsupported artifact tool.

    Only executable code counts: the module docstring is allowed to *name* the
    dependency it replaces, so imports/attribute use are checked, not prose.
    """
    src = (CAREER_OPS / "tracker_writer.py").read_text(encoding="utf-8")
    code_lines = [ln for ln in src.splitlines()
                  if not ln.strip().startswith("#")]
    body = "\n".join(code_lines)
    # strip the module docstring
    import ast
    tree = ast.parse(src)
    doc = ast.get_docstring(tree)
    if doc:
        body = body.replace(doc, "")
    assert "@oai/artifact-tool" not in body
    assert "SpreadsheetFile" not in body
    assert "FileBlob" not in body
    import importlib
    mod = importlib.import_module("tracker_writer")
    for banned in ("SpreadsheetFile", "FileBlob"):
        assert not hasattr(mod, banned)


# --------------------------------------------------------------------------- #
# dedupe behaviour
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("region", REGIONS)
def test_existing_canonical_row_is_a_duplicate(region, workspace):
    cfg = dict(tw.region_config(PROFILES, region))
    cfg["region"] = region
    copy = copy_tracker(region, workspace)
    tr = tw.Tracker(copy, cfg)
    existing = tr.data_rows()[0]
    url_idx = tw.column_index_from_string(cfg["dedupe"]["url_column"]) - 1
    comp_idx = tw.column_index_from_string(cfg["dedupe"]["company_column"]) - 1
    title_idx = tw.column_index_from_string(cfg["dedupe"]["title_column"]) - 1
    tr.wb.close()
    rec = make_record(url=tw.extract_url(existing[url_idx]),
                      company=existing[comp_idx], title=existing[title_idx])
    result = tw.write_records(PROFILES, region, [rec], apply=False, tracker_override=str(copy))
    assert result["counts"]["appended"] == 0
    assert result["counts"]["duplicates"] >= 1


def test_url_normalization_handles_tracking_params_and_hyperlinks():
    a = tw.normalize_url("https://Jobs.Lever.co/acme/abc123?utm_source=x&gh_src=y")
    b = tw.normalize_url("https://jobs.lever.co/acme/abc123")
    assert a == b
    c = tw.normalize_url('=HYPERLINK("https://jobs.lever.co/acme/abc123","Apply")')
    assert c == b


def test_url_normalization_drops_any_utm_prefixed_parameter():
    """A re-share that adds an arbitrary utm_* marker is the same posting."""
    canonical = tw.normalize_url("https://jobs.lever.co/acme/abc123")
    assert tw.normalize_url(
        "https://jobs.lever.co/acme/abc123?utm_campaign=alert&utm_content=7788"
    ) == canonical
    # a non-tracking query parameter still distinguishes two postings
    assert tw.normalize_url("https://jobs.lever.co/acme/abc123?team=soc") != canonical


def test_cross_month_duplicate_detected_from_archive(workspace, tmp_path):
    """A URL recorded only in a previous month's workbook is not re-added."""
    cfg = dict(tw.region_config(PROFILES, "japan"))
    cfg["region"] = "japan"
    copy = copy_tracker("japan", tmp_path)

    # Build a synthetic previous-month archive holding one known posting URL.
    archive_dir = tmp_path / "archive"
    archive_dir.mkdir()
    archived_url = "https://hrmos.co/pages/example/jobs/archived-2026-08-01"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg["sheet"]
    ws["E2"] = archived_url
    wb.save(archive_dir / "Japan_Cybersecurity_Job_Tracker.2026-08.xlsx")
    wb.close()

    rec = make_record(url=archived_url, company="Archive Example", title="Archived Role")
    result = tw.write_records(PROFILES, "japan", [rec], apply=False,
                              tracker_override=str(copy),
                              extra_archive_dirs=[str(archive_dir)])
    assert result["counts"]["appended"] == 0
    decisions = [o["decision"] for o in result["outcomes"]]
    assert "duplicate-cross-month" in decisions, decisions
    assert "Japan_Cybersecurity_Job_Tracker.2026-08.xlsx" in result["cross_month_index"]["sources"]


def test_ledger_only_url_is_deduplicated(tmp_path):
    """URLs present only in the region scan ledger are treated as already seen."""
    cfg = dict(tw.region_config(PROFILES, "uk"))
    cfg["region"] = "uk"
    copy = copy_tracker("uk", tmp_path)
    ledger = Path(PROFILES["career_ops_root"]) / cfg["dedupe_ledger"]["path"]
    ledger_urls = tw.urls_from_scan_history(ledger, cfg["dedupe_ledger"]["url_col"])
    tr = tw.Tracker(copy, cfg)
    in_workbook = {tw.normalize_url(r[tw.column_index_from_string(cfg["dedupe"]["url_column"]) - 1])
                   for r in tr.data_rows()}
    tr.wb.close()
    ledger_only = sorted(ledger_urls - in_workbook)
    assert ledger_only, "expected the UK scan ledger to hold URLs beyond the workbook rows"
    rec = make_record(url=ledger_only[0], company="Ledger Example", title="Ledger Example Role")
    result = tw.write_records(PROFILES, "uk", [rec], apply=False, tracker_override=str(copy))
    assert result["counts"]["appended"] == 0
    assert any(o["decision"] == "duplicate-cross-month" for o in result["outcomes"])


def test_manifest_cannot_set_application_state(tmp_path):
    """A malicious/careless manifest must not be able to mark a role Applied."""
    cfg = dict(tw.region_config(PROFILES, "uk"))
    cfg["region"] = "uk"
    copy = copy_tracker("uk", tmp_path)
    rec = make_record(application_status="Applied", status="Applied", selected="Yes",
                      url="https://job-boards.greenhouse.io/exampleanalytics/jobs/999000112")
    res = tw.write_records(PROFILES, "uk", [rec], apply=True, tracker_override=str(copy),
                          backup_dir=str(tmp_path))
    assert res["applied"], res.get("error")
    wb = openpyxl.load_workbook(copy)
    ws = wb[cfg["sheet"]]
    last = max(r for r in range(cfg["first_data_row"], ws.max_row + 1)
               if ws.cell(row=r, column=15).value not in (None, ""))
    assert ws[f"J{last}"].value == "To Review", "manifest smuggled an application status"
    wb.close()


# --------------------------------------------------------------------------- #
# write path: dry-run safety and reversible acceptance
# --------------------------------------------------------------------------- #

def test_dry_run_changes_nothing(tmp_path):
    copy = copy_tracker("uk", tmp_path)
    before = h(copy)
    res = tw.write_records(PROFILES, "uk", [make_record()], apply=False,
                           tracker_override=str(copy))
    assert res["applied"] is False
    assert res["counts"]["appended"] == 1
    assert h(copy) == before
    assert not list(tmp_path.glob("*.pre-*.xlsx"))


@pytest.mark.parametrize("region", REGIONS)
def test_apply_is_reversible_and_verified(region, workspace, tmp_path):
    """Full acceptance: append -> verify -> backup exists -> rollback restores."""
    cfg = dict(tw.region_config(PROFILES, region))
    cfg["region"] = region
    copy = copy_tracker(region, tmp_path)
    backups = tmp_path / f"backups-{region}"
    before_hash = h(copy)

    tr = tw.Tracker(copy, cfg)
    pre_rows = tr.last_data_row() - cfg["first_data_row"] + 1
    pre_owned = tr.owned_snapshot()
    pre_dv = {col: len(v) for col, v in tw._dv_map(tr.ws).items()}
    tr.wb.close()

    records = [
        make_record(company=f"Evidence Run {i} Ltd",
                    title=f"Security Analyst Intern {i}",
                    url=f"https://job-boards.greenhouse.io/evidencerun{i}/jobs/5550000{i}")
        for i in (1, 2)
    ]
    res = tw.write_records(PROFILES, region, records, apply=True,
                           tracker_override=str(copy), backup_dir=str(backups))

    assert res["applied"], res.get("error")
    assert res["counts"]["appended"] == 2
    assert Path(res["backup"]).exists()
    assert res["verification"]["ok"], res["verification"]["problems"]
    assert res["post_data_rows"] == pre_rows + 2
    assert res["verification"]["owner_columns_unchanged"] is True
    assert res["verification"]["duplicate_urls_in_workbook"] == []

    # ids are unique and follow the region convention
    tr2 = tw.Tracker(copy, cfg)
    ids = tr2.ids()
    assert len(ids) == len(set(ids))
    for rid in res["planned_rows"]:
        assert rid in ids
    # fit-tier formula present on every appended row where the region uses one
    if cfg.get("formula_map"):
        col = next(iter(cfg["formula_map"].values()))["column"]
        last = tr2.last_data_row()
        for r in range(last - 1, last + 1):
            assert str(tr2.ws[f"{col}{r}"].value).startswith("=")
    # data validations preserved
    assert {c: len(v) for c, v in tw._dv_map(tr2.ws).items()} == pre_dv
    tr2.wb.close()

    # --- rollback test: restoring the backup returns the workbook byte-identically
    shutil.copy2(res["backup"], copy)
    assert h(copy) == before_hash
    res3 = tw.verify_workbook(copy, cfg, expect_data_rows=pre_rows)
    assert res3["ok"], res3["problems"]


def test_second_identical_run_is_idempotent(tmp_path):
    copy = copy_tracker("uk", tmp_path)
    recs = [make_record(url="https://job-boards.greenhouse.io/exampleanalytics/jobs/999000777")]
    first = tw.write_records(PROFILES, "uk", recs, apply=True, tracker_override=str(copy),
                             backup_dir=str(tmp_path))
    assert first["applied"] and first["counts"]["appended"] == 1
    after_first = h(copy)
    second = tw.write_records(PROFILES, "uk", recs, apply=True, tracker_override=str(copy),
                              backup_dir=str(tmp_path))
    # nothing to write: no append, reported as a duplicate, workbook untouched
    assert second["applied"] is False
    assert second["counts"]["appended"] == 0
    assert second["counts"]["duplicates"] == 1
    assert second["reason"] == "no changes required: nothing to append or refresh"
    assert h(copy) == after_first


def test_summary_is_chief_readable():
    from career_ops_cli import cmd_summary
    import argparse
    args = argparse.Namespace(region=None, profiles=None)
    assert cmd_summary(args) == 0


@pytest.mark.parametrize("region", REGIONS)
def test_summary_reports_real_ids_not_dates(region):
    """Summary actionable ids must be workbook ids for every region's layout."""
    import argparse
    import io
    from contextlib import redirect_stdout
    from career_ops_cli import cmd_summary

    buf = io.StringIO()
    with redirect_stdout(buf):
        assert cmd_summary(argparse.Namespace(region=region, profiles=None)) == 0
    data = json.loads(buf.getvalue())["regions"][region]
    cfg = dict(tw.region_config(PROFILES, region))
    cfg["region"] = region
    tr = tw.Tracker(Path(cfg["tracker"]), cfg)
    real_ids = set(tr.ids())
    tr.wb.close()
    for aid in data["actionable_ids"]:
        assert aid in real_ids, f"{aid!r} is not a real id in region {region}"
