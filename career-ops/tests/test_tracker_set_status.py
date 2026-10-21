"""Status changes come from employer evidence, never from inference.

The failure this pins: a tracker row said "Applied" long after the employer had
sent a rejection. The canonical writer deliberately refuses to touch the Status
column (it is the owner's field), so the only path is an explicit, backed-up,
verified edit. These tests pin the safety properties of that path.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import tracker_set_status as tss  # noqa: E402

SHEET = 'Jobs'
HEADER_ROW = 2


def make_workbook(tmp_path: Path, rows: list[tuple]) -> Path:
    """rows: (company, title, url, status) starting at row 3."""
    p = tmp_path / 'uk.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.cell(row=1, column=1, value='UK Cyber Job Tracker')
    for c, h in enumerate(['Date Found', 'Company', 'Role Title', 'Apply Link',
                           'Application Deadline', 'Status'], start=1):
        ws.cell(row=HEADER_ROW, column=c, value=h)
    for i, (company, title, url, status) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value='2026-10-21')
        ws.cell(row=i, column=2, value=company)
        ws.cell(row=i, column=3, value=title)
        cell = ws.cell(row=i, column=4, value=url)
        cell.hyperlink = url
        ws.cell(row=i, column=6, value=status)
    wb.save(p)
    return p


class _StubTracker:
    """Only the surface tracker_set_status uses: a data-row count."""

    def __init__(self, path, cfg):
        self.path = Path(path)

    def data_rows(self):
        wb = openpyxl.load_workbook(self.path, read_only=True)
        rows = [r for r in wb[SHEET].iter_rows(min_row=3, values_only=True) if any(r)]
        wb.close()
        return rows


@pytest.fixture(autouse=True)
def _stub_tracker_writer(monkeypatch):
    monkeypatch.setattr(tss.tw, 'Tracker', _StubTracker)
    monkeypatch.setattr(tss.tw, 'region_config', lambda profiles, region: {'sheet': SHEET})


@pytest.fixture
def profiles():
    return {}


def test_dry_run_reports_the_change_without_writing(tmp_path, profiles):
    p = make_workbook(tmp_path, [('Acme', 'Graduate Analyst', 'https://x/1', 'Applied')])
    out = tss.apply_changes(p, [{'row': 3, 'status': 'Rejected'}],
                            profiles=profiles, backup_dir=tmp_path / 'b', dry_run=True)
    assert out['written'] is False
    assert out['mode'] == 'dry-run'
    assert out['applied'][0]['from'] == 'Applied'
    assert out['applied'][0]['to'] == 'Rejected'
    wb = openpyxl.load_workbook(p)
    assert wb[SHEET].cell(row=3, column=6).value == 'Applied'


def test_wrong_current_status_is_refused_and_nothing_is_written(tmp_path, profiles):
    """A guard against marking the wrong row: if it is not where we expect, stop."""
    p = make_workbook(tmp_path, [('Acme', 'Graduate Analyst', 'https://x/1', 'Not Applied')])
    out = tss.apply_changes(
        p, [{'row': 3, 'status': 'Rejected', 'expect_status': 'Applied'}],
        profiles=profiles, backup_dir=tmp_path / 'b', dry_run=False)
    assert out['written'] is False
    assert 'refused' in out and out['refused']
    assert out['aborted']
    wb = openpyxl.load_workbook(p)
    assert wb[SHEET].cell(row=3, column=6).value == 'Not Applied'


def test_a_row_with_no_company_is_refused(tmp_path, profiles):
    p = make_workbook(tmp_path, [('', '', 'https://x/1', 'Applied')])
    out = tss.apply_changes(p, [{'row': 3, 'status': 'Rejected'}],
                            profiles=profiles, backup_dir=tmp_path / 'b', dry_run=False)
    assert out['written'] is False
    assert out['refused'][0]['reason'] == 'no company in column B'


def test_successful_write_sets_status_and_takes_a_backup(tmp_path, profiles, monkeypatch):
    p = make_workbook(tmp_path, [('Acme', 'Graduate Analyst', 'https://x/1', 'Applied')])
    monkeypatch.setattr(tss.tw, 'verify_workbook',
                        lambda *a, **k: {'ok': True, 'problems': []})
    out = tss.apply_changes(p, [{'row': 3, 'status': 'Rejected', 'expect_status': 'Applied'}],
                            profiles=profiles, backup_dir=tmp_path / 'b', dry_run=False)
    assert out['written'] is True
    assert out['verification_ok'] is True
    assert Path(out['backup']).exists()
    wb = openpyxl.load_workbook(p)
    ws = wb[SHEET]
    assert ws.cell(row=3, column=6).value == 'Rejected'
    assert ws.cell(row=3, column=4).hyperlink is not None, 'hyperlink must survive the edit'


def test_a_failed_verification_restores_the_original(tmp_path, profiles, monkeypatch):
    """Fail closed: a workbook that cannot be proven intact is rolled back."""
    p = make_workbook(tmp_path, [('Acme', 'Graduate Analyst', 'https://x/1', 'Applied')])
    monkeypatch.setattr(tss.tw, 'verify_workbook',
                        lambda *a, **k: {'ok': False, 'problems': ['link mismatch']})
    out = tss.apply_changes(p, [{'row': 3, 'status': 'Rejected'}],
                            profiles=profiles, backup_dir=tmp_path / 'b', dry_run=False)
    assert out['written'] is False
    assert out['aborted'] == 'post-write verification failed; workbook restored'
    wb = openpyxl.load_workbook(p)
    assert wb[SHEET].cell(row=3, column=6).value == 'Applied'


def test_only_the_named_rows_change(tmp_path, profiles, monkeypatch):
    p = make_workbook(tmp_path, [
        ('Acme', 'Graduate Analyst', 'https://x/1', 'Applied'),
        ('Beta', 'SOC Analyst', 'https://x/2', 'Applied'),
        ('Gamma', 'Cyber Intern', 'https://x/3', 'Not Applied'),
    ])
    monkeypatch.setattr(tss.tw, 'verify_workbook',
                        lambda *a, **k: {'ok': True, 'problems': []})
    tss.apply_changes(p, [{'row': 4, 'status': 'Rejected'}],
                      profiles=profiles, backup_dir=tmp_path / 'b', dry_run=False)
    wb = openpyxl.load_workbook(p)
    ws = wb[SHEET]
    assert ws.cell(row=3, column=6).value == 'Applied'
    assert ws.cell(row=4, column=6).value == 'Rejected'
    assert ws.cell(row=5, column=6).value == 'Not Applied'
