import datetime as dt
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
import tracker_writer as tw
import tracker_dedupe_repair as repair
from tracker_mail_reconcile import evidence_status, reconcile_file


def workbook(tmp_path, rows):
    path = tmp_path / 'uk.xlsx'
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Jobs'
    ws.append(['UK jobs'])
    ws.append(['Date Found', 'Company', 'Role Title', 'Apply Link', 'Application Deadline', 'Status'])
    for row in rows:
        ws.append(row)
        ws.cell(ws.max_row, 4).hyperlink = row[3]
    rule = DataValidation(type='list', formula1='"Applied,Assessment Pending,Interview Scheduled,Rejected,Not Applied"')
    ws.add_data_validation(rule); rule.add('F3:F500')
    wb.save(path); wb.close()
    return path


def test_year_prefix_company_variant_and_qualifier_identity():
    assert tw.pair_key('Tencent', 'Cyber Security Intern') == tw.pair_key('Tencent UK - Prosple', 'Cyber Security Intern')
    assert tw.pair_key('Barclays', 'Technology Cyber and Security Summer Internship Programme') == tw.pair_key('Barclays', '2027 Technology Cyber and Security Summer Internship Programme')
    assert tw.pair_key('KnowBe4', 'Graduate Cybersecurity Threat Researcher') == tw.pair_key('KnowBe4', 'Graduate Cybersecurity Threat Researcher (Position located in Cheltenham, United Kingdom)')


def test_same_title_different_requisition_cannot_append(tmp_path):
    old = ['2026-10-20', 'Apple', 'Information Security Internship', 'https://jobs.example/jobs/123456', None, 'Applied']
    path = workbook(tmp_path, [old])
    record = {'company': 'Apple', 'title': old[2], 'url': 'https://www.linkedin.com/jobs/view/987654', 'date_found': '2026-10-25'}
    profile = tw.load_profiles()
    outcome = tw.write_records(profile, 'uk', [record], tracker_override=str(path), apply=False, cleanup_existing=False)
    assert outcome['counts']['duplicates'] == 1
    assert outcome['counts']['appended'] == 0


def test_merge_precedence_and_idempotency(tmp_path):
    base = ['2026-10-20', 'Tencent', 'Cyber Security Intern', 'https://jobs.example/jobs/123456', None]
    path = workbook(tmp_path, [base + ['Not Applied'],
                               ['2026-10-21', 'Tencent UK', base[2], 'https://www.linkedin.com/jobs/view/789012', None, 'Rejected'],
                               ['2026-10-22', 'Tencent UK - Prosple', base[2], 'https://www.prosple.com/jobs/345678', None, 'Applied']])
    dry = repair.repair(path)
    assert not dry['applied'] and dry['removed_rows'] == 2
    result = repair.repair(path, apply=True, backup_dir=tmp_path / 'backups')
    assert result['applied'] and pathlib.Path(result['backup']).exists()
    wb = openpyxl.load_workbook(path); ws = wb['Jobs']
    assert ws.max_row == 3 and ws['F3'].value == 'Applied'
    assert tw.parse_date(ws['A3'].value, False) == dt.date(2026,10,20)
    wb.close()
    before = tw.sha256_file(path)
    assert repair.repair(path, apply=True)['removed_rows'] == 0
    assert tw.sha256_file(path) == before


def test_sort_status_then_newest_date(tmp_path):
    rows = [['2026-10-20', 'A', 'Role A', 'https://x.example/jobs/111111', None, 'Rejected'],
            ['2026-10-21', 'B', 'Role B', 'https://x.example/jobs/222222', None, 'Not Applied'],
            ['2026-10-25', 'C', 'Role C', 'https://x.example/jobs/333333', None, 'Applied'],
            ['2026-10-24', 'D', 'Role D', 'https://x.example/jobs/444444', None, 'Not Applied']]
    path = workbook(tmp_path, rows)
    wb = openpyxl.load_workbook(path); ws = wb['Jobs']
    tw.sort_status_date(ws, dict(tw.region_config(tw.load_profiles(), 'uk'), region='uk'))
    assert [ws.cell(r,2).value for r in range(3,7)] == ['D','B','C','A']
    wb.close()


def test_writer_places_new_open_row_above_rejected(tmp_path):
    path = workbook(tmp_path, [['2026-10-20', 'Old', 'Rejected Role',
                                'https://x.example/jobs/111111', None, 'Rejected']])
    result = tw.write_records(tw.load_profiles(), 'uk', [{
        'company': 'New', 'title': 'Open Role', 'url': 'https://x.example/jobs/222222',
        'date_found': '2026-10-25'}], tracker_override=str(path), apply=True,
        backup_dir=str(tmp_path / 'backups'), cleanup_existing=False)
    assert result['applied']
    wb = openpyxl.load_workbook(path); ws = wb['Jobs']
    assert [(ws.cell(r, 2).value, ws.cell(r, 6).value) for r in (3, 4)] == [
        ('New', 'Not Applied'), ('Old', 'Rejected')]
    wb.close()


def test_explicit_mail_only_and_replay(tmp_path):
    assert evidence_status('Thank you for applying to Canonical', '') == 'Applied'
    assert evidence_status('We have received your application', '') == 'Applied'
    assert evidence_status('Thank You for Your Application', '') == 'Applied'
    assert evidence_status('Application update', 'We have decided not to move forward with your application.') == 'Rejected'
    assert evidence_status('Application update', 'Your application is under review.') is None
    path = workbook(tmp_path, [['2026-10-20', 'Canonical', 'Cyber Security Intern', 'https://jobs.example/jobs/123456', None, 'Not Applied']])
    message = {'id': 'm1', 'thread_id': 't1', 'subject': 'Thank you for applying to Canonical',
               'body': 'Your application for Cyber Security Intern at Canonical has been received.',
               'from': 'Canonical <jobs@canonical.com>', 'received_at': '2026-10-22T10:00:00Z'}
    audit = tmp_path / 'audit.json'
    first = reconcile_file(path, [message], audit_path=audit, apply=True)
    assert first['receipts'][0]['decision'] == 'update'
    again = reconcile_file(path, [message], audit_path=audit, apply=True)
    assert again['receipts'] == []
    wb = openpyxl.load_workbook(path); assert wb['Jobs']['F3'].value == 'Applied'; wb.close()
    rejection = dict(message, id='m2', subject='Your Canonical application update',
                     body='We will not be moving forward with your application for Cyber Security Intern at Canonical.')
    changed = reconcile_file(path, [rejection], audit_path=audit, apply=True)
    assert changed['receipts'][0]['to'] == 'Rejected'
    wb = openpyxl.load_workbook(path); assert wb['Jobs']['F3'].value == 'Rejected'; wb.close()
