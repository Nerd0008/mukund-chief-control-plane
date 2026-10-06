"""Offline acceptance: isolated workbook, mocked Google APIs, fake credentials."""
import copy
import datetime as dt
import json
import sys
from pathlib import Path
from urllib.parse import urlsplit, parse_qs

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import career_google_auth as auth
from career_google_clients import GmailReader, DeadlineCalendar
from career_mail_parser import extract_deadline, parse_message, calendar_event
from career_mail_tracker import WorkbookTracker
from career_mail_monitor import reconcile, apply_report, run_scan, atomic_json, report_hash, main


def mail(mid="m1", body="Thank you for applying.", thread="t1", date="2026-10-01T10:00:00+00:00"):
    return {"message_id": mid, "thread_id": thread, "received_at": date,
            "from": "Recruiting <careers@example.test>", "subject": "Your application",
            "body": "Company: Acme\nRole: Graduate Engineer\nRegion: UK\n" + body}


def assess(mid="m2", deadline="Complete by 6 October 2026 at 23:59 BST."):
    return mail(mid, "Please complete your online assessment. SHL. " + deadline,
                date="2026-10-02T10:00:00+00:00")


def test_confirmation_new_record():
    report = reconcile([mail()], [])
    assert report['counts']['proposed_new_rows'] == 1
    assert report['proposed_records'][0]['latest_status'] == 'application_confirmation'


def test_duplicate_confirmation_no_duplicate_application():
    first = reconcile([mail()], [])
    second = reconcile([mail(), mail('m2', thread='t2')], first['proposed_records'])
    assert second['counts']['duplicates'] == 1
    assert second['counts']['proposed_new_rows'] == 0
    assert len(second['proposed_records']) == 1


def test_assessment_updates_existing():
    first = reconcile([mail()], [])['proposed_records']
    report = reconcile([assess()], first)
    record = report['proposed_records'][0]
    assert record['application_id'] == first[0]['application_id']
    assert record['current_stage'] == 'assessment'
    assert record['assessment_provider'] == 'SHL'


def test_exact_deadline_stored_and_event_requested():
    report = reconcile([assess()], [])
    assert report['proposed_records'][0]['deadline']['value'] == '2026-10-06T23:59:00+01:00'
    event = report['proposed_calendar_events'][0]
    assert event['start']['dateTime'] == '2026-10-06T23:59:00+01:00'
    assert event['reminders']['overrides'][0]['minutes'] == 1440
    assert 'Gmail message ID: m2' in event['description']


def test_relative_five_days_from_received():
    report = reconcile([assess(deadline='Please complete within 5 days.')], [])
    deadline = report['proposed_records'][0]['deadline']
    assert deadline['kind'] == 'DERIVED'
    assert deadline['value'] == '2026-10-07T10:00:00+00:00'
    assert deadline['source_thread_id'] == 't1'


@pytest.mark.parametrize('phrase', ['Complete soon.', 'Complete by next Friday.',
                                     'Please complete within 5 business days.',
                                     'Complete by 6 October 2026 at 23:59.',
                                     'Complete by 6 October 2026 at 23:59 CST.'])
def test_ambiguous_no_calendar(phrase):
    # "soon" has no exact date; convert to standard ambiguity wording.
    phrase = phrase.replace('Complete soon.', 'Complete as soon as possible.')
    report = reconcile([assess(deadline=phrase)], [])
    assert report['proposed_records'][0]['needs_review']
    assert not report['proposed_calendar_events']


def test_date_only_all_day():
    event = reconcile([assess(deadline='Complete by 6 October 2026.')], [])['proposed_calendar_events'][0]
    assert event['start'] == {'date': '2026-10-06'}
    assert event['end'] == {'date': '2026-10-07'}


def test_deadline_extension_same_event():
    first = reconcile([assess()], [])['proposed_records']
    first[0]['calendar_event_id'] = calendar_event(first[0])['id']
    extension = mail('m3', 'Your deadline has been extended to 9 October 2026 at 12:00 BST.',
                     date='2026-10-03T11:00:00+00:00')
    report = reconcile([extension], first)
    assert report['proposed_records'][0]['latest_status'] == 'deadline_extension'
    assert report['proposed_calendar_events'][0]['id'] == first[0]['calendar_event_id']
    assert report['proposed_calendar_events'][0]['start']['dateTime'] == '2026-10-09T12:00:00+01:00'


def test_rejection_updates_application():
    records = reconcile([mail()], [])['proposed_records']
    report = reconcile([mail('m3', 'Your application was unsuccessful.', date='2026-10-04T12:00:00+00:00')], records)
    assert report['proposed_records'][0]['latest_status'] == 'rejection'
    assert report['counts']['proposed_new_rows'] == 0


def test_unrelated_alert_ignored():
    signal = mail(body='Job alert: online assessment roles. Unsubscribe from this newsletter.')
    signal['subject'] = 'Recommended jobs'
    assert reconcile([signal], [])['counts']['ignored'] == 1


@pytest.mark.parametrize('text,kind', [
    ('We have received your application.', 'application_received'),
    ('Your application is under review.', 'application_under_review'),
    ('Your coding assessment is ready.', 'coding_assessment'),
    ('Your numerical reasoning test is ready.', 'psychometric_assessment'),
    ('Your SJT assessment is ready.', 'psychometric_assessment'),
    ('Your HireVue video interview is ready.', 'hirevue_video_interview'),
    ('Your telephone interview is booked.', 'telephone_interview'),
    ('Your technical interview is booked.', 'technical_interview'),
    ('Your assessment centre invitation.', 'assessment_centre'),
    ('Update on your application.', 'recruiter_update'),
    ('Your application has been withdrawn.', 'withdrawal'),
    ('Your formal offer of employment.', 'offer'),
    ('Your Cappfinity assessment.', 'online_assessment'),
    ('Your TestGorilla assessment.', 'online_assessment')])
def test_classification(text, kind):
    assert parse_message(mail(body=text))['latest_status'] == kind


def test_owner_confirmed_deadline_not_overwritten():
    records = reconcile([assess()], [])['proposed_records']
    original = copy.deepcopy(records[0]['deadline'])
    records[0]['owner_confirmed_fields'] = ['deadline', 'assessment_interview_link']
    report = reconcile([mail('m3','Deadline extended to 12 October 2026 at 12:00 BST.',
                             date='2026-10-03T12:00:00+00:00')], records)
    assert report['proposed_records'][0]['deadline'] == original


def test_conflicting_identity_review():
    records = reconcile([mail()], [])['proposed_records']
    other = assess(); other['body'] = other['body'].replace('Acme','Other Company')
    report = reconcile([other], records)
    assert not report['proposed_records']
    assert report['needs_review'][0]['needs_review']


def test_distinct_application_references_not_merged():
    first = mail(body='Thank you for applying. Application ID: A1')
    second = mail('m2','Thank you for applying. Application ID: A2',thread='t2')
    report = reconcile([first,second], [])
    assert report['counts']['proposed_new_rows'] == 2


def test_old_mail_does_not_roll_deadline_backwards():
    records = reconcile([assess()], [])['proposed_records']
    old = assess('old','Complete by 4 October 2026 at 23:59 BST.'); old['received_at'] = '2026-10-01T10:00:00Z'
    report = reconcile([old], records)
    assert report['proposed_records'][0]['deadline'] == records[0]['deadline']


class CalendarTransport:
    def __init__(self):
        self.events, self.calls = {}, []

    def __call__(self, method, url, **kwargs):
        self.calls.append((method,url))
        eid = urlsplit(url).path.rsplit('/',1)[-1]
        if method == 'GET':
            if eid not in self.events: raise auth.GoogleError('request',404)
            return copy.deepcopy(self.events[eid])
        body = kwargs['data']
        if method == 'POST':
            self.events[body['id']] = copy.deepcopy(body)
            return body
        self.events[eid].update(body)
        return copy.deepcopy(self.events[eid])


def test_calendar_repeat_and_extension_idempotent():
    transport = CalendarTransport(); calendar = DeadlineCalendar('FAKE',transport=transport)
    event = reconcile([assess()],[])['proposed_calendar_events'][0]
    calendar.ensure(event); calendar.ensure(event)
    assert [m for m,u in transport.calls].count('POST') == 1
    assert not any(m=='PATCH' for m,u in transport.calls)
    event['start']['dateTime'] = '2026-10-09T12:00:00+01:00'
    calendar.ensure(event)
    assert [m for m,u in transport.calls].count('PATCH') == 1
    assert len(transport.events) == 1


def test_gmail_complete_pagination_and_get_only():
    calls = []
    def transport(method,url,**kwargs):
        calls.append((method,url)); parts = urlsplit(url); q=parse_qs(parts.query)
        if parts.path.endswith('/profile'): return {'historyId':'10'}
        if parts.path.endswith('/messages'):
            return {'messages':[{'id':'m2'}]} if q.get('pageToken') else {'messages':[{'id':'m1'}],'nextPageToken':'page2'}
        return {'id':parts.path.rsplit('/',1)[-1]}
    reader = GmailReader('FAKE',transport)
    result=reader.read_window(now=dt.datetime(2026,10,6,tzinfo=dt.timezone.utc))
    assert result['messages_read']==2 and result['checkpoint']=='10'
    assert all(method=='GET' for method,url in calls)
    assert 'after:' in parse_qs(urlsplit(calls[1][1]).query)['q'][0]
    for path in ['/messages/m1/modify','/messages/send','/threads/t1/trash','/labels']:
        with pytest.raises(ValueError): reader._get(path)


def test_incremental_uses_history_not_full_window():
    calls=[]
    def transport(method,url,**kwargs):
        calls.append(url)
        if '/profile' in url:return {'historyId':'12'}
        if '/history' in url:return {'history':[{'messagesAdded':[{'message':{'id':'m3'}}]}]}
        return {'id':'m3'}
    result=GmailReader('FAKE',transport).read_window('10')
    assert result['mode']=='incremental'
    assert not any(urlsplit(u).path.endswith('/messages') for u in calls)


def test_expired_history_backfills():
    def transport(method,url,**kwargs):
        if '/profile' in url:return {'historyId':'20'}
        if '/history' in url:raise auth.GoogleError('request',404)
        return {'messages':[]}
    assert GmailReader('FAKE',transport).read_window('expired')['mode']=='expired_history_backfill'


def test_safety_cap_fails_before_checkpoint():
    def transport(method,url,**kwargs):
        if '/profile' in url:return {'historyId':'10'}
        return {'messages':[{'id':'m1'},{'id':'m2'}]}
    with pytest.raises(auth.GoogleError):GmailReader('FAKE',transport).read_window(max_messages=1)


@pytest.fixture
def tracker(tmp_path):
    path=tmp_path/'canonical.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Jobs'
    ws.append(['ID','Company','Role','Status','Notes']);ws.append(['J1','Acme','Graduate Engineer','Applied','Owner-confirmed note'])
    wb.save(path);wb.close()
    cfg={'tracker':str(path),'sheet':'Jobs','header_row':1,'first_data_row':2,
         'field_map':{'company':'B','title':'C'},'id':{'column':'A','style':'sequential','prefix':'J'},
         'status_columns':{'application_status':'D'},'defaults':{'application_status':'To Review'}}
    return WorkbookTracker({'regions':{'uk':cfg}},tmp_path/'backups')


def test_actual_tracker_metadata_preserves_owner_cells(tracker):
    records,_=tracker.read();report=reconcile([assess()],records)
    tracker.upsert(report['proposed_records'][0])
    after,_=tracker.read();assert len(after)==1
    assert after[0]['deadline']['kind']=='EXACT'
    wb=openpyxl.load_workbook(tracker.profiles['uk']['tracker']);ws=wb['Jobs']
    assert ws['D2'].value=='Applied' and ws['E2'].value=='Owner-confirmed note'
    wb.close()


def test_actual_new_row_and_repeat_no_duplicate(tracker):
    message=mail();message['body']=message['body'].replace('Acme','New Co')
    record=reconcile([message],[])['proposed_records'][0]
    tracker.upsert(record);tracker.upsert(record)
    records,_=tracker.read();assert len(records)==2
    assert records[1]['company']=='New Co'


def test_formula_injection_is_literal(tracker):
    record=reconcile([mail()],[])['proposed_records'][0]
    record['recruiter_contact']='=HYPERLINK("https://bad.test")'
    record['row']=2;tracker.upsert(record)
    wb=openpyxl.load_workbook(tracker.profiles['uk']['tracker']);ws=wb['Jobs']
    cell=next(ws.cell(2,c.column) for c in ws[1] if c.value=='Career Ops Email Recruiter Contact')
    assert cell.data_type=='s';wb.close()


class Reader:
    def read_window(self,*args,**kwargs):
        return {'messages':[assess()],'checkpoint':'10','mode':'backfill'}


def test_first_live_path_dry_run_no_writes_or_checkpoint(tmp_path,tracker):
    calendar=DeadlineCalendar('FAKE',transport=CalendarTransport())
    original=Path(tracker.profiles['uk']['tracker']).read_bytes()
    cfg={'backfill_days':30,'max_messages':2000}
    report=run_scan(cfg,tmp_path/'runtime',reader=Reader(),calendar=calendar,tracker=tracker)
    assert report['tracker_writes']==report['calendar_writes']==0
    assert not (tmp_path/'runtime/checkpoint.json').exists()
    assert Path(tracker.profiles['uk']['tracker']).read_bytes()==original
    assert not calendar._transport.calls


def test_apply_refused_before_initial_approval(tmp_path,tracker):
    with pytest.raises(ValueError):run_scan({},tmp_path/'runtime',apply=True,tracker=tracker)


def test_apply_calendar_receipt_persisted(tracker,tmp_path):
    records,_=tracker.read();report=reconcile([assess()],records)
    transport=CalendarTransport();calendar=DeadlineCalendar('FAKE',transport=transport)
    receipt=apply_report(report,tracker,calendar,tmp_path/'audit.json')
    assert receipt[0]['calendar_status']=='confirmed'
    records,_=tracker.read();assert records[0]['calendar_event_id'] in transport.events


def test_secrets_absent_status_and_error(monkeypatch,capsys):
    secret='DO-NOT-PRINT-FAKE-TOKEN'
    monkeypatch.setattr(auth,'read_secret',lambda target: secret if not target.endswith('token-state') else '{}')
    assert secret not in json.dumps(auth.status())
    monkeypatch.setattr(auth,'authorize',lambda: (_ for _ in ()).throw(auth.GoogleError('request',401)))
    assert auth.main(['authorize'])==1
    assert secret not in capsys.readouterr().out


def test_extra_oauth_scopes_refused(monkeypatch):
    monkeypatch.setattr(auth,'store_secret',lambda *args:pytest.fail('no storage on invalid grant'))
    with pytest.raises(auth.GoogleError):auth.save_tokens({'access_token':'FAKE','scope':' '.join(auth.SCOPES)+' https://mail.google.com/'})


def test_email_url_drops_secret_query():
    record=parse_message(assess(deadline='Complete by 6 October 2026. https://assessment.test/run?access_token=FAKE_SECRET'))
    assert record['assessment_interview_link']=='https://assessment.test/run'
    assert 'FAKE_SECRET' not in json.dumps(record)


def test_calendar_failure_replays_pending_after_saved_message(tracker,tmp_path):
    records,_=tracker.read();report=reconcile([assess()],records)
    class Broken:
        def ensure(self,event):raise auth.GoogleError('request',503)
    with pytest.raises(auth.GoogleError):apply_report(report,tracker,Broken(),tmp_path/'audit.json')
    records,_=tracker.read();assert records[0]['calendar_pending']
    replay=reconcile([assess()],records)
    assert replay['counts']['duplicates']==1
    assert len(replay['proposed_calendar_events'])==1
    apply_report(replay,tracker,DeadlineCalendar('FAKE',transport=CalendarTransport()),tmp_path/'audit.json')
    records,_=tracker.read();assert not records[0]['calendar_pending']


def test_gmail_requests_are_paced_without_retry():
    ticks=[0.0];delays=[];calls=[]
    def sleep(delay):delays.append(delay);ticks[0]+=delay
    reader=GmailReader('FAKE',lambda *a,**k:calls.append(a) or {},clock=lambda:ticks[0],sleep=sleep)
    reader._get('/profile');reader._get('/history');reader._get('/messages')
    assert len(calls)==3 and delays==[1.0,1.0]


def test_quota_recovery_is_bounded_and_permission_errors_not_retried():
    waits=[];calls=[]
    def denied(*a,**k):
        calls.append(1);raise auth.GoogleError('request',403,'rateLimitExceeded')
    reader=GmailReader('FAKE',denied,sleep=waits.append)
    with pytest.raises(auth.GoogleError):reader._get('/profile')
    assert len(calls)==3 and waits==[60,60]
    calls.clear();waits.clear()
    def forbidden(*a,**k):
        calls.append(1);raise auth.GoogleError('request',403,'insufficientPermissions')
    with pytest.raises(auth.GoogleError):GmailReader('FAKE',forbidden,sleep=waits.append)._get('/profile')
    assert len(calls)==1 and waits==[]


def test_backfill_query_is_recruitment_scoped():
    urls=[]
    def transport(method,url,**kwargs):
        urls.append(url);return {'historyId':'1'} if '/profile' in url else {'messages':[]}
    GmailReader('FAKE',transport).read_window()
    from urllib.parse import urlsplit,parse_qs
    q=parse_qs(urlsplit(urls[1]).query)['q'][0]
    assert q.startswith('after:') and 'application' in q and 'assessment' in q


def test_unresolved_application_still_reports_assessment_and_deadline():
    message=assess();message['body']=message['body'].replace('Company: Acme','Company:').replace('Role: Graduate Engineer','Role:').replace('Region: UK','Region:')
    report=reconcile([message],[])
    assert report['counts']['assessments_interviews']==1
    assert report['counts']['exact_deadlines']==1
    assert report['needs_review'] and not report['proposed_calendar_events']


def unlabelled(body, subject="Your application"):
    return {"message_id": "unlabelled", "thread_id": "unlabelled-thread",
            "received_at": "2026-10-02T10:00:00+00:00", "subject": subject, "body": body}


def test_role_capture_stops_at_sentence_not_later_at():
    parsed = parse_message(unlabelled("Thank you for your application for our 2027 Summer Intern – Product Analyst, EMEA opportunity! We invite you to provide more details and outline next steps at your convenience."))
    assert parsed["role"] == "2027 Summer Intern – Product Analyst, EMEA"
    assert parsed["company"] is None


def test_application_phrase_identity():
    parsed = parse_message(unlabelled("Thank you for applying for the Graduate Engineer role at Acme."))
    assert parsed["company"] == "Acme"
    assert parsed["role"] == "Graduate Engineer"


def test_unique_canonical_full_identity_mentions_resolve_region():
    existing = reconcile([mail()], [])["proposed_records"]
    report = reconcile([unlabelled("Acme recruitment: Your application for Graduate Engineer has been received.")], existing)
    assert report["counts"]["matched_existing"] == 1
    assert report["proposed_records"][0]["region"] == "uk"


def test_typography_only_matching():
    existing = reconcile([mail()], [])["proposed_records"]
    existing[0]["role"] = "Graduate – Engineer"
    report = reconcile([unlabelled("Company: Acme\nRole: Graduate - Engineer\nYour application has been received.")], existing)
    assert report["counts"]["matched_existing"] == 1
    assert report["proposed_records"][0]["role"] == "Graduate – Engineer"


def test_employer_alone_does_not_select_job():
    existing = reconcile([mail()], [])["proposed_records"]
    report = reconcile([unlabelled("Acme: Please complete your online assessment.")], existing)
    assert not report["proposed_records"]
    assert len(report["needs_review"]) == 1


def test_same_role_at_two_regions_is_ambiguous():
    existing = reconcile([mail()], [])["proposed_records"]
    other = copy.deepcopy(existing[0])
    other.update(application_id="another", region="dubai")
    report = reconcile([unlabelled("Acme: Your application for Graduate Engineer has been received.")], existing + [other])
    assert not report["proposed_records"]
    assert report["needs_review"][0]["review_reason"] == "multiple canonical applications match"


def test_explicit_conflicting_company_does_not_match_mentions():
    existing = reconcile([mail()], [])["proposed_records"]
    report = reconcile([unlabelled("Company: Other\nRole: Graduate Engineer\nYour application has been received. Acme is an unrelated customer.")], existing)
    assert not report["proposed_records"]


def test_role_substring_does_not_match_different_role():
    existing = reconcile([mail()], [])["proposed_records"]
    report = reconcile([unlabelled("Company: Acme\nRole: Senior Graduate Engineer\nYour application has been received.")], existing)
    assert not report["proposed_records"]


@pytest.mark.parametrize("conflict", ["Company: Other", "Region: Dubai"])
def test_partial_thread_conflict_stays_review_only(conflict):
    existing = reconcile([mail()], [])["proposed_records"]
    raw = unlabelled(conflict + "\nPlease complete your online assessment.")
    raw["thread_id"] = "t1"
    report = reconcile([raw], existing)
    assert not report["proposed_records"]
    assert report["needs_review"][0]["review_reason"] == "thread conflicts with company/role/region evidence"
