"""Set the owner Status column (F) on specific UK tracker rows, safely.

The canonical writer refuses to touch application state, by design: Status is
the owner's field, set by the owner from evidence. This tool exists for the
narrow case where an employer's own email proves a state change, and it does
exactly that and nothing else.

Safety: backup first, atomic replace, then the same integrity gate the pipeline
uses. Any mismatch aborts and leaves the workbook untouched.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import tracker_writer as tw  # noqa: E402

DEFAULT_TRACKER = Path(r'C:\Users\mukun\Downloads\codex\uk-cyber-job-tracker.xlsx')
DEFAULT_BACKUP_DIR = HERE.parent / 'runtime' / 'career-ops' / 'backups'
SHEET = 'Jobs'
HEADER_ROW = 2
STATUS_COL = 6


def allowed_statuses(path: Path) -> list[str]:
    """Read the Status dropdown so we never write a value Excel would reject."""
    wb = openpyxl.load_workbook(path, read_only=False)
    ws = wb[SHEET]
    out: list[str] = []
    for dv in ws.data_validations.dataValidation:
        if dv.type == 'list' and dv.formula1:
            out = [s.strip() for s in dv.formula1.strip('"').split(',')]
            break
    wb.close()
    return out


def apply_changes(path: Path, changes: list[dict], *, profiles: dict, backup_dir: Path,
                  dry_run: bool = False) -> dict:
    cfg = tw.region_config(profiles, 'uk')
    before_rows = len(tw.Tracker(path, cfg).data_rows())

    wb = openpyxl.load_workbook(path, keep_links=True)
    ws = wb[SHEET]

    applied, refused = [], []
    for ch in changes:
        row = int(ch['row'])
        want = str(ch['status'])
        company = ws.cell(row=row, column=2).value
        title = ws.cell(row=row, column=3).value
        current = ws.cell(row=row, column=STATUS_COL).value
        expect = ch.get('expect_status')
        if expect is not None and str(current) != str(expect):
            refused.append({'row': row, 'reason': f'status is {current!r}, expected {expect!r}'})
            continue
        if not company:
            refused.append({'row': row, 'reason': 'no company in column B'})
            continue
        applied.append({'row': row, 'company': company, 'title': title,
                        'from': current, 'to': want})
        ws.cell(row=row, column=STATUS_COL, value=want)

    result = {'applied': applied, 'refused': refused, 'rows_before': before_rows}

    if refused:
        wb.close()
        result['written'] = False
        result['aborted'] = 'refused rows present; workbook untouched'
        return result

    if dry_run:
        wb.close()
        result['written'] = False
        result['mode'] = 'dry-run'
        return result

    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
    backup = backup_dir / f'{path.stem}.pre-status-{stamp}{path.suffix}'
    shutil.copy2(path, backup)

    tmp = path.with_suffix('.tmp.xlsx')
    wb.save(tmp)
    wb.close()
    tmp.replace(path)

    v = tw.verify_workbook(path, cfg, expect_data_rows=before_rows)
    ok = bool(v.get('ok'))
    problems = list(v.get('problems') or [])
    result.update({'written': True, 'backup': str(backup),
                   'verification_ok': ok, 'problems': problems})
    if not ok:
        shutil.copy2(backup, path)
        result['written'] = False
        result['aborted'] = 'post-write verification failed; workbook restored'
    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--tracker', default=str(DEFAULT_TRACKER))
    ap.add_argument('--backup-dir', default=str(DEFAULT_BACKUP_DIR))
    ap.add_argument('--changes', required=True, help='JSON file: [{"row":120,"status":"Rejected"}]')
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    path = Path(args.tracker)
    changes = json.loads(Path(args.changes).read_text(encoding='utf-8'))

    allowed = allowed_statuses(path)
    print('allowed statuses:', allowed)
    for ch in changes:
        if allowed and ch['status'] not in allowed:
            print(f"REFUSING: {ch['status']!r} is not in the dropdown {allowed}")
            return 2

    profiles = tw.load_profiles(None)
    out = apply_changes(path, changes, profiles=profiles,
                        backup_dir=Path(args.backup_dir), dry_run=args.dry_run)
    print(json.dumps(out, ensure_ascii=False, indent=2))
    return 0 if not out.get('aborted') else 1


if __name__ == '__main__':
    raise SystemExit(main())
