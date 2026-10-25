"""Dry-run-first repair for duplicate UK vacancies."""
from __future__ import annotations
import argparse
import copy
import datetime as dt
import json
import os
from pathlib import Path
import shutil
import uuid
from urllib.parse import urlsplit

import openpyxl
import tracker_writer as tw
from posting_identity import identity_keys, is_aggregator, company_key
from tracker_integrity import bind_links


def _date(value):
    parsed = tw.parse_date(value, False)
    return parsed if isinstance(parsed, dt.date) else None


def plan_groups(ws, first=3):
    groups = []
    keys = {}
    for row in range(first, ws.max_row + 1):
        company, title, url = (ws.cell(row, c).value for c in (2, 3, 4))
        if not company or not title:
            continue
        identity = {"pair:" + tw.pair_key(company, title)}
        if tw.normalize_url(url):
            identity.add("url:" + tw.normalize_url(url))
        identity |= identity_keys(company, url)
        hits = set()
        for key in identity:
            if key not in keys:
                continue
            candidate = keys[key]
            if key.startswith('pair:') and not is_aggregator(url):
                direct = [r for r in groups[candidate] if not is_aggregator(ws.cell(r, 4).value)]
                if direct and all(urlsplit(str(url)).hostname == urlsplit(str(ws.cell(r, 4).value)).hostname and
                                  company_key(company) != 'tencent' and
                                  identity_keys(company, url) and
                                  identity_keys(ws.cell(r, 2).value, ws.cell(r, 4).value) and
                                  not (identity_keys(company, url) & identity_keys(ws.cell(r, 2).value, ws.cell(r, 4).value))
                                  for r in direct):
                    continue
            hits.add(candidate)
        group = min(hits) if hits else len(groups)
        if not hits:
            groups.append(set())
        for other in hits - {group}:
            groups[group].update(groups[other])
            groups[other].clear()
            for key, index in list(keys.items()):
                if index == other:
                    keys[key] = group
        groups[group].add(row)
        for key in identity:
            keys[key] = group
    return [sorted(g) for g in groups if len(g) > 1]


def merge_sheet(ws, first=3):
    receipts = []
    deleted = []
    for group in plan_groups(ws, first):
        def quality(row):
            url = ws.cell(row, 4).value
            return (bool(tw.normalize_url(url) and not tw.listing_page_reason(url)),
                    bool(_date(ws.cell(row, 1).value)), -row)
        survivor = max(group, key=quality)
        states = [str(ws.cell(r, 6).value or '') for r in group]
        state = next((s for s in ('Applied', 'Rejected') if s in states),
                     ws.cell(survivor, 6).value)
        dates = [_date(ws.cell(r, 1).value) for r in group]
        earliest = min((d for d in dates if d), default=None)
        old_date = _date(ws.cell(survivor, 1).value)
        if earliest and (not old_date or earliest < old_date):
            ws.cell(survivor, 1).value = earliest
        ws.cell(survivor, 6).value = state
        removed = [r for r in group if r != survivor]
        deleted.extend(removed)
        receipts.append({'rows': group, 'survivor': survivor, 'removed': removed,
                         'status': state, 'date_found': str(ws.cell(survivor, 1).value),
                         'apply_link': ws.cell(survivor, 4).value})
    for row in sorted(deleted, reverse=True):
        ws.delete_rows(row)
    return receipts


def repair(path, *, apply=False, backup_dir=None):
    path = Path(path)
    cfg = dict(tw.region_config(tw.load_profiles(), 'uk'), region='uk')
    lock = path.with_suffix(path.suffix + '.career-mail.lock')
    fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    staged = path.with_name(path.stem + '.dedupe-' + uuid.uuid4().hex + '.xlsx')
    try:
        before = tw.sha256_file(path)
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[cfg['sheet']]
            if ws.max_column != 6:
                raise ValueError('Expected six UK columns')
            receipts = merge_sheet(ws, cfg['first_data_row'])
            result = {'applied': False, 'groups': receipts, 'removed_rows': sum(len(r['removed']) for r in receipts)}
            if not receipts:
                return result
            bind_links(ws, cfg)
            tw.sort_status_date(ws, cfg)
            wb.save(staged)
        finally:
            wb.close()
        verification = tw.verify_workbook(staged, cfg)
        if not verification['ok']:
            raise ValueError('Staged workbook failed verification: ' + repr(verification['problems']))
        result['verification'] = verification
        if not apply:
            return result
        backup_dir = Path(backup_dir or path.parent)
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / (path.stem + '.pre-dedupe-' + uuid.uuid4().hex + '.xlsx')
        shutil.copy2(path, backup)
        if tw.sha256_file(backup) != before or tw.sha256_file(path) != before:
            raise ValueError('Workbook changed before repair')
        os.replace(staged, path)
        post = tw.verify_workbook(path, cfg)
        if not post['ok']:
            shutil.copy2(backup, path)
            raise ValueError('Post-write verification failed; backup restored')
        result.update(applied=True, backup=str(backup), verification=post)
        return result
    finally:
        staged.unlink(missing_ok=True)
        os.close(fd)
        lock.unlink()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--tracker', default=str(tw.region_config(tw.load_profiles(), 'uk')['tracker']))
    parser.add_argument('--backup-dir')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    print(json.dumps(repair(args.tracker, apply=args.apply, backup_dir=args.backup_dir), default=str, indent=2))


if __name__ == '__main__':
    main()
