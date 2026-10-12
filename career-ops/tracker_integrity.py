"""Canonical workbook schema and link integrity; run before committing staged files."""
from urllib.parse import urlsplit
from openpyxl.utils import column_index_from_string as ci

def problems(ws,cfg):
    errors=[];ceiling=cfg.get('column_ceiling')
    if ceiling and ws.max_column>ceiling:errors.append('tracker exceeds fixed column ceiling')
    for col,label in cfg.get('required_headers',{}).items():
        if ws.cell(cfg['header_row'],ci(col)).value!=label:errors.append('required header changed: '+col)
    col=ci(cfg['dedupe']['url_column'])
    for row in range(cfg['first_data_row'],ws.max_row+1):
        cell=ws.cell(row,col);link=getattr(cell,'hyperlink',None)
        value=str(cell.value or '').strip()
        if value.startswith(('https://','http://')) and not all(ws.cell(row,ci(cfg['field_map'][key])).value for key in ('company','title')):
            errors.append('URL-only row has no application identity: '+cell.coordinate)
        if link and (link.target!=value or link.ref!=cell.coordinate):errors.append('hyperlink identity mismatch: '+cell.coordinate)
        if cfg.get('hyperlinks_required') and value.startswith(('https://','http://')) and not link and not getattr(ws,'_read_only',False):errors.append('missing hyperlink: '+cell.coordinate)
    return errors

def assert_valid(ws,cfg):
    errors=problems(ws,cfg)
    if errors:raise ValueError('tracker integrity failed: '+ '; '.join(errors[:10]))

def bind_links(ws,cfg):
    """Bind from the row's own visible URL, never a copied Hyperlink/ref object."""
    col=ci(cfg['dedupe']['url_column'])
    for row in range(cfg['first_data_row'],ws.max_row+1):
        cell=ws.cell(row,col);value=str(cell.value or '').strip()
        if value.startswith(('https://','http://')):
            cell.hyperlink=None
            cell.hyperlink=value
        elif cell.hyperlink:raise ValueError('URL cell has a link but no explicit URL text: '+cell.coordinate)
