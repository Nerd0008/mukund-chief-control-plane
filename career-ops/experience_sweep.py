#!/usr/bin/env python3
"""Re-screen tracker rows for prior-experience requirements.

The discovery gate now reads the JD, but every row already in a tracker was
admitted by the old title-only check. This walks the live workbook and reports
what those rows actually ask for.

    python experience_sweep.py --region uk
    python experience_sweep.py --region uk --limit 20

Writes a JSON report next to the other runtime artefacts and prints a summary.
Read-only: it never edits a tracker.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import tracker_writer as tw  # noqa: E402
from experience_screening import experience_hits  # noqa: E402
from jd_fetcher import fetch_jd  # noqa: E402

FIRST_DATA_ROW = 3
LAST_COL = 6
RUNTIME = HERE.parent / 'runtime' / 'career-ops'


def read_rows(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[cfg_sheet()]
    out = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, LAST_COL + 1)]
        if not any(vals):
            continue
        out.append({'row': r, 'values': vals})
    wb.close()
    return out


def cfg_sheet():
    return 'Jobs'


def screen(url, timeout, cache):
    """Return (hits, readable) for one posting, using the shared cache."""
    if not isinstance(url, str) or not url.startswith(('http://', 'https://')):
        return [], False
    if url in cache:
        c = cache[url]
        return c.get('hits') or [], bool(c.get('readable'))
    try:
        jd = fetch_jd(url, timeout=timeout)
    except Exception:  # noqa: BLE001 - a failed fetch is data, not a crash
        jd = {}
    text = jd.get('text') if isinstance(jd, dict) else None
    hits = experience_hits(text) if isinstance(text, str) and text.strip() else []
    readable = isinstance(text, str) and bool(text.strip())
    cache[url] = {'hits': hits, 'readable': readable,
                  'checked_at': dt.datetime.now(dt.timezone.utc).isoformat()}
    return hits, readable


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--region', default='uk')
    ap.add_argument('--limit', type=int, default=0)
    ap.add_argument('--timeout', type=int, default=25)
    ap.add_argument('--cache', default=str(RUNTIME / 'experience-cache-uk.json'))
    ap.add_argument('--out', default=str(RUNTIME / 'experience-sweep-uk.json'))
    args = ap.parse_args()

    profiles = tw.load_profiles(None)
    cfg = tw.region_config(profiles, args.region)
    path = pathlib.Path(cfg['tracker'])

    cache_path = pathlib.Path(args.cache)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache = json.loads(cache_path.read_text(encoding='utf-8')) if cache_path.exists() else {}

    rows = read_rows(path)
    if args.limit:
        rows = rows[:args.limit]

    report = []
    for i, r in enumerate(rows, 1):
        company, title, url, status = r['values'][1], r['values'][2], r['values'][3], r['values'][5]
        hits, readable = screen(url, args.timeout, cache)
        report.append({'row': r['row'], 'company': company, 'title': title,
                       'url': url, 'status': status,
                       'needs_experience': bool(hits), 'hits': hits,
                       'screened': readable})
        if i % 10 == 0:
            print(f'  ... {i}/{len(rows)}', flush=True)
            cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding='utf-8')

    cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding='utf-8')
    pathlib.Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=1),
                                      encoding='utf-8')

    flagged = [r for r in report if r['needs_experience']]
    unscreened = [r for r in report if not r['screened']]
    print()
    print(f'rows screened      : {len(report)}')
    print(f'require experience : {len(flagged)}')
    print(f'could not screen   : {len(unscreened)}')
    print()
    for r in flagged:
        print(f"  row {r['row']:>3} | {str(r['status']):<12} | {str(r['company'])[:26]:<26} | {str(r['title'])[:38]}")
        if r['hits']:
            print(f"        -> {r['hits'][0][:110]}")
    print()
    print(f'written {args.out}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
