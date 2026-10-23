"""Extend the canonical regional trackers, never create a competing workbook.

Uses the repository's existing openpyxl backend. Email-owned columns are added
to the right of each existing layout; owner-owned cells are never overwritten.
"""
from __future__ import annotations

import copy
import contextlib
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile

import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter
from openpyxl.utils.cell import range_boundaries

from career_mail_parser import stable_id
from career_tracker_layout import resolve_profiles, validate_sheet

FIELDS = ["application_id", "application_date", "source_application_channel", "current_stage",
          "latest_status", "assessment_type", "assessment_provider", "assessment_received_date",
          "assessment_deadline", "deadline_timezone", "assessment_interview_link", "recruiter_contact",
          "gmail_message_id", "gmail_thread_id", "last_update_timestamp", "calendar_event_id",
          "confidence", "needs_review", "deadline_evidence", "state_json"]
HEADERS = {k: "Career Ops Email " + k.replace("_", " ").title() for k in FIELDS}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def posting_evidence(url):
    """One canonical projection for public posting identity (never signed queries)."""
    import re
    from urllib.parse import urlsplit,parse_qs
    parsed=urlsplit(str(url or ''))
    refs=re.findall(r"\b(?:WD\d{6,}|JR-\d{6,}|R-\d{6,}|R\d{6,}|SYS-\d{4,})\b",parsed.path,re.I)
    for key,values in parse_qs(parsed.query).items():
        if key.lower() in {'jobid','requisitionid','reqid'}:
            refs.extend(v for v in values if re.fullmatch(r'[A-Za-z0-9_-]{3,40}',v))
    return {'canonical_job_host':parsed.hostname or '', 'canonical_job_path':parsed.path,
            'job_reference_ids':sorted(set(refs))}


class WorkbookTracker:
    def __init__(self, profiles, backup_dir):
        self.profile = resolve_profiles(profiles)
        self.profiles, self.backup_dir = self.profile["regions"], Path(backup_dir)
        self.accepted_state = (Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / profiles["accepted_mail_state"]
                               if profiles.get("accepted_mail_state") else None)

    def read(self):
        records, fingerprints = [], {}
        for region, cfg in self.profiles.items():
            path = Path(cfg["tracker"])
            if not path.exists():
                continue
            fingerprints[region] = digest(path)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                ws = wb[cfg["sheet"]]
                validate_sheet(ws, cfg)
                headers = {str(c.value): c.column for c in ws[cfg["header_row"]] if c.value}
                state_column = headers.get(HEADERS["state_json"])
                fm = cfg["field_map"]
                for row, cells in enumerate(ws.iter_rows(min_row=cfg["first_data_row"], values_only=True), cfg["first_data_row"]):
                    cells = tuple(cells) + (None,) * max(0, max(ci(fm["company"]), ci(fm["title"]), ci(fm.get("url") or "A"), ci(cfg["status_columns"]["application_status"]), state_column or 0) - len(cells))
                    company = cells[ci(fm["company"]) - 1]
                    role = cells[ci(fm["title"]) - 1]
                    if not company or not role:
                        continue
                    saved = cells[state_column - 1] if state_column else None
                    record = json.loads(saved) if saved else {
                        "application_id": stable_id(company, role, "canonical:" + region + ":" + str(row)),
                        "company": str(company), "role": str(role), "seen_message_ids": [],
                        "thread_ids": [], "owner_confirmed_fields": ["company", "role"]}
                    record.update(region=region, row=row)
                    # Public posting references only; never retain URL queries or fragments.
                    url = str(cells[ci(fm["url"]) - 1] or "") if fm.get("url") else ""
                    record.update(posting_evidence(url))
                    record["canonical_owner_status"] = cells[ci(cfg["status_columns"]["application_status"]) - 1]
                    records.append(record)
            finally:
                wb.close()
        records=self._with_sidecars(records)
        if self.accepted_state:
            records = self._with_accepted_state(records)
        return records, fingerprints

    def _with_sidecars(self,records):
        for region,cfg in self.profiles.items():
            if not cfg.get('mail_metadata_sidecar'):continue
            path=Path(os.environ.get('LOCALAPPDATA',str(Path.home())))/cfg['mail_metadata_sidecar']
            if not path.exists():continue
            saved=json.loads(path.read_text(encoding='utf-8'))
            for old in saved.get('records',{}).values():
                matches=[r for r in records if r['region']==region and r['company'].casefold()==old['company'].casefold() and r['role'].casefold()==old['role'].casefold()
                    and r.get('canonical_job_host')==old.get('canonical_job_host') and r.get('canonical_job_path')==old.get('canonical_job_path')
                    and set(r.get('job_reference_ids',[]))==set(old.get('job_reference_ids',[]))]
                if len(matches)>1:raise ValueError('ambiguous mailbox sidecar identity')
                if not matches:continue  # Owner-deleted applications must never be resurrected.
                current=matches[0];protected={k:current.get(k) for k in ('company','role','region','row','canonical_owner_status','canonical_job_host','canonical_job_path','job_reference_ids')}
                current.update(copy.deepcopy(old));current.update(protected)
        return records

    def _upsert_sidecar(self,record,cfg,path):
        lock=path.with_suffix(path.suffix+'.career-mail.lock');fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY);temp=None
        try:
            before=digest(path); records,_=self.read()
            workbook_after=before
            if digest(path)!=before:raise ValueError("canonical workbook changed during identity read")
            store=Path(os.environ.get('LOCALAPPDATA',str(Path.home())))/cfg['mail_metadata_sidecar']
            metadata_before=digest(store) if store.exists() else None
            state=json.loads(store.read_text(encoding='utf-8')) if store.exists() else {'records':{}}
            candidates=[r for r in records if r['region']==record['region']
                        and r['company'].casefold()==record['company'].casefold()
                        and r['role'].casefold()==record['role'].casefold()]
            if len(candidates)>1:
                raise ValueError('ambiguous mailbox application identity; reconciliation required')
            if not candidates:
                if (record.get('owner_confirmed_application') is not True or record.get('row')
                        or record['application_id'] in state['records'] or record.get('needs_review')):
                    raise ValueError('new or owner-deleted mailbox application requires owner reconciliation; no row restored')
                evidence=posting_evidence(record.get('job_url') or record.get('url'))
                for key,value in evidence.items():
                    incoming=record.get(key)
                    if incoming and (set(incoming)!=set(value) if key=='job_reference_ids' else incoming!=value):
                        raise ValueError('new mailbox posting evidence requires owner reconciliation before insertion')
                # Shared writer, under the already-held shared lock. Missing URLs
                # are allowed only for this explicitly confirmed mailbox path.
                import tracker_writer as writer
                writer_profile={**self.profile, 'canonical_dir':self.profile.get('canonical_dir',str(path.parent)),
                                'career_ops_root':self.profile.get('career_ops_root',str(path.parent))}
                result=writer._write_records_unlocked(writer_profile,record['region'],[{
                    'company':record['company'],'title':record['role'],
                    'url':record.get('job_url') or record.get('url'),
                    'date_found':record.get('application_date') or record.get('last_update_timestamp'),
                    'owner_confirmed_application':True}],apply=True,confirmed_mail=True,
                    cleanup_existing=False,backup_dir=str(self.backup_dir))
                if result.get('error') or result['counts']['appended']!=1:
                    raise ValueError('confirmed mailbox application writer reconciliation failed')
                workbook_after=result['tracker_sha256_after']
                records,_=self.read()
                candidates=[r for r in records if r['region']==record['region']
                            and r['company'].casefold()==record['company'].casefold()
                            and r['role'].casefold()==record['role'].casefold()]
            if len(candidates)!=1:
                raise ValueError('ambiguous mailbox application identity; reconciliation required')
            current=candidates[0]
            if record.get('row') and record['row']!=current['row']:
                raise ValueError('canonical row changed identity; reconciliation required')
            for key in ('canonical_job_host','canonical_job_path','job_reference_ids'):
                if record.get(key) and record[key]!=current.get(key):
                    raise ValueError('conflicting mailbox posting evidence; reconciliation required')
            if (current.get('seen_message_ids') and current['application_id']!=record['application_id']):
                raise ValueError('conflicting mailbox application ID; reconciliation required')
            store.parent.mkdir(parents=True,exist_ok=True)
            saved=copy.deepcopy(record);saved.pop('row',None)
            for key in ('canonical_job_host','canonical_job_path','job_reference_ids'):saved[key]=current.get(key)
            state['records'][record['application_id']]=saved
            if (digest(store) if store.exists() else None)!=metadata_before:
                raise ValueError('metadata changed concurrently; sidecar write refused')
            if store.exists():
                self.backup_dir.mkdir(parents=True,exist_ok=True)
                backup=self.backup_dir/(store.name+'.'+metadata_before+'.json')
                shutil.copy2(store,backup)
                if digest(backup)!=metadata_before:raise ValueError('metadata backup verification failed')
            handle,name=tempfile.mkstemp(dir=store.parent,suffix='.tmp');os.close(handle);temp=Path(name)
            temp.write_text(json.dumps(state,sort_keys=True),encoding='utf-8');json.loads(temp.read_text(encoding='utf-8'))
            if digest(path)!=workbook_after:raise ValueError('canonical workbook changed concurrently; sidecar write refused')
            if (digest(store) if store.exists() else None)!=metadata_before:
                raise ValueError('metadata changed concurrently; sidecar write refused')
            os.replace(temp,store);record['row']=current['row']
            return {'region':record['region'],'row':current['row'],'before_sha256':before,'after_sha256':digest(path),'metadata_store':'private_sidecar','metadata_path':str(store)}
        finally:
            os.close(fd);lock.unlink()
            if temp and temp.exists():temp.unlink()

    def _with_accepted_state(self, records):
        """Recover dedupe evidence from the hash-verified, already-applied snapshot.

        This does not restore or write a workbook, acknowledge review-only mail,
        authorize new writes, or replace current owner fields.
        """
        root = self.accepted_state
        receipt = json.loads((root / "result.json").read_text())
        audit = json.loads((root / "apply-audit.json").read_text())
        if receipt.get("status") != "PASS" or audit.get("state") != "applied":
            raise ValueError("accepted mailbox state has no successful apply receipt")
        prior_profiles = {"regions": {}}
        for region, cfg in self.profiles.items():
            snapshot = root / "staged" / (region + ".xlsx")
            if not snapshot.exists():
                continue
            if digest(snapshot) != receipt.get("after_hashes", {}).get(region):
                raise ValueError("accepted mailbox snapshot fingerprint differs from receipt")
            prior_profiles["regions"][region] = {**cfg, "tracker": str(snapshot)}
        prior, _ = WorkbookTracker(prior_profiles, self.backup_dir).read()
        restored = 0
        for old in prior:
            if not old.get("seen_message_ids"):
                continue
            matches = [r for r in records if r["region"] == old["region"]
                       and r["company"].casefold() == old["company"].casefold()
                       and r["role"].casefold() == old["role"].casefold()
                       and r.get("canonical_job_host") == old.get("canonical_job_host")
                       and r.get("canonical_job_path") == old.get("canonical_job_path")
                       and set(r.get("job_reference_ids", [])) == set(old.get("job_reference_ids", []))]
            if len(matches) != 1:
                raise ValueError("accepted mailbox identity needs fresh owner reconciliation")
            current = matches[0]
            if current.get("seen_message_ids"):
                if current["application_id"] != old["application_id"]:
                    raise ValueError("accepted mailbox application identity conflicts with current metadata")
                # Current persisted metadata is stronger than the historical snapshot.
                current["seen_message_ids"] = sorted(set(current["seen_message_ids"]) | set(old["seen_message_ids"]))
                continue
            protected = {k: current.get(k) for k in ("company", "role", "region", "row", "canonical_owner_status", "job_reference_ids", "canonical_job_host", "canonical_job_path")}
            current.update(copy.deepcopy(old))
            current.update(protected)
            current["accepted_state_source"] = "verified previously applied snapshot"
            restored += 1
        self.accepted_state_recovered_records = restored
        return records

    def upsert(self, record):
        region = record.get("region")
        if region not in self.profiles:
            raise ValueError("application region is unresolved; no workbook write")
        cfg = self.profiles[region]
        path = Path(cfg["tracker"])
        if cfg.get('mail_metadata_sidecar'):
            return self._upsert_sidecar(record,cfg,path)
        # Owner mandate: the UK tracker is a fixed six-column sheet. Refuse to
        # append email metadata columns to it; that is what turned a readable
        # tracker into a 31-column log. Non-UK regions keep their own layouts.
        if cfg.get("column_ceiling"):
            raise ValueError(
                f"{region} tracker is capped at {cfg['column_ceiling']} columns; "
                "email metadata columns are not permitted"
            )
        # A fail-fast file lock excludes our own overlapping scheduled scans.
        lock = path.with_suffix(path.suffix + ".career-mail.lock")
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        temp = None
        try:
            before = digest(path)
            wb = openpyxl.load_workbook(path)
            try:
                ws = wb[cfg["sheet"]]
                validate_sheet(ws, cfg)
                headers = {str(c.value): c.column for c in ws[cfg["header_row"]] if c.value}
                identity = cfg.get("id") or cfg.get("mail_identity")
                if cfg.get("mail_identity"):
                    spec = cfg["mail_identity"]
                    cell = ws.cell(cfg["header_row"], ci(spec["column"]))
                    if cell.value not in (None, spec["header"]):
                        raise ValueError("reserved email identity column is occupied")
                    cell.value = spec["header"]
                cols = {}
                for key, header in HEADERS.items():
                    if header not in headers:
                        column = ws.max_column + 1
                        ws.cell(cfg["header_row"], column, header)
                        headers[header] = column
                    cols[key] = headers[header]
                row = record.get("row")
                # Recover prior write if the local checkpoint was never saved.
                for n in range(cfg["first_data_row"], ws.max_row + 1):
                    if ws.cell(n, cols["application_id"]).value == record["application_id"]:
                        row = n
                        break
                if row:
                    fm = cfg["field_map"]
                    if (str(ws.cell(row, ci(fm["company"])).value or "").casefold() != record["company"].casefold()
                            or str(ws.cell(row, ci(fm["title"])).value or "").casefold() != record["role"].casefold()):
                        raise ValueError("canonical row changed identity; reconciliation required")
                else:
                    id_col = ci(identity["column"]) if identity else None
                    occupied = [n for n in range(cfg["first_data_row"], ws.max_row + 1)
                                if ws.cell(n, ci(cfg["field_map"]["company"])).value]
                    row = max(occupied, default=cfg["first_data_row"] - 1) + 1
                    if row > cfg.get("license_free_rows", 1000):
                        raise ValueError("canonical tracker capacity reached")
                    style = identity.get("style") if identity else None
                    if style == "sequential":
                        prefix = identity["prefix"]
                        import re
                        nums = [int(m.group(1)) for n in occupied if
                                (m := re.fullmatch(re.escape(prefix) + r"(\d+)", str(ws.cell(n, id_col).value or "")))]
                        job_id = prefix + str(max(nums, default=0) + 1)
                    else:
                        job_id = (identity or {}).get("prefix", region.upper()) + "-MAIL-" + record["application_id"][:12]
                    if id_col:
                        ws.cell(row, id_col, job_id)
                    for key, value in (("company", record["company"]), ("title", record["role"])):
                        cell = ws.cell(row, ci(cfg["field_map"][key]), value)
                        cell.data_type = "s"
                    # Existing validation vocabulary, never invent Assessment in owner status.
                    status_col = cfg["status_columns"]["application_status"]
                    ws.cell(row, ci(status_col), cfg.get("defaults", {}).get("application_status"))
                    if cfg.get("table") in ws.tables:
                        table = ws.tables[cfg["table"]]
                        left, top, right, bottom = range_boundaries(table.ref)
                        table.ref = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{max(bottom, row)}"
                saved = {k: v for k, v in record.items() if k != "row"}
                deadline = saved.get("deadline") or {}
                values = {**saved, "assessment_deadline": deadline.get("value"),
                          "deadline_timezone": deadline.get("timezone"),
                          "deadline_evidence": json.dumps(deadline), "state_json": json.dumps(saved, sort_keys=True)}
                for key, column in cols.items():
                    value = values.get(key)
                    cell = ws.cell(row, column)
                    cell.value = value
                    if isinstance(value, str):
                        if len(value) > 32767:
                            raise ValueError("application metadata exceeds Excel cell capacity")
                        cell.data_type = "s"  # email text cannot become a formula
                self.backup_dir.mkdir(parents=True, exist_ok=True)
                backup = self.backup_dir / f"{path.stem}-{before}.xlsx"
                if not backup.exists():
                    shutil.copy2(path, backup)
                if digest(backup) != before:
                    raise ValueError("workbook backup verification failed")
                handle, name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
                os.close(handle)
                temp = Path(name)
                wb.save(temp)
            finally:
                wb.close()
            verify = openpyxl.load_workbook(temp, read_only=True)
            try:
                if verify[cfg["sheet"]].cell(row, cols["application_id"]).value != record["application_id"]:
                    raise ValueError("saved workbook verification failed")
            finally:
                verify.close()
            if digest(path) != before:
                raise ValueError("workbook changed concurrently; write refused")
            os.replace(temp, path)
            record["row"] = row
            return {"region": region, "row": row, "before_sha256": before, "after_sha256": digest(path)}
        finally:
            os.close(fd)
            if temp and temp.exists():
                temp.unlink()
            lock.unlink()
