"""Validate configured layouts; regional_profiles.json alone defines columns."""
import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci

def validate_sheet(ws, cfg):
    if cfg.get("column_ceiling") and ws.max_column > cfg["column_ceiling"]:
        raise ValueError("canonical tracker exceeds fixed column ceiling")
    for column, expected in cfg.get("required_headers", {}).items():
        if ws.cell(cfg["header_row"], ci(column)).value != expected:
            raise ValueError("canonical tracker layout unrecognized: headers differ from the authoritative profile")

def resolve_profiles(profiles):
    result = copy.deepcopy(profiles)
    for cfg in result["regions"].values():
        if cfg.get("mail_metadata_sidecar") and not cfg.get("column_ceiling"):
            raise ValueError("sidecar tracker requires an explicit fixed schema ceiling")
        if cfg.get("column_ceiling") and cfg.get("mail_identity"):
            raise ValueError("fixed tracker schema cannot reserve an email identity column")
        path = Path(cfg["tracker"])
        if path.exists() and cfg.get("required_headers"):
            wb = openpyxl.load_workbook(path, read_only=True)
            try:
                validate_sheet(wb[cfg["sheet"]], cfg)
            finally:
                wb.close()
    return result

def normalize_status(value, cfg):
    return cfg.get("status_aliases", {}).get(value, value)
